"""Single Receipt Report -- Excel export (openpyxl). Mirrors the VFP Excel macro:
CHALLAN NO, DATE, ITEM DESC, UNIT, QTY, INV. RATE."""
from io import BytesIO


def _date(iso):
    if not iso:
        return None
    from datetime import date
    y, m, d = iso.split("-")
    return date(int(y), int(m), int(d))


SINGLE_RECEIPT_REPORT_COLUMNS = [
    ("CHALLAN NO", 7, "@"),
    ("DATE", 8, "dd/mm/yy"),
    ("ITEM DESC", 30, "@"),
    ("UNIT", 4, "@"),
    ("QTY", 10, "0.000"),
    ("INV. RATE", 10, "0.00"),
]


def build_single_receipt_report_workbook(rows, company_name, start_date, end_date):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side
    from openpyxl.utils import get_column_letter

    font_name = "Arial"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Single Receipt Report"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=8, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"SINGLE RECEIPT REPORT FROM DATE {start_date.strftime('%d/%m/%Y')} TO {end_date.strftime('%d/%m/%Y')}",
    ).font = Font(name=font_name, size=8, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(SINGLE_RECEIPT_REPORT_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=8, bold=True)
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for line in rows:
        values = [line["challan_no"], _date(line["date"]), line["item_name"], line["unit"],
                  float(line["qty"]), float(line["inv_rate"])]
        for col, (value, (_h, _w, fmt)) in enumerate(zip(values, SINGLE_RECEIPT_REPORT_COLUMNS), start=1):
            cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
            cell.font = Font(name=font_name, size=8)
            cell.number_format = fmt
        row += 1

    last_row = max(row - 1, header_row)
    last_col = get_column_letter(len(SINGLE_RECEIPT_REPORT_COLUMNS))
    for r in range(header_row, row):
        for c in range(1, len(SINGLE_RECEIPT_REPORT_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
