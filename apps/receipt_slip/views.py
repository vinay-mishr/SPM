import json
from datetime import date, datetime

from django.contrib.auth import authenticate
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import JsonResponse
from django.views.decorators.http import require_GET, require_POST
from django.views.generic import TemplateView

from . import services


def _can_edit(user):
    return services.writes_enabled() and user.has_perm(services.PERMISSION)


def _forbidden():
    if not services.writes_enabled():
        return JsonResponse({"message": "Saving to the SP_MACHINE database is switched off. Nothing was changed."}, status=403)
    return JsonResponse({"message": "You don't have permission to change receipt slips."}, status=403)


def _parse_date(raw, default=None):
    try:
        return datetime.strptime(raw, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return default


class StockReceiptView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/stock_receipt.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["can_edit"] = _can_edit(self.request.user)
        context["form_title"] = "STOCK RECEIVE"
        context["initial_data"] = {
            "rows": services.list_receipt_slips(today, today),
            "start_date": today.isoformat(), "end_date": today.isoformat(),
        }
        return context


class ReceiptSlipEntryView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/receipt_slip_entry.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["can_edit"] = _can_edit(self.request.user)
        context["initial_data"] = {
            "rows": services.list_receipt_slips(today, today),
            "start_date": today.isoformat(), "end_date": today.isoformat(),
        }
        return context


@login_required
@require_GET
def receipt_slip_list(request):
    today = date.today()
    start = _parse_date(request.GET.get("start_date"), today)
    end = _parse_date(request.GET.get("end_date"), today)
    return JsonResponse({"rows": services.list_receipt_slips(start, end)})


@login_required
@require_GET
def receipt_slip_next_no(request):
    return JsonResponse({"mis_no": services.next_receipt_no()})


@login_required
@require_GET
def receipt_slip_detail(request, mis_no):
    order = services.get_receipt_slip(mis_no)
    if not order:
        return JsonResponse({"message": "Receipt not found."}, status=404)
    return JsonResponse(order)


@login_required
@require_POST
def reauthenticate(request):
    """Re-checks the CURRENT user's own account password before Modify/Delete."""
    try:
        payload = json.loads(request.body or "{}")
    except ValueError:
        payload = {}
    user = authenticate(request, username=request.user.get_username(), password=payload.get("password") or "")
    if user is None:
        return JsonResponse({"ok": False, "message": "Invalid password."}, status=403)
    return JsonResponse({"ok": True})


@login_required
@require_POST
def receipt_slip_save(request):
    if not _can_edit(request.user):
        return _forbidden()
    try:
        payload = json.loads(request.body or b"{}")
    except ValueError:
        return JsonResponse({"message": "Invalid request."}, status=400)
    mis_date = None
    try:
        y, m, d = (payload.get("mis_date") or "").split("-")
        mis_date = date(int(y), int(m), int(d))
    except (ValueError, AttributeError):
        pass
    try:
        result = services.save_receipt_slip(
            mis_date=mis_date,
            items=payload.get("items"),
            remarks=payload.get("remarks", ""),
            key=payload.get("key") or None,
            username=request.user.get_username()[:10],
            chl_no=payload.get("mis_no") or "",
        )
    except services.ReceiptSlipError as exc:
        return JsonResponse({"message": str(exc), "errors": exc.errors}, status=400)
    return JsonResponse({"message": f"Receipt {result['mis_no']} saved.", "mis_no": result["mis_no"]})


@login_required
@require_POST
def receipt_slip_delete(request):
    if not _can_edit(request.user):
        return _forbidden()
    mis_no = request.POST.get("mis_no", "")
    try:
        services.delete_receipt_slip(mis_no)
    except services.ReceiptSlipError as exc:
        return JsonResponse({"message": str(exc)}, status=400)
    return JsonResponse({"message": f"Receipt {mis_no} deleted."})


# --- Single Receipt Report ----------------------------------------------------------

class SingleReceiptReportView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/single_receipt_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {
            "rows": services.single_receipt_report(today, today),
            "start_date": today.isoformat(), "end_date": today.isoformat(),
        }
        return context


def _report_params(request):
    today = date.today()
    return _parse_date(request.GET.get("start_date"), today), _parse_date(request.GET.get("end_date"), today)


@login_required
@require_GET
def single_receipt_report_data(request):
    start, end = _report_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse({"rows": services.single_receipt_report(start, end)})


@login_required
@require_GET
def single_receipt_report_export(request):
    from django.conf import settings
    from django.http import HttpResponse
    from . import exports

    start, end = _report_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    rows = services.single_receipt_report(start, end)
    data = exports.build_single_receipt_report_workbook(rows, settings.BRAND["name"].upper(), start, end)
    response = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    filename = f"Single_Receipt_Report_{start.strftime('%d-%m-%Y')}_to_{end.strftime('%d-%m-%Y')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


class StockReportView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/stock_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def stock_report_data(request):
    start, end = _report_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse({"rows": services.stock_report(start, end)})


@login_required
@require_GET
def stock_report_export(request):
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from io import BytesIO

    start, end = _report_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append(["STOCK RECEIPT REPORT"])
    ws.append([f"FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["SNO", "BILL NO", "BILL DATE", "ITEM NAME", "QTY"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 12, 12, 40, 12], "ABCDE"):
        ws.column_dimensions[col].width = w
    for r in services.stock_report(start, end):
        ws.append([r["sno"], r["bill_no"], r["date"], r["item_name"], float(r["qty"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Stock_Receipt_Report_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response


class CurrentStockView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/current_stock.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


@login_required
@require_GET
def current_stock_data(request):
    return JsonResponse({"rows": services.current_stock_report(_parse_date(request.GET.get("run_date"), date.today()))})


@login_required
@require_GET
def current_stock_export(request):
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from io import BytesIO

    today = _parse_date(request.GET.get("run_date"), date.today())
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"CURRENT STOCK AS ON DATE {today:%d/%m/%Y}"])
    ws.append([])
    ws.append(["CODE", "ITEM NAME", "QTY."])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([14, 40, 12], "ABC"):
        ws.column_dimensions[col].width = w
    for r in services.current_stock_report(today):
        ws.append([r["code"], r["name"], float(r["qty"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Current_Stock_{today:%d-%m-%Y}.xlsx"'
    return response


class StockLedgerView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/stock_ledger.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        fy_start = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
        context["initial_data"] = {"start_date": fy_start.isoformat(), "end_date": today.isoformat()}
        return context


class ProductionScheduleView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/production_schedule.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        fy_start = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
        context["initial_data"] = {"start_date": fy_start.isoformat(), "end_date": today.isoformat(), "days": 30}
        return context


def _schedule_params(request):
    from . import production_schedule
    start, end = _report_params(request)
    try:
        days = int(request.GET.get("days") or 30)
    except ValueError:
        days = 30
    return start, end, days


@login_required
@require_GET
def production_schedule_data(request):
    from . import production_schedule
    start, end, days = _schedule_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse(production_schedule.production_schedule(start, end, days))


@login_required
@require_GET
def production_schedule_export(request):
    from django.conf import settings
    from django.http import HttpResponse
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from . import production_schedule

    start, end, days = _schedule_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    data = production_schedule.production_schedule(start, end, days)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"Production Schedule Based Upon Inventory for {days} Days"])
    ws.append([])
    headers = ["SNO", "CODE", "PRODUCT DESCRIPTION", "PENDING SALE ORD", f"AVG CURR. SALE FOR {days} DAYS",
               f"AVG LAST YEAR SALE FOR {days} DAYS", "QTY IN PROCESS", "CLOSING STOCK",
               "QTY REQD. TO BE PRODUCED", "MAN HRS REQD.", "MAN HRS PER UNIT"]
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 12, 36, 14, 16, 18, 14, 14, 18, 14, 14], "ABCDEFGHIJK"):
        ws.column_dimensions[col].width = w
    for r in data["rows"]:
        if r["pending_only"]:
            ws.append([r["sno"], r["code"], r["name"], float(r["pending"])])
        else:
            ws.append([r["sno"], r["code"], r["name"], float(r["pending"]), float(r["cur_sale"]),
                       float(r["last_sale"]), float(r["in_process"]), float(r["closing"]), float(r["required"]),
                       float(r["man_hrs"]), float(r["time_per_unit"])])
    ws.append(["", "Total", "", None, None, None, None, None, None, float(data["total_man_hrs"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Production_Schedule_{end:%d-%m-%Y}.xlsx"'
    return response


class ProductionPendingView(LoginRequiredMixin, TemplateView):
    template_name = "receipt_slip/production_schedule.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        fy_start = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
        context["pending_mode"] = True
        context["initial_data"] = {"start_date": fy_start.isoformat(), "end_date": today.isoformat(), "days": 30}
        return context


@login_required
@require_GET
def production_pending_data(request):
    from . import production_schedule
    start, end = _report_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse(production_schedule.production_pending(start, end))


@login_required
@require_GET
def production_pending_export(request):
    from django.conf import settings
    from django.http import HttpResponse
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from . import production_schedule

    start, end = _report_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    data = production_schedule.production_pending(start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"Production Schedule Based Upon Pending Sales Order Requiring Execution Upto {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["SNO", "CODE", "PRODUCT DESCRIPTION", f"PENDING SALE ORDER UPTO {end:%d/%m/%Y}", "QTY IN PROCESS",
               "CLOSING STOCK", "QTY REQD. TO BE PRODUCED", "MAN HRS REQD.", "MAN HRS PER UNIT"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 12, 36, 18, 14, 14, 18, 14, 14], "ABCDEFGHI"):
        ws.column_dimensions[col].width = w
    for r in data["rows"]:
        if r["pending_only"]:
            ws.append([r["sno"], r["code"], r["name"], float(r["pending"])])
        else:
            ws.append([r["sno"], r["code"], r["name"], float(r["pending"]), float(r["in_process"]),
                       float(r["closing"]), float(r["required"]), float(r["man_hrs"]), float(r["time_per_unit"])])
    ws.append(["", "", "Total", None, None, None, None, float(data["total_man_hrs"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Production_Schedule_Pending_{end:%d-%m-%Y}.xlsx"'
    return response
