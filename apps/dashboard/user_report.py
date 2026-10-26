"""Reports > Sales Reports > User Wise Report (VFP "USER WISE REPORT" form,
OLE-Excel export block).

Three independent sections for ONE required user (``uname``, the 10-char
login name every save routine in this project already stamps -- see e.g.
apps.gst_invoice.services' own ``username[:10]``), each scoped to the same
date range but reading a different table/date column:
  * PURCHASE: ``mrn_dat`` rows (``b_date``).
  * SALE BILL: ``bill`` rows (``b_date``).
  * ACCOUNT HEAD: ``trans`` rows (``tdate``) -- ledger postings.

The VFP source's own ALL/SINGLE toggle (``tsa``) is NOT actually load-
bearing: every one of its three ``SCAN FOR`` clauses hard-codes
``ALLTRIM(uname) == ALLTRIM(tuser.value)`` regardless of ``tsa``, and the
one filter block that reads ``tsa = 'S'`` sets a VFP ``SET FILTER`` on
whatever work area happens to be current at that point (not clearly
``mrn_dat``) and is never read back by the actual ``SCAN`` that follows --
dead code. So a user is always required here, matching what the three
scans actually do; "ALL" is not reproduced as a real option.

The VFP source's own password gate (``DO FORM password`` against
``office.mpassword``, before the form even opens) is not reproduced --
this project's own login (``LoginRequiredMixin``, same as every other
report) already gates access to this view.

ACCOUNT HEAD's own AMOUNT column (G) is a bug in the VFP source, preserved
exactly: ``tamt`` is written to every row's cell and to the final TOTAL,
but is NEVER incremented anywhere in this section's own scan (the one line
that would increment it, ``tamt=tamt+n_amount``, is commented out) -- so
every Account Head amount, including the total, is always 0.
"""
from datetime import date

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import connections
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView


def _rows(sql, params=None):
    with connections["legacy"].cursor() as cur:
        cur.execute(sql, params or [])
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]


def _clean(value):
    return (value or "").strip() if isinstance(value, str) or value is None else value


def _d(value):
    try:
        return float(value) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def user_report_data(user, start, end):
    purchase_rows = _rows(
        "SELECT m.bill_no, m.b_date, m.n_amount, c.cname, c.cadd3 FROM mrn_dat m "
        "LEFT JOIN cust_mst c ON c.ccode = m.p_code "
        "WHERE m.b_date >= %s AND m.b_date <= %s AND LTRIM(RTRIM(m.uname)) = %s ORDER BY m.b_date",
        [start, end, user],
    )
    sale_rows = _rows(
        "SELECT b.bill_no, b.b_date, b.amount, c.cname, c.cadd3 FROM bill b "
        "LEFT JOIN cust_mst c ON c.ccode = b.p_code "
        "WHERE b.b_date >= %s AND b.b_date <= %s AND LTRIM(RTRIM(b.uname)) = %s ORDER BY b.b_date",
        [start, end, user],
    )
    trans_rows = _rows(
        "SELECT t.tdoc, t.tdate, t.tnar1, t.tcode, t.tccode FROM trans t "
        "WHERE t.tdate >= %s AND t.tdate <= %s AND LTRIM(RTRIM(t.uname)) = %s ORDER BY t.tdate",
        [start, end, user],
    )

    ledger_codes = {_clean(r["tcode"]) for r in trans_rows if _clean(r["tcode"])}
    ledgers = {_clean(r["gcode"]): _clean(r["gdesc"]) for r in _rows(
        "SELECT gcode, gdesc FROM gled_mst WHERE gcode IN (%s)" % ",".join(["%s"] * len(ledger_codes)),
        list(ledger_codes),
    )} if ledger_codes else {}
    party_codes = {_clean(r["tccode"]) for r in trans_rows if _clean(r["tccode"])}
    parties = {_clean(r["ccode"]): _clean(r["cname"]) for r in _rows(
        "SELECT ccode, cname FROM cust_mst WHERE ccode IN (%s)" % ",".join(["%s"] * len(party_codes)),
        list(party_codes),
    )} if party_codes else {}

    purchase = [{
        "sno": i + 1, "bill_no": _clean(r["bill_no"]) or "", "bill_date": r["b_date"].isoformat() if r["b_date"] else "",
        "party_name": _clean(r["cname"]) or "", "city": _clean(r["cadd3"]) or "", "amount": f"{_d(r['n_amount']):.2f}",
    } for i, r in enumerate(purchase_rows)]
    sale = [{
        "sno": i + 1, "bill_no": _clean(r["bill_no"]) or "", "bill_date": r["b_date"].isoformat() if r["b_date"] else "",
        "party_name": _clean(r["cname"]) or "", "city": _clean(r["cadd3"]) or "", "amount": f"{_d(r['amount']):.2f}",
    } for i, r in enumerate(sale_rows)]
    account_head = [{
        "v_no": _clean(r["tdoc"]) or "", "ledger": ledgers.get(_clean(r["tcode"]), "") or "",
        "vou_date": r["tdate"].isoformat() if r["tdate"] else "", "narration": _clean(r["tnar1"]) or "",
        "party": parties.get(_clean(r["tccode"]), "") or "",
    } for r in trans_rows]

    return {
        "purchase": purchase, "purchase_total": f"{sum(float(r['amount']) for r in purchase):.2f}",
        "sale": sale, "sale_total": f"{sum(float(r['amount']) for r in sale):.2f}",
        "account_head": account_head,
    }


class UserReportView(LoginRequiredMixin, TemplateView):
    template_name = "dashboard/user_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    user = (request.GET.get("user") or "").strip()[:10]
    return user, start, end


@login_required
@require_GET
def ur_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    user, start, end = _params(request)
    if not user:
        return JsonResponse({"message": "Enter a user."}, status=400)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)

    data = user_report_data(user, start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append([f"USERWISE REPORT FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])

    ws.append(["PURCHASE"])
    ws["A" + str(ws.max_row)].font = Font(bold=True)
    header = ["User", "SNO", "BLL NO", "BILL DATE", "PARTY NAME", "CITY", "AMOUNT"]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    fl = False
    for r in data["purchase"]:
        ws.append([
            user if not fl else "", r["sno"], r["bill_no"],
            date.fromisoformat(r["bill_date"]) if r["bill_date"] else "", r["party_name"], r["city"],
            float(r["amount"]),
        ])
        fl = True
    ws.append(["", "", "", "", "TOTAL", "", float(data["purchase_total"])])
    ws.cell(row=ws.max_row, column=5).font = Font(bold=True)

    ws.append([])
    ws.append(["SALE BILL"])
    ws["A" + str(ws.max_row)].font = Font(bold=True)
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    fl = False
    for r in data["sale"]:
        ws.append([
            user if not fl else "", r["sno"], r["bill_no"],
            date.fromisoformat(r["bill_date"]) if r["bill_date"] else "", r["party_name"], r["city"],
            float(r["amount"]),
        ])
        fl = True
    ws.append(["", "", "", "", "TOTAL", "", float(data["sale_total"])])
    ws.cell(row=ws.max_row, column=5).font = Font(bold=True)

    ws.append([])
    ws.append(["ACCOUNT HEAD"])
    ws["A" + str(ws.max_row)].font = Font(bold=True)
    ws.append(["User", "V. NO", "LEDGER", "VOU DATE", "NARRATION", "PARTY", "AMOUNT"])
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    fl = False
    for r in data["account_head"]:
        ws.append([
            user if not fl else "", r["v_no"], r["ledger"],
            date.fromisoformat(r["vou_date"]) if r["vou_date"] else "", r["narration"], r["party"], 0,
        ])
        fl = True
    ws.append(["", "", "", "", "TOTAL", "", 0])
    ws.cell(row=ws.max_row, column=5).font = Font(bold=True)

    for col, w in zip("ABCDEFG", [10, 5, 15, 12, 40, 20, 12]):
        ws.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="User_Report_{user}_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
