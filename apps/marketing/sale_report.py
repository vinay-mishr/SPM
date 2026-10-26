"""Reports > Sales Reports > Mktg Managerwise Sale (VFP "MKTG MANAGERWISE
SALE" form, OLE-Excel export block).

Same shape as apps.reps.sale_report.rep_wise_sale_groups -- see that
module's own docstring for the shared rep/state/bill grouping and rolling
totals -- grouped Marketing Manager (``cust_mst.cmm``, ALL or one via the
``tpcode``/``mmhelp`` picker) then STATE (``cust_mst.cstate``) then bill,
instead of rep. Two differences from Rep Wise Sale, both preserved from
this pasted source rather than reconciled:
  * No ``bill_type = 'G'`` filter anywhere -- every bill in the date range
    counts, not just GST bills (apps.reps.sale_report's own newer OLE
    block added that filter; this one never did).
  * The manager for a bill comes from ``cust_mst.cmm`` (the party's own
    stored marketing manager), not ``bill.rep_cod`` like Rep Wise Sale's
    own newer block reads straight off the bill -- this report never
    picked up that same fix.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import connections
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView


def _rows(sql, params=None):
    with connections["legacy"].cursor() as cur:
        cur.execute(sql, params or [])
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]


def _clean(value):
    return (value or "").strip() if isinstance(value, str) or value is None else value


def _d(value):
    try:
        return Decimal(str(value if value is not None else 0))
    except Exception:
        return Decimal("0")


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


def mm_wise_sale_groups(start, end, mm_code=None):
    sql = (
        "SELECT b.div, b.bill_no, b.b_date, b.amount, c.cmm, c.cstate, c.cname, c.cadd3 "
        "FROM bill b LEFT JOIN cust_mst c ON c.ccode = b.p_code "
        "WHERE b.b_date >= %s AND b.b_date <= %s"
    )
    params = [start, end]
    if mm_code:
        sql += " AND c.cmm = %s"
        params.append(mm_code)
    sql += " ORDER BY c.cmm, c.cstate, b.b_date"
    bills = _rows(sql, params)

    mm_names = {_clean(r["mcode"]): _clean(r["mname"]) for r in _rows("SELECT mcode, mname FROM mm_mst")}
    state_names = {_clean(r["stcode"]): _clean(r["stname"]) for r in _rows("SELECT stcode, stname FROM state")}

    managers = []
    grand_total = Decimal("0")
    current_mm = None
    for b in bills:
        mm_code_val = _clean(b["cmm"]) or ""
        state_code_val = _clean(b["cstate"]) or ""
        if current_mm is None or current_mm["code"] != mm_code_val:
            current_mm = {
                "code": mm_code_val, "name": mm_names.get(mm_code_val, "") or mm_code_val,
                "states": [], "total": Decimal("0"),
            }
            managers.append(current_mm)
            current_state = None
        if current_state is None or current_state["code"] != state_code_val:
            current_state = {
                "code": state_code_val, "name": state_names.get(state_code_val, "") or state_code_val,
                "bills": [], "total": Decimal("0"),
            }
            current_mm["states"].append(current_state)

        amount = _d(b["amount"])
        current_state["bills"].append({
            "invoice_no": _clean(b["bill_no"]) or "",
            "invoice_date": b["b_date"].isoformat() if b["b_date"] else "",
            "party_name": _clean(b["cname"]) or "", "party_city": _clean(b["cadd3"]) or "",
            "amount": f"{amount:.2f}",
        })
        current_state["total"] += amount
        current_mm["total"] += amount
        grand_total += amount

    for m in managers:
        m["total"] = f"{m['total']:.2f}"
        for s in m["states"]:
            s["total"] = f"{s['total']:.2f}"
    return managers, f"{grand_total:.2f}"


class MmSaleReportView(LoginRequiredMixin, TemplateView):
    template_name = "marketing/sale_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": _fy_start(today).isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), _fy_start(today))
    end = parse(request.GET.get("end_date"), today)
    scope = (request.GET.get("scope") or "all").strip().lower()
    mm_code = (request.GET.get("mm_code") or "").strip().upper() if scope == "single" else None
    return start, end, scope, mm_code


@login_required
@require_GET
def mms_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end, scope, mm_code = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    if scope == "single" and not mm_code:
        return JsonResponse({"message": "Select a marketing manager first."}, status=400)

    managers, grand_total = mm_wise_sale_groups(start, end, mm_code)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([settings.BRAND.get("address_line1", "")])
    ws.append([f"MKTG MANAGER WISE SALE FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])

    header = ["INV.NO", "INV.DATE", "PARTY NAME", "PARTY CITY", "SALE AMT"]
    widths = [20, 10, 30, 20, 12]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip(widths, "ABCDE"):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["B"].number_format = "dd/mm/yy"
    ws.column_dimensions["E"].number_format = "0.00"

    for m in managers:
        ws.append([m["name"]])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for s in m["states"]:
            ws.append([s["name"]])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            for bi in s["bills"]:
                ws.append([
                    bi["invoice_no"], date.fromisoformat(bi["invoice_date"]) if bi["invoice_date"] else None,
                    bi["party_name"], bi["party_city"], float(bi["amount"]),
                ])
                ws.cell(row=ws.max_row, column=2).number_format = "dd/mm/yy"
                ws.cell(row=ws.max_row, column=5).number_format = "0.00"
            ws.append(["", "", "", "TOTAL", float(s["total"])])
            ws.cell(row=ws.max_row, column=4).font = Font(bold=True)
        ws.append(["", "", "", "MM. TOTAL", float(m["total"])])
        ws.cell(row=ws.max_row, column=4).font = Font(bold=True)

    ws.append(["", "", "", "GRAND TOTAL", float(grand_total)])
    ws.cell(row=ws.max_row, column=4).font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Mktg_Managerwise_Sale_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
