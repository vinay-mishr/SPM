"""Part Master Excel export (EXCEL button). Read-only: current part-mst records."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import services
from .models import Part

FONT_NAME = "Arial"
COLUMNS = [
    ("Part Number", "part_number", 16, "@"),
    ("Part Description 1", "description_1", 30, "@"),
    ("Part Description 2", "description_2", 30, "@"),
    ("Quantity in Order", "quantity_in_order", 14, "0.000"),
    ("Drawing Number", "drawing_no", 14, "@"),
    ("Gauge", "gauge", 10, "@"),
    ("Packing", "packing", 10, "0"),
    ("Grade", "grade", 8, "@"),
    ("Gross Weight", "gross_weight", 12, "0.0000"),
    ("Purchase Rate", "purchase_rate", 12, "0.0000"),
    ("Sale Rate", "sale_rate", 12, "0.00"),
    ("Opening Stock", "opening_stock", 12, "0.000"),
    ("Reorder Quantity", "reorder_qty", 14, "0.000"),
    ("Material", "material", 14, "@"),
    ("Weight by 100", "weight_by_100", 12, "0.000"),
    ("Current Stock", "current_stock", 12, "0"),
    ("Location", "location", 10, "@"),
    ("Unit", "unit", 8, "@"),
    ("HSN Code", "hsn_code", 12, "@"),
    ("GST %", "gst_percent", 8, "0.00"),
    ("Part Type", "part_type", 10, "@"),
    ("Sub-Assembly", "sub_assembly", 12, "@"),
    ("Supplier 1", "supplier_1", 10, "@"),
    ("Rate 1", "rate_1", 12, "0.0000"),
    ("Supplier 2", "supplier_2", 10, "@"),
    ("Rate 2", "rate_2", 12, "0.0000"),
    ("Supplier 3", "supplier_3", 10, "@"),
    ("Rate 3", "rate_3", 12, "0.0000"),
    ("Supplier 4", "supplier_4", 10, "@"),
    ("Rate 4", "rate_4", 12, "0.0000"),
]
PART_TYPE_LABELS = {"R": "RM", "C": "CP", "S": "CS", "": ""}


def _sanitize(value):
    """Old data sometimes has stray control characters Excel can't store; strip them."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def build_workbook():
    """All part-mst records, as the Django Part Master screen shows them."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Part Master"

    for col, (heading, _field, width, _fmt) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=heading)
        cell.font = Font(name=FONT_NAME, size=10, bold=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    row = 2
    for part in Part.objects.order_by("part_number"):
        data = services.serialize(part)
        data["sub_assembly"] = "Y" if data["sub_assembly"] else ""
        data["part_type"] = PART_TYPE_LABELS.get(data["part_type"], data["part_type"])
        for col, (_heading, field, _width, fmt) in enumerate(COLUMNS, start=1):
            value = _sanitize(data[field])
            if fmt != "@" and value not in ("", None):
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    pass
            cell = ws.cell(row=row, column=col, value=value if value != "" else None)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.number_format = fmt
        row += 1

    ws.freeze_panes = "A2"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
