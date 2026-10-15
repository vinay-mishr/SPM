from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.parties.models import Party

from . import exports, services
from .models import Part


class PartMasterTests(TestCase):
    """part-mst and cust_mst are the old software's tables (unmanaged); tests create them."""

    @classmethod
    def setUpClass(cls):
        with connection.schema_editor() as editor:
            editor.create_model(Party)
        with connection.cursor() as cursor:
            # Like the real part-mst: no primary key / unique constraint (part numbers repeat).
            cursor.execute(
                'CREATE TABLE "part-mst" (icod varchar(10), ides varchar(35), rname1 varchar(35), '
                "qnty decimal(11,3), drgno varchar(10), gauge varchar(10), std_pack decimal(5,0), "
                "grade varchar(2), g_wt decimal(10,4), rate decimal(10,4), srate decimal(10,2), "
                "op_stk decimal(12,3), rop decimal(11,3), material varchar(15), wt decimal(7,3), "
                "cur_stk decimal(10,0), loc varchar(10), uom varchar(3), hs_code varchar(8), "
                "gs_rt decimal(5,2), s_cod varchar(1), item_type varchar(1), party1 varchar(4), "
                "rate1 decimal(10,4), party2 varchar(4), rate2 decimal(10,4), party3 varchar(4), "
                "rate3 decimal(10,4), party4 varchar(4), rate4 decimal(10,4), uname varchar(10))"
            )
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        with connection.cursor() as cursor:
            cursor.execute('DROP TABLE "part-mst"')
        with connection.schema_editor() as editor:
            editor.delete_model(Party)

    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="S3cure-pass!")

    def setUp(self):
        Party.objects.create(party_code="S048", name="STEEL SUPPLIERS")
        Part.objects.create(part_number="ABM01", description_1="INFEED ASSY", hsn_code="8485",
                            gst_percent=Decimal("18"), sub_assembly="Y")
        Part.objects.create(part_number="ABM01_001", description_1="FLAT", drawing_no="ABM01_001",
                            hsn_code="8445", gst_percent=Decimal("18"), purchase_rate=Decimal("1500"),
                            sale_rate=Decimal("3100"), supplier_1="S048", rate_1=Decimal("1500"))
        self.client.force_login(self.user)

    def post(self, url=None, **overrides):
        data = {
            "part_number": "ABM01_099", "description_1": "  test part  ", "description_2": "",
            "drawing_no": "ABM01_099", "gauge": "", "packing": "10", "grade": "", "gross_weight": "1.5",
            "quantity_in_order": "0", "purchase_rate": "100.50", "sale_rate": "250", "opening_stock": "0",
            "reorder_qty": "5", "material": "ms", "weight_by_100": "0.5", "current_stock": "0",
            "location": "", "unit": "pcs", "hsn_code": "8445", "gst_percent": "18",
            "sub_assembly": "false", "part_type": "C",
            "supplier_1": "S048", "rate_1": "100.50", "supplier_2": "", "rate_2": "0",
            "supplier_3": "", "rate_3": "0", "supplier_4": "", "rate_4": "0",
        }
        data.update(overrides)
        return self.client.post(url or reverse("parts:api_create"), data)

    def test_grid_shows_reference_columns(self):
        response = self.client.get(reverse("parts:master"))
        for heading in ("Part No", "Item Description", "Drawing No", "HSN", "Unit", "GST%", "P.Rate", "S.Rate", "S.A"):
            self.assertContains(response, heading)
        parts = self.client.get(reverse("parts:api_list")).json()["parts"]
        self.assertEqual([p["part_number"] for p in parts], ["ABM01", "ABM01_001"])
        self.assertTrue(next(p for p in parts if p["part_number"] == "ABM01")["sub_assembly"])

    def test_add_saves_all_fields(self):
        response = self.post()
        self.assertEqual(response.status_code, 200, response.content)
        part = Part.objects.get(part_number="ABM01_099")
        self.assertEqual((part.description_1, part.unit, part.material), ("TEST PART", "PCS", "MS"))
        self.assertEqual((part.purchase_rate, part.packing, part.supplier_1), (Decimal("100.50"), 10, "S048"))
        self.assertEqual(part.part_type, "C")
        self.assertFalse(part.sub_assembly)

    def test_part_number_required_and_unique(self):
        self.assertEqual(self.post(part_number="").status_code, 400)
        self.assertEqual(self.post(part_number="ABM01").status_code, 409)

    def test_numeric_validation(self):
        response = self.post(purchase_rate="abc", gst_percent="150")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(response.json()["errors"]), {"purchase_rate", "gst_percent"})

    def test_unknown_supplier_rejected(self):
        response = self.post(supplier_1="ZZZZ")
        self.assertEqual(response.status_code, 400)
        self.assertIn("supplier_1", response.json()["errors"])

    def test_modify_protects_part_number(self):
        response = self.post(url=reverse("parts:api_update", args=["ABM01_001"]),
                             part_number="CHANGED", sale_rate="9999")
        self.assertEqual(response.status_code, 200, response.content)
        part = Part.objects.get(part_number="ABM01_001")  # unchanged key
        self.assertEqual(part.sale_rate, Decimal("9999"))
        self.assertFalse(Part.objects.filter(part_number="CHANGED").exists())

    def test_delete(self):
        self.assertEqual(self.client.post(reverse("parts:api_delete", args=["ABM01_099"])).status_code, 404)
        self.post()
        self.assertEqual(self.client.post(reverse("parts:api_delete", args=["ABM01_099"])).status_code, 200)
        self.assertFalse(Part.objects.filter(part_number="ABM01_099").exists())

    def test_delete_blocked_when_used(self):
        with mock.patch.object(services, "tables_using", return_value=["bhu", "parthis"]):
            response = self.client.post(reverse("parts:api_delete", args=["ABM01_001"]))
        self.assertEqual(response.status_code, 409)
        self.assertIn("bhu, parthis", response.json()["message"])
        self.assertTrue(Part.objects.filter(part_number="ABM01_001").exists())

    def test_duplicate_part_number_protected_from_modify_and_delete(self):
        Part.objects.create(part_number="DUP01", description_1="First")
        Part.objects.create(part_number="DUP01", description_1="Second")
        response = self.post(url=reverse("parts:api_update", args=["DUP01"]))
        self.assertEqual(response.status_code, 409)
        response = self.client.post(reverse("parts:api_delete", args=["DUP01"]))
        self.assertEqual(response.status_code, 409)
        self.assertEqual(Part.objects.filter(part_number="DUP01").count(), 2)

    def test_party_search(self):
        response = self.client.get(reverse("parts:api_parties"), {"q": "steel"})
        self.assertEqual(response.json()["parties"], [{"code": "S048", "name": "STEEL SUPPLIERS"}])

    def test_login_and_post_only(self):
        self.assertEqual(self.client.get(reverse("parts:api_delete", args=["ABM01"])).status_code, 405)
        self.client.logout()
        self.assertEqual(self.client.get(reverse("parts:master")).status_code, 302)

    def test_writes_switch(self):
        with mock.patch.object(services, "writes_enabled", return_value=False):
            self.assertEqual(self.post().status_code, 403)

    def test_excel_export(self):
        response = self.client.get(reverse("parts:export"))
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(BytesIO(response.content)).active
        self.assertEqual(ws["A1"].value, "Part Number")
        rows = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(rows, ["ABM01", "ABM01_001"])
        # sale rate column (K) for the second row should be a real number, not text
        self.assertIsInstance(ws["K3"].value, (int, float))
