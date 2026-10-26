from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.parties import services as party_services
from apps.parties.models import Party

from . import services

VALID_GSTIN = "03ABCDE1234F1Z" + party_services.gstin_check_digit("03ABCDE1234F1Z")


class CustMstTableMixin:
    """Party (vendors are cust_mst rows) maps the old software's table; tests create it."""

    @classmethod
    def setUpClass(cls):
        with connection.schema_editor() as editor:
            editor.create_model(Party)
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        with connection.schema_editor() as editor:
            editor.delete_model(Party)


@override_settings(COMPANY_STATE_CODE="03", PARTY_DEBTOR_CODE="S001", PARTY_CREDITOR_CODE="S002")
class VendorMasterTests(CustMstTableMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="S3cure-pass!")

    def setUp(self):
        self.client.force_login(self.user)

    def post_vendor(self, url=None, **overrides):
        data = {
            "name": "  A.S. Engineers  ",
            "address_1": "Industrial Area Phase 2",
            "city": "Mohali",
            "pin": "160055",
            "state_code": "03",
            "gst_no": VALID_GSTIN.lower(),
            "gst_category": "REGULAR",
            "mobile_1": "+91 98765 43210",
            "contact_person_1": "Ramesh",
            "mobile_person_1": "9988776655",
            "contact_person_2": "",
            "mobile_person_2": "",
            "email": "Accounts@Ase.in",
            "opening_balance": "1500.50",
            "transport": "ABC Transport",
            "credit_days": "30",
            "association": "",
            "grade": "a",
            "supply_type": "",
            "is_active": "true",
        }
        data.update(overrides)
        return self.client.post(url or reverse("vendors:api_create"), data)

    def test_page_renders(self):
        response = self.client.get(reverse("vendors:master"))
        self.assertContains(response, "Vendor Code")
        self.assertContains(response, "Vendor Name")

    def test_create_maps_to_cust_mst_columns_as_creditor(self):
        response = self.post_vendor()
        self.assertEqual(response.status_code, 200, response.content)
        vendor = Party.objects.get()
        self.assertEqual(vendor.party_code, "A001")  # first letter of the name + 001
        self.assertEqual(vendor.group_code, "S002")  # always Sundry Creditor
        self.assertEqual(vendor.name, "A.S. ENGINEERS")
        self.assertEqual((vendor.city, vendor.pin), ("MOHALI", "160055"))
        self.assertEqual((vendor.state_gst_code, vendor.state_name), ("03", "PUNJAB"))
        self.assertEqual((vendor.gst_no, vendor.gst_category), (VALID_GSTIN, "REGULAR"))
        self.assertEqual(vendor.mobile_1, "9876543210")
        self.assertEqual(vendor.contact_person_1, "Ramesh")
        self.assertEqual(vendor.phone_2, "9988776655")  # "Mobile Number 1" -> phone_2 slot
        self.assertEqual(vendor.email, "accounts@ase.in")
        self.assertEqual(vendor.opening_balance, Decimal("1500.50"))
        self.assertEqual(vendor.transport, "ABC Transport")
        self.assertEqual(vendor.credit_days, Decimal("30"))
        self.assertEqual(vendor.grade, "A")
        self.assertEqual(vendor.active, "Y")

    def test_vendor_not_visible_to_party_type_debtor(self):
        """A debtor (S001) must never show up as a vendor, and vice-versa."""
        Party.objects.create(party_code="D001", name="A DEBTOR ONLY", group_code="S001")
        self.post_vendor()
        codes = [v["party_code"] for v in services.list_vendors()]
        self.assertEqual(codes, ["A001"])

    def test_codes_share_the_same_namespace_as_parties(self):
        """Vendor codes come from the same A001/A002... sequence as Party Master (shared ccode)."""
        Party.objects.create(party_code="A001", name="EXISTING DEBTOR", group_code="S001")
        response = self.post_vendor()
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(Party.objects.get(group_code="S002").party_code, "A002")

    def test_validation_messages(self):
        response = self.post_vendor(name=" ", pin="1234", mobile_1="12345", email="bad", credit_days="-2")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(response.json()["errors"]), {"name", "pin", "mobile_1", "email", "credit_days"})

    def test_gstin_must_match_state(self):
        response = self.post_vendor(state_code="06")
        self.assertEqual(response.status_code, 400)
        self.assertIn("selected state's code is 06", response.json()["errors"]["gst_no"][0]["message"])

    def test_modify_keeps_code_and_updates_fields(self):
        self.post_vendor()
        response = self.client.post(reverse("vendors:api_update", args=["A001"]), {
            "name": "A.S. Engineers", "city": "Khanna", "pin": "141401", "state_code": "",
            "gst_no": "", "gst_category": "", "mobile_1": "", "opening_balance": "0",
            "transport": "", "contact_person_1": "", "mobile_person_1": "", "contact_person_2": "",
            "mobile_person_2": "", "email": "", "credit_days": "0", "association": "", "grade": "",
            "supply_type": "", "is_active": "true",
        })
        self.assertEqual(response.status_code, 200, response.content)
        vendor = Party.objects.get(party_code="A001")
        self.assertEqual((vendor.city, vendor.active, Party.objects.count()), ("KHANNA", "Y", 1))

    def test_delete_blocked_by_transactions(self):
        self.post_vendor()
        with mock.patch.object(party_services, "has_transactions", return_value=True):
            response = self.client.post(reverse("vendors:api_delete", args=["A001"]))
        self.assertEqual(response.status_code, 409)
        self.assertIn("Transactions Exist", response.json()["message"])
        self.assertTrue(Party.objects.exists())

    def test_delete_blocked_by_other_tables(self):
        self.post_vendor()
        with mock.patch.object(party_services, "tables_using", return_value=["order", "bill"]):
            response = self.client.post(reverse("vendors:api_delete", args=["A001"]))
        self.assertEqual(response.status_code, 409)
        self.assertIn("order, bill", response.json()["message"])

    def test_delete(self):
        self.post_vendor()
        self.assertEqual(self.client.post(reverse("vendors:api_delete", args=["A001"])).status_code, 200)
        self.assertFalse(Party.objects.exists())

    def test_cannot_delete_or_modify_a_debtor_through_vendor_urls(self):
        Party.objects.create(party_code="D001", name="A DEBTOR ONLY", group_code="S001")
        self.assertEqual(self.client.post(reverse("vendors:api_delete", args=["D001"])).status_code, 404)
        response = self.client.post(reverse("vendors:api_update", args=["D001"]), {"name": "X"})
        self.assertEqual(response.status_code, 404)

    def test_next_code_preview(self):
        self.post_vendor()
        response = self.client.get(reverse("vendors:api_next_code"), {"name": "Alfa Tools"})
        self.assertEqual(response.json()["code"], "A002")

    def test_writes_can_be_switched_off(self):
        with mock.patch.object(services, "writes_enabled", return_value=False):
            self.assertEqual(self.post_vendor().status_code, 403)
        self.assertFalse(Party.objects.exists())

    def test_excel_export(self):
        self.post_vendor()
        response = self.client.get(reverse("vendors:export"))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        headers = [c.value for c in ws[4]]
        self.assertEqual(headers[:3], ["Vendor Code", "Vendor Name", "Address 1"])
        self.assertEqual(ws.cell(row=5, column=1).value, "A001")
        self.assertEqual(ws.cell(row=5, column=2).value, "A.S. ENGINEERS")


class VendorGstinLookupTests(TestCase):
    """GST No. auto-fill on the Vendor form reuses the Customer Master lookup;
    the provider call is mocked so these don't reach the network."""

    def test_view_requires_login(self):
        res = self.client.get(reverse("vendors:api_gstin_lookup"), {"gstin": VALID_GSTIN})
        self.assertEqual(res.status_code, 302)

    def test_returns_only_name_address_and_pin(self):
        self.client.force_login(get_user_model().objects.create_user(username="vendor_gst", password="x"))
        with mock.patch("apps.eway_bill.provider.lookup_gstin") as mocked:
            mocked.return_value = {
                "Gstin": VALID_GSTIN, "TradeName": "Acme", "AddrBno": "12", "AddrBnm": "Industrial Estate",
                "AddrSt": "Main Road", "AddrLoc": "Mohali", "AddrPncd": "160055", "StateCode": 3,
            }
            res = self.client.get(reverse("vendors:api_gstin_lookup"), {"gstin": VALID_GSTIN})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["fields"], {
            "name": "ACME", "address_1": "12, Industrial Estate", "address_2": "Main Road, Mohali", "pin": "160055",
        })

    def test_invalid_gstin_is_an_error_without_provider_call(self):
        self.client.force_login(get_user_model().objects.create_user(username="vendor_gst2", password="x"))
        with mock.patch("apps.eway_bill.provider.lookup_gstin") as mocked:
            res = self.client.get(reverse("vendors:api_gstin_lookup"), {"gstin": "03ABC"})
            mocked.assert_not_called()
        self.assertEqual(res.status_code, 400)
        self.assertIn("message", res.json())
