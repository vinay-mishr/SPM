"""Vendor Master Excel export (openpyxl). Reads the real cust_mst rows (group S002)."""
from io import BytesIO

from django.conf import settings

from apps.parties.models import Party

from . import services

COLUMNS = [
    ("Vendor Code", "party_code", 8, "@"),
    ("Vendor Name", "name", 32, "@"),
    ("Address 1", "address_1", 26, "@"),
    ("Address 2", "address_2", 26, "@"),
    ("Address 3", "address_3", 26, "@"),
    ("City", "city", 18, "@"),
    ("PIN", "pin", 8, "@"),
    ("GST State Name", "state_name", 18, "@"),
    ("GST Number", "gst_no", 18, "@"),
    ("GST Category", "gst_category", 14, "@"),
    ("Phone Number", "phone_no", 14, "@"),
    ("Mobile", "mobile_1", 14, "@"),
    ("Contact Person 1", "contact_person_1", 18, "@"),
    ("Mobile Number 1", "mobile_person_1", 14, "@"),
    ("Contact Person 2", "contact_person_2", 18, "@"),
    ("Mobile Number 2", "mobile_person_2", 14, "@"),
    ("Email Address", "email", 26, "@"),
    ("Opening Balance", "opening_balance", 14, "0.00"),
    ("Transport", "transport", 18, "@"),
    ("Credit Days", "credit_days", 10, "0"),
    ("Association", "association", 12, "@"),
    ("Grade", "grade", 8, "@"),
    ("Interstate / Within State", "supply_type_label", 20, "@"),
    ("Active", "active_label", 8, "@"),
]
SUPPLY_TYPE_LABELS = {"W": "Within State", "I": "Interstate"}


def fetch_vendors():
    """Vendors (Sundry Creditors), sorted by name -- the real database, no dummy rows."""
    return services.list_vendors()


def build_workbook(vendors, company_name):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial Narrow"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=12, bold=True)
    ws.cell(row=2, column=1, value="VENDOR MASTER LIST").font = Font(name=font_name, size=12, bold=True)

    header_row = 4
    for col, (heading, _field, width, _fmt) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, bold=True)
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for vendor in vendors:
        data = dict(vendor)
        data["supply_type_label"] = SUPPLY_TYPE_LABELS.get(vendor.get("supply_type"), "")
        data["active_label"] = "Yes" if vendor.get("is_active") else "No"
        for col, (_heading, field, _width, fmt) in enumerate(COLUMNS, start=1):
            value = sanitize(data.get(field))
            if fmt != "@" and value not in ("", None):
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    pass
            cell = ws.cell(row=row, column=col, value=value if value != "" else None)
            cell.font = Font(name=font_name, size=10)
            cell.number_format = fmt
        row += 1

    for r in range(header_row, row):
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
