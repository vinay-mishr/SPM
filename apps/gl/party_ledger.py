"""Accounts > Ledger Printing > Party Ledger (VFP "PARTY LEDGER" report).

Identical shape to apps.gl.main_ledger's own "MAIN LEDGER" -- same per-
account opening/closing balance formula, same running-balance sign
convention (Debit increases, Credit decreases), same "only shown if it has
a non-zero opening balance or at least one transaction" ``fl`` flag, same
VNO/V.DATE/NARRATION/DEBIT/CREDIT/BALANCE columns -- just keyed by
``cust_mst`` (party/sub-ledger master) and ``trans.tccode`` instead of
``gled_mst``/``trans.tcode``. ALL mode here has no exclusion filter at all
(every party in ``cust_mst``, unlike Main Ledger's ALL mode which drops the
Cash/Sundry-Debtor/Sundry-Creditor codes) -- confirmed against the pasted
VFP source, which filters ALL mode's ``cust_mst`` scan by nothing but
``set orde to cname``.

See apps.gl.main_ledger's own docstring for the parts this doesn't repeat
(the opening-balance formula shape, the sign convention, the per-page
layout). Deterministic order within a party: ``tdate`` then ``tdoc``, same
tie-breaker the rest of this app's reports use; parties themselves ordered
by ``cname`` (the VFP source's own ``set orde to cname``).

No PRINT button -- the VFP form's own screenshot only has VIEW / EXCEL /
EXIT (not ported here as "missing", it's genuinely absent from the source).
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows, search_parties


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


def _parties(mode, ccode):
    if mode == "S":
        if not ccode:
            return []
        return _rows("SELECT ccode, cname, cyob FROM cust_mst WHERE ccode = %s", [ccode])
    return _rows("SELECT ccode, cname, cyob FROM cust_mst ORDER BY cname")


def party_ledger_blocks(mode, ccode, start, end):
    blocks = []
    for party in _parties(mode, ccode):
        code = _clean(party["ccode"])
        balance = _d(party["cyob"])
        for r in _rows(
            "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tccode = %s AND tdate < %s GROUP BY tdbcr",
            [code, start],
        ):
            if _clean(r["tdbcr"]) == "D":
                balance += _d(r["total"])
            elif _clean(r["tdbcr"]) == "C":
                balance -= _d(r["total"])
        opening_balance = balance
        has_opening = opening_balance != 0

        lines = []
        total_debit = total_credit = Decimal("0")
        for r in _rows(
            "SELECT tdoc, tdate, tdbcr, tamt, tnar1, tnar2 FROM trans "
            "WHERE tccode = %s AND tdate >= %s AND tdate <= %s ORDER BY tdate, tdoc",
            [code, start, end],
        ):
            amt = _d(r["tamt"])
            tdbcr = _clean(r["tdbcr"])
            if tdbcr == "D":
                balance += amt
                total_debit += amt
                debit, credit = f"{amt:.2f}", ""
            else:
                balance -= amt
                total_credit += amt
                debit, credit = "", f"{amt:.2f}"
            narration1 = _clean(r["tnar1"]) or ""
            narration2 = _clean(r["tnar2"]) or ""
            lines.append({
                "vno": _clean(r["tdoc"]) or "", "date": r["tdate"].isoformat(),
                "narration1": narration1, "narration2": narration2 if narration2 and narration2 != narration1 else "",
                "debit": debit, "credit": credit,
                "balance": f"{abs(balance):.2f}", "side": "Cr" if balance < 0 else "Dr",
            })

        if not has_opening and not lines:
            continue

        blocks.append({
            "code": code, "name": _clean(party["cname"]) or "",
            "has_opening": has_opening,
            "opening_balance": f"{abs(opening_balance):.2f}", "opening_side": "Cr" if opening_balance < 0 else "Dr",
            "lines": lines,
            "total_debit": f"{total_debit:.2f}", "total_credit": f"{total_credit:.2f}",
            "closing_balance": f"{abs(balance):.2f}", "closing_side": "Cr" if balance < 0 else "Dr",
        })
    return blocks


class PartyLedgerView(LoginRequiredMixin, TemplateView):
    template_name = "gl/party_ledger.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"mode": "S", "start_date": _fy_start(today).isoformat(), "end_date": today.isoformat()}
        return context


class PartyLedgerGraphicalView(PartyLedgerView):
    """Accounts > Ledger Printing > Party Ledger Graphical -- the VFP
    source's own ``repo form pledger prev`` (on-screen report preview) entry
    point for this same report: identical data (``gl:party_ledger_data``/
    ``_parties``/``_export``, all unchanged) and identical ALL/SINGLE +
    opening/closing balance logic as "PARTY LEDGER" above -- the only UI
    difference is a PRINT button (``repo form pledger prev`` here vs. no
    print option on the plain Party Ledger form), matching the VFP form's
    own screenshot (VIEW / PRINT / EXCEL / EXIT).
    """
    template_name = "gl/party_ledger_graphical.html"


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    mode = (request.GET.get("mode") or "S").strip().upper()
    mode = mode if mode in ("A", "S") else "S"
    ccode = (request.GET.get("ccode") or "").strip()
    start = parse(request.GET.get("start_date"), _fy_start(today))
    end = parse(request.GET.get("end_date"), today)
    return mode, ccode, start, end


@login_required
@require_GET
def pl_parties(request):
    return JsonResponse({"parties": search_parties(request.GET.get("q", ""))})


@login_required
@require_GET
def pl_data(request):
    mode, ccode, start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    if mode == "S" and not ccode:
        return JsonResponse({"message": "Select a party."}, status=400)
    return JsonResponse({
        "start_date": start.isoformat(), "end_date": end.isoformat(),
        "blocks": party_ledger_blocks(mode, ccode, start, end),
    })


@login_required
@require_GET
def pl_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    mode, ccode, start, end = _params(request)
    if start > end or (mode == "S" and not ccode):
        return JsonResponse({"message": "Enter a valid Start/End Date and Party."}, status=400)

    blocks = party_ledger_blocks(mode, ccode, start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([])
    ws.append([f"PARTY LEDGER FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["Party Name", "Vou.Date", "Narration", "Debit", "Credit", "Balance", ""])
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([20, 11, 30, 11, 11, 11, 4], "ABCDEFG"):
        ws.column_dimensions[col].width = w

    for b in blocks:
        ws.append([b["name"]])
        ws[ws.max_row][0].font = Font(bold=True)
        if b["has_opening"]:
            row = ["", "", "Opening Balance", float(b["opening_balance"]) if b["opening_side"] == "Dr" else None,
                   float(b["opening_balance"]) if b["opening_side"] == "Cr" else None]
            ws.append(row)
            for cell in ws[ws.max_row][2:5]:
                cell.font = Font(bold=True)
        for l in b["lines"]:
            narration = l["narration1"] + (" " + l["narration2"] if l["narration2"] else "")
            ws.append(["", l["date"], narration, float(l["debit"]) if l["debit"] else None,
                       float(l["credit"]) if l["credit"] else None, float(l["balance"]), l["side"]])
        if b["has_opening"] or b["lines"]:
            ws.append(["", "", "Total Amount", float(b["total_debit"]), float(b["total_credit"])])
            row = ["", "", "Closing Balance", float(b["closing_balance"]) if b["closing_side"] == "Dr" else None,
                   float(b["closing_balance"]) if b["closing_side"] == "Cr" else None]
            ws.append(row)
            for cell in ws[ws.max_row][2:5]:
                cell.font = Font(bold=True)
        ws.append([])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Party_Ledger_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
