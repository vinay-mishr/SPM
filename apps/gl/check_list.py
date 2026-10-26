"""Accounts > Data Entry > Check-List Printing (VFP "CHECK-LIST" report).

A data-integrity audit, not an accounting entry screen: every ``trans`` line
for one voucher type (CV/BV/JV -- the same ``ttype`` values
apps.gl.cash_entry/bank_entry/journal_entry post) in a date range, flagged
per line when its own ``tcode`` isn't a real ``gled_mst`` row ("GL Code Not
Found") or its ``tccode`` is set but isn't a real ``cust_mst`` row ("Party
Code Not Found"), and per voucher (grouped by date+voucher no) when its own
lines' debit and credit totals don't match ("Amount mismatch") -- a direct
port of the VFP PRN report's own three checks.

No write path: this only reads ``trans``/``gled_mst``/``cust_mst``, same
confirmed columns apps.gl.cash_entry already established.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows

VOUCHER_TYPES = {"CV": "Cash Entry", "BV": "Bank Entry", "JV": "Journal Entry"}


def check_list_rows(ttype, start, end):
    sql = (
        "SELECT tdate, tdoc, tnar1, tcode, tccode, tdbcr, tamt FROM trans "
        "WHERE ttype = %s AND tdate >= %s AND tdate <= %s ORDER BY tdate, tdoc"
    )
    lines = _rows(sql, [ttype, start, end])
    if not lines:
        return []

    codes = {_clean(l["tcode"]) for l in lines if _clean(l["tcode"])}
    ccodes = {_clean(l["tccode"]) for l in lines if _clean(l["tccode"])}
    gl_found = set()
    if codes:
        gl_found = {_clean(r["gcode"]) for r in _rows(
            "SELECT gcode FROM gled_mst WHERE gcode IN (%s)" % ",".join(["%s"] * len(codes)), list(codes),
        )}
    party_found = set()
    if ccodes:
        party_found = {_clean(r["ccode"]) for r in _rows(
            "SELECT ccode FROM cust_mst WHERE ccode IN (%s)" % ",".join(["%s"] * len(ccodes)), list(ccodes),
        )}

    # Group consecutive lines into vouchers by (tdate, tdoc) -- same as the
    # VFP source's own "do while tdate=mtdate and tdoc=mtdoc" inner loop.
    groups = []
    current_key = None
    current = None
    for l in lines:
        key = (l["tdate"], _clean(l["tdoc"]))
        if key != current_key:
            current = []
            groups.append((key, current))
            current_key = key
        current.append(l)

    out = []
    for (tdate, tdoc), group in groups:
        total_c = total_d = Decimal("0")
        group_out = []
        for l in group:
            tcode, tccode = _clean(l["tcode"]), _clean(l["tccode"])
            tdbcr = _clean(l["tdbcr"])
            amt = _d(l["tamt"])
            if tdbcr == "C":
                total_c += amt
            elif tdbcr == "D":
                total_d += amt
            errors = []
            if tcode and tcode not in gl_found:
                errors.append("GL Code Not Found")
            if tccode and tccode not in party_found:
                errors.append("Party Code Not Found")
            group_out.append({
                "tdate": tdate.isoformat() if tdate else "", "tdoc": tdoc,
                "narration": _clean(l["tnar1"]) or "", "code": tcode, "scode": tccode,
                "dc": tdbcr, "amount": f"{amt:.2f}", "error": ", ".join(errors),
            })
        if total_c != total_d and group_out:
            last = group_out[-1]
            last["error"] = (last["error"] + ", " if last["error"] else "") + "Amount mismatch"
        out.extend(group_out)
    return out


class CheckListView(LoginRequiredMixin, TemplateView):
    template_name = "gl/check_list.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        context["voucher_types"] = VOUCHER_TYPES
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    ttype = (request.GET.get("ttype") or "CV").strip().upper()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    return ttype, start, end


@login_required
@require_GET
def cl_data(request):
    ttype, start, end = _params(request)
    if ttype not in VOUCHER_TYPES:
        return JsonResponse({"message": "Enter a valid Voucher Type (CV, BV or JV)."}, status=400)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse({"rows": check_list_rows(ttype, start, end)})


@login_required
@require_GET
def cl_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ttype, start, end = _params(request)
    rows = check_list_rows(ttype, start, end) if ttype in VOUCHER_TYPES else []

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"CHECK-LIST FOR THE PERIOD OF {start:%d/%m/%Y} TO {end:%d/%m/%Y} ({VOUCHER_TYPES.get(ttype, ttype)})"])
    ws.append([])
    header = ["VOU. DATE", "V.NO.", "NARRATION", "CODE", "SCODE", "D/C", "AMOUNT", "ERROR MESSAGE"]
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([12, 8, 30, 8, 8, 6, 13, 30], "ABCDEFGH"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([r["tdate"], r["tdoc"], r["narration"], r["code"], r["scode"], r["dc"],
                   float(r["amount"]), r["error"]])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Check_List_{ttype}_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
