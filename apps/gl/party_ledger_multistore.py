"""Accounts > Ledger Printing > Party Ledger Multistore (VFP "PARTY LEDGER"
bill-wise report, ``pled.PRN``).

A different data model from every other Ledger Printing report in this app:
not the ``trans`` ledger-posting table at all, but a bill-wise outstanding/
collection view -- one row per sales bill (``bill``) for the one selected
party, each with its amount received and last receipt date pulled from the
collection table (``col``). Both tables and every column used here are
already confirmed elsewhere in this project, not guessed from the pasted
VFP source alone:

  * ``bill`` (p_code, div, bill_no, b_date, n_amount) -- apps.money_receipt.
    overdue's own confirmed ``bill`` query.
  * ``col`` (div, bill_no, b_date, amount, slip_dt) -- apps.money_receipt.
    services/dishonour's own confirmed ``col`` queries (the same table
    apps.money_receipt.services.save_receipt writes to).

For each bill in the date range: amount received = sum of ``col.amount``
for rows matching (div, bill_no, b_date) against that bill; receipt date =
that bill's latest ``col.slip_dt``. The VFP source's own inner ``scan...
endscan`` over ``col`` just keeps whatever ``slip_dt`` the loop last visits
(no explicit order), which isn't reproducible in SQL -- ``MAX(slip_dt)``
(the most recent receipt) is used here instead; flagged as a deliberate,
documented choice rather than a guess at the VFP's own undefined scan
order. Balance = bill amount - amount received. Delay days = receipt date -
bill date, left blank when nothing has been received yet.

No opening/closing balance here at all (unlike Main Ledger/Party Ledger) --
every bill in the date range is listed and totalled, full stop.

VIEW / EXCEL / EXIT only, matching the VFP form's own screenshot (no
PRINT).
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows, search_parties


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


def party_ledger_multistore_rows(ccode, start, end):
    bills = _rows(
        "SELECT div, bill_no, b_date, n_amount FROM bill "
        "WHERE p_code = %s AND b_date >= %s AND b_date <= %s ORDER BY b_date, bill_no",
        [ccode, start, end],
    )
    rows = []
    total_bill = total_received = total_balance = Decimal("0")
    for b in bills:
        div = _clean(b["div"]) or ""
        bill_no = _clean(b["bill_no"]) or ""
        b_date = b["b_date"]
        bill_amount = _d(b["n_amount"])

        col = _rows(
            "SELECT SUM(amount) AS total, MAX(slip_dt) AS last_dt FROM col "
            "WHERE div = %s AND bill_no = %s AND b_date = %s",
            [div, bill_no, b_date],
        )
        amount_received = _d(col[0]["total"]) if col and col[0]["total"] is not None else Decimal("0")
        receipt_date = col[0]["last_dt"] if col else None

        balance = bill_amount - amount_received
        delay_days = (receipt_date - b_date).days if receipt_date else None

        total_bill += bill_amount
        total_received += amount_received
        total_balance += balance

        rows.append({
            "div": div, "bill_no": bill_no, "bill_date": b_date.isoformat(),
            "bill_amount": f"{bill_amount:.2f}", "amount_received": f"{amount_received:.2f}",
            "receipt_date": receipt_date.isoformat() if receipt_date else "",
            "balance": f"{balance:.2f}", "delay_days": delay_days if delay_days is not None else "",
        })
    return rows, f"{total_bill:.2f}", f"{total_received:.2f}", f"{total_balance:.2f}"


class PartyLedgerMultistoreView(LoginRequiredMixin, TemplateView):
    template_name = "gl/party_ledger_multistore.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": _fy_start(today).isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    ccode = (request.GET.get("ccode") or "").strip()
    start = parse(request.GET.get("start_date"), _fy_start(today))
    end = parse(request.GET.get("end_date"), today)
    return ccode, start, end


@login_required
@require_GET
def plm_parties(request):
    return JsonResponse({"parties": search_parties(request.GET.get("q", ""))})


@login_required
@require_GET
def plm_data(request):
    ccode, start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    if not ccode:
        return JsonResponse({"message": "Select a party."}, status=400)
    rows, total_bill, total_received, total_balance = party_ledger_multistore_rows(ccode, start, end)
    return JsonResponse({
        "start_date": start.isoformat(), "end_date": end.isoformat(),
        "rows": rows, "total_bill": total_bill, "total_received": total_received, "total_balance": total_balance,
    })


@login_required
@require_GET
def plm_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    ccode, start, end = _params(request)
    if start > end or not ccode:
        return JsonResponse({"message": "Enter a valid Start/End Date and Party."}, status=400)

    party_name = ""
    party_city = ""
    party = _rows("SELECT cname, cadd3 FROM cust_mst WHERE ccode = %s", [ccode])
    if party:
        party_name = _clean(party[0]["cname"]) or ""
        party_city = _clean(party[0]["cadd3"]) or ""

    rows, total_bill, total_received, total_balance = party_ledger_multistore_rows(ccode, start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"PARTY LEDGER FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([f"PARTY NAME : {party_name}"])
    ws.append([f"PARTY CITY : {party_city}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["BILL NO", "BILL DATE", "BILL AMOUNT", "AMT RECEIVED", "RECEIPT DATE", "BALANCE AMT", "DELAY DAYS"])
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([14, 12, 13, 13, 13, 13, 11], "ABCDEFG"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([f"{r['div']},{r['bill_no']}", r["bill_date"], float(r["bill_amount"]),
                   float(r["amount_received"]), r["receipt_date"], float(r["balance"]), r["delay_days"]])

    total_row = ["", "TOTAL", float(total_bill), float(total_received), "", float(total_balance), ""]
    ws.append(total_row)
    ws[ws.max_row][1].font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Party_Ledger_Multistore_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
