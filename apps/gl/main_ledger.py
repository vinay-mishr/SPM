"""Accounts > Ledger Printing > Main Ledger (VFP "MAIN LEDGER" report).

One block per General Ledger account (``gled_mst``), either every account
(ALL/SINGLE toggle, the VFP source's own ``tsa`` field) excluding the
configured Cash ledger and the Sundry Debtor/Creditor GROUP codes
themselves (``psdcode``/``psccode`` -- apps.gl.cash_entry's own confirmed
``settings.PARTY_DEBTOR_CODE``/``PARTY_CREDITOR_CODE``), or one specific
account picked via the DESCRIPTION lookup.

Unlike every other Book Printing report in this app, Main Ledger has NO
``ttype`` filter at all -- every ``trans`` row whose ``tcode`` matches the
ledger's own code counts, across CV/BV/JV alike (this is the ledger's own
side of its postings, not a voucher's contra side).

Opening balance: ``gled_mst.gyob`` + sum(Debit tamt) - sum(Credit tamt) for
that ledger's own ``tcode``, dated before Start Date (same formula shape as
every other Book Printing report's opening balance, just scoped per-ledger
here instead of to the Cash ledger specifically).

Debit/credit sign convention for the ledger's own running balance (this is
the ledger's own side, same direction as the Cash-ledger opening-balance
calc those other reports share -- NOT the "contra" sign Cash Book/Combined
Book use for their movement lines): Debit increases the balance, Credit
decreases it.

A ledger is only included in the output when it has a non-zero opening
balance OR at least one transaction in the date range -- exactly matching
the VFP source's own ``fl`` flag, which never prints a ledger's header at
all otherwise.

Deterministic order within a ledger: ``tdate`` then ``tdoc`` (voucher
number), the same tie-breaker every other Book Printing report here uses;
ledgers themselves are ordered by ``gdesc`` (the VFP source's own
``set orde to gdesc``).
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import CASH_LEDGER_CODE, _clean, _d, _rows, search_ledgers

PARTY_DEBTOR_CODE = getattr(settings, "PARTY_DEBTOR_CODE", "S001")
PARTY_CREDITOR_CODE = getattr(settings, "PARTY_CREDITOR_CODE", "S002")


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


def _ledgers(mode, gcode):
    if mode == "S":
        if not gcode:
            return []
        return _rows("SELECT gcode, gdesc, gyob FROM gled_mst WHERE gcode = %s", [gcode])
    return _rows(
        "SELECT gcode, gdesc, gyob FROM gled_mst WHERE gcode <> %s AND gcode <> %s AND gcode <> %s ORDER BY gdesc",
        [CASH_LEDGER_CODE, PARTY_DEBTOR_CODE, PARTY_CREDITOR_CODE],
    )


def _build_ledger_block(code, desc, opening_balance, line_rows):
    """The per-ledger math itself (opening carried into a running balance,
    one line per ``trans`` row, closing totals) -- pulled out of
    ``ledger_block`` so both that single-ledger function AND
    ``main_ledger_blocks``'s own batched path (see its docstring) build the
    exact same block shape from whichever opening balance / line rows they
    were handed, without duplicating the formula."""
    balance = opening_balance
    has_opening = opening_balance != 0

    lines = []
    total_debit = total_credit = Decimal("0")
    for r in line_rows:
        amt = _d(r["tamt"])
        tdbcr = _clean(r["tdbcr"])
        if tdbcr == "D":
            balance += amt
            total_debit += amt
            debit, credit = f"{amt:.2f}", ""
        else:
            balance -= amt
            total_credit += amt
            debit, credit = "", f"{amt:.2f}"
        narration1 = _clean(r["tnar1"]) or ""
        narration2 = _clean(r["tnar2"]) or ""
        lines.append({
            "vno": _clean(r["tdoc"]) or "", "date": r["tdate"].isoformat(),
            "narration1": narration1, "narration2": narration2 if narration2 and narration2 != narration1 else "",
            "debit": debit, "credit": credit,
            "balance": f"{abs(balance):.2f}", "side": "Cr" if balance < 0 else "Dr",
        })

    if not has_opening and not lines:
        return None

    return {
        "code": code, "desc": desc,
        "has_opening": has_opening,
        "opening_balance": f"{abs(opening_balance):.2f}", "opening_side": "Cr" if opening_balance < 0 else "Dr",
        "lines": lines,
        "total_debit": f"{total_debit:.2f}", "total_credit": f"{total_credit:.2f}",
        "closing_balance": f"{abs(balance):.2f}", "closing_side": "Cr" if balance < 0 else "Dr",
    }


def ledger_block(code, desc, start, end):
    """One ledger's opening/lines/closing block, or None if it has neither a
    non-zero opening balance nor any transaction in range (the VFP source's
    own ``fl`` flag). Shared by every "one block per ``gled_mst`` account"
    report in this app (Main Ledger Group Wise's gl_type-filtered scan --
    Main Ledger's own ALL/SINGLE scan uses the batched path in
    ``main_ledger_blocks`` instead, for a single account's own two queries
    this is no slower, just the single-ledger entry point other callers
    already depend on).
    """
    balance = _d(_rows("SELECT gyob FROM gled_mst WHERE gcode = %s", [code])[0]["gyob"])
    for r in _rows(
        "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tcode = %s AND tdate < %s GROUP BY tdbcr",
        [code, start],
    ):
        if _clean(r["tdbcr"]) == "D":
            balance += _d(r["total"])
        elif _clean(r["tdbcr"]) == "C":
            balance -= _d(r["total"])

    line_rows = _rows(
        "SELECT tdoc, tdate, tdbcr, tamt, tnar1, tnar2 FROM trans "
        "WHERE tcode = %s AND tdate >= %s AND tdate <= %s ORDER BY tdate, tdoc, rec_no",
        [code, start, end],
    )
    return _build_ledger_block(code, desc, balance, line_rows)


def main_ledger_blocks(mode, gcode, start, end):
    """Batched version of "one ledger_block() call per account": the ALL
    scan was previously 3 sequential round-trips to the legacy SQL Server
    PER LEDGER (gled_mst.gyob lookup + opening-balance aggregate + line
    rows) -- for the whole chart of accounts that's 300+ round-trips for
    one page load, the actual cause of "Main Ledger is slow". Opening
    balances and transaction lines for every ledger in scope are now
    fetched in exactly two queries total (one aggregate GROUP BY tcode,
    tdbcr for everything before Start Date; one ordered row set for
    everything in range), grouped in Python by ``tcode``, and fed through
    the same ``_build_ledger_block`` formula ``ledger_block`` uses -- same
    numbers, same output shape, just not one DB round-trip per account.
    """
    ledgers = _ledgers(mode, gcode)
    if not ledgers:
        return []

    codes = [_clean(led["gcode"]) for led in ledgers]
    opening_balances = {code: _d(led["gyob"]) for code, led in zip(codes, ledgers)}
    placeholders = ",".join(["%s"] * len(codes))

    for r in _rows(
        f"SELECT tcode, tdbcr, SUM(tamt) AS total FROM trans "
        f"WHERE tcode IN ({placeholders}) AND tdate < %s GROUP BY tcode, tdbcr",
        codes + [start],
    ):
        code = _clean(r["tcode"])
        if code not in opening_balances:
            continue
        if _clean(r["tdbcr"]) == "D":
            opening_balances[code] += _d(r["total"])
        elif _clean(r["tdbcr"]) == "C":
            opening_balances[code] -= _d(r["total"])

    lines_by_code = {}
    for r in _rows(
        f"SELECT tcode, tdoc, tdate, tdbcr, tamt, tnar1, tnar2 FROM trans "
        f"WHERE tcode IN ({placeholders}) AND tdate >= %s AND tdate <= %s ORDER BY tcode, tdate, tdoc, rec_no",
        codes + [start, end],
    ):
        lines_by_code.setdefault(_clean(r["tcode"]), []).append(r)

    blocks = []
    for code, led in zip(codes, ledgers):
        block = _build_ledger_block(code, _clean(led["gdesc"]) or "", opening_balances[code], lines_by_code.get(code, []))
        if block:
            blocks.append(block)
    return blocks


class MainLedgerView(LoginRequiredMixin, TemplateView):
    template_name = "gl/main_ledger.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"mode": "S", "start_date": _fy_start(today).isoformat(), "end_date": today.isoformat()}
        return context


class MainLedgerGraphicalView(MainLedgerView):
    """Accounts > Ledger Printing > Main Ledger Graphical -- the VFP source's
    own ``repo form ledger prev`` (on-screen report preview) entry point for
    this same report: identical data (``gl:main_ledger_data``/``_ledgers``,
    unchanged) and identical ALL/SINGLE + opening/closing balance logic as
    "MAIN LEDGER" above, just without the EXCEL/PRINT buttons -- the VFP
    form's own screenshot only has VIEW and EXIT.
    """
    template_name = "gl/main_ledger_graphical.html"


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    mode = (request.GET.get("mode") or "S").strip().upper()
    mode = mode if mode in ("A", "S") else "S"
    gcode = (request.GET.get("gcode") or "").strip()
    start = parse(request.GET.get("start_date"), _fy_start(today))
    end = parse(request.GET.get("end_date"), today)
    return mode, gcode, start, end


@login_required
@require_GET
def ml_ledgers(request):
    return JsonResponse({"ledgers": search_ledgers(request.GET.get("q", ""))})


@login_required
@require_GET
def ml_data(request):
    mode, gcode, start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    if mode == "S" and not gcode:
        return JsonResponse({"message": "Select an account."}, status=400)
    return JsonResponse({
        "start_date": start.isoformat(), "end_date": end.isoformat(),
        "blocks": main_ledger_blocks(mode, gcode, start, end),
    })


@login_required
@require_GET
def ml_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    mode, gcode, start, end = _params(request)
    if start > end or (mode == "S" and not gcode):
        return JsonResponse({"message": "Enter a valid Start/End Date and Account."}, status=400)

    blocks = main_ledger_blocks(mode, gcode, start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([])
    ws.append([f"MAIN LEDGER FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["V.NO.", "Vou.Date", "Narration", "Debit", "Credit", "Balance", ""])
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([7, 11, 50, 11, 11, 11, 4], "ABCDEFG"):
        ws.column_dimensions[col].width = w

    for b in blocks:
        ws.append([b["desc"]])
        ws[ws.max_row][0].font = Font(bold=True)
        if b["has_opening"]:
            row = ["", "", "Opening Balance", float(b["opening_balance"]) if b["opening_side"] == "Dr" else None,
                   float(b["opening_balance"]) if b["opening_side"] == "Cr" else None]
            ws.append(row)
            for cell in ws[ws.max_row][2:5]:
                cell.font = Font(bold=True)
        for l in b["lines"]:
            narration = l["narration1"] + (" " + l["narration2"] if l["narration2"] else "")
            ws.append([l["vno"], l["date"], narration, float(l["debit"]) if l["debit"] else None,
                       float(l["credit"]) if l["credit"] else None, float(l["balance"]), l["side"]])
        if b["has_opening"] or b["lines"]:
            ws.append(["", "", "Total Amount", float(b["total_debit"]), float(b["total_credit"])])
            row = ["", "", "Closing Balance", float(b["closing_balance"]) if b["closing_side"] == "Dr" else None,
                   float(b["closing_balance"]) if b["closing_side"] == "Cr" else None]
            ws.append(row)
            for cell in ws[ws.max_row][2:5]:
                cell.font = Font(bold=True)
        ws.append([])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Main_Ledger_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
