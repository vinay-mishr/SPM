"""Accounts > Data Entry > Voucher Printing (VFP "VOUCHER PRINTING" report).

Renders one already-saved voucher (CV/BV/JV -- the same ``trans`` rows
apps.gl.cash_entry/bank_entry/journal_entry post, keyed the same way: ttype
+ tdate + tdoc) as a traditional two-column accounting voucher: each line's
account (``gled_mst`` if it has no ``tccode``, else ``cust_mst``), prefixed
"TO " for a Debit line / "BY " for a Credit line (the VFP source's own
``tnar1``/``tnar2`` narration convention), Debit/Credit totals, and a
signature line.

The VFP source's own voucher-type label is literally
``iif(ttype='CV','CASH','JOURNAL')`` -- i.e. anything that isn't CV prints
as "JOURNAL VOUCHER", Bank Entry (BV) included. Kept exactly as-is (not
"fixed" to say "BANK VOUCHER" for BV), flagged here rather than guessed.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _date, _rows

VOUCHER_TYPES = {"CV": "Cash Entry", "BV": "Bank Entry", "JV": "Journal Entry"}


def voucher_label(ttype):
    return "CASH VOUCHER" if ttype == "CV" else "JOURNAL VOUCHER"


def get_voucher(ttype, voucher_date, voucher_no):
    rows = _rows(
        "SELECT t.*, g.gdesc, c.cname FROM trans t LEFT JOIN gled_mst g ON g.gcode = t.tcode "
        "LEFT JOIN cust_mst c ON c.ccode = t.tccode "
        "WHERE t.ttype = %s AND t.tdate = %s AND t.tdoc = %s ORDER BY t.tdbcr DESC",
        [ttype, voucher_date, voucher_no],
    )
    if not rows:
        return None
    lines = []
    total_debit = total_credit = Decimal("0")
    for r in rows:
        tccode = _clean(r["tccode"])
        account = _clean(r["cname"]) if tccode else _clean(r["gdesc"])
        account = account or ("PARTY NOT FOUND" if tccode else "ACCOUNT NOT FOUND")
        tdbcr = _clean(r["tdbcr"])
        amt = _d(r["tamt"])
        if tdbcr == "D":
            total_debit += amt
        else:
            total_credit += amt
        lines.append({
            "code": _clean(r["tcode"]) or "", "scode": tccode or "", "account": account,
            "dc": tdbcr, "narration1": _clean(r["tnar1"]) or "", "narration2": _clean(r.get("tnar2")) or "",
            "debit": f"{amt:.2f}" if tdbcr == "D" else "", "credit": f"{amt:.2f}" if tdbcr == "C" else "",
        })
    return {
        "ttype": ttype, "label": voucher_label(ttype), "voucher_date": voucher_date.isoformat(),
        "voucher_no": voucher_no, "lines": lines,
        "total_debit": f"{total_debit:.2f}", "total_credit": f"{total_credit:.2f}",
    }


class VoucherPrintingView(LoginRequiredMixin, TemplateView):
    template_name = "gl/voucher_printing.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"voucher_date": date.today().isoformat()}
        context["voucher_types"] = VOUCHER_TYPES
        return context


def _params(request):
    ttype = (request.GET.get("ttype") or "CV").strip().upper()
    voucher_date = _date(request.GET.get("voucher_date")) or date.today()
    voucher_no = (request.GET.get("voucher_no") or "").strip()
    return ttype, voucher_date, voucher_no


@login_required
@require_GET
def vp_data(request):
    ttype, voucher_date, voucher_no = _params(request)
    if ttype not in VOUCHER_TYPES:
        return JsonResponse({"message": "Enter a valid Voucher Type (CV, BV or JV)."}, status=400)
    if not voucher_no:
        return JsonResponse({"message": "Enter a Voucher No."}, status=400)
    data = get_voucher(ttype, voucher_date, voucher_no)
    if not data:
        return JsonResponse({"message": "Voucher not found."}, status=404)
    return JsonResponse(data)


@login_required
@require_GET
def vp_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side

    ttype, voucher_date, voucher_no = _params(request)
    data = get_voucher(ttype, voucher_date, voucher_no) if ttype in VOUCHER_TYPES and voucher_no else None
    if not data:
        return JsonResponse({"message": "Voucher not found."}, status=404)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([settings.BRAND.get("address_line1", "")])
    ws.append([f"VOUCHER DATE {voucher_date:%d/%m/%Y}"])
    ws.append([f"VOUCHER NO   {voucher_no}"])
    ws.append([f"VOUCHER TYPE {data['label']}"])
    ws.append([])
    header_row = ws.max_row + 1
    header = ["CODE", "S.CODE", "A/C DESCRIPTION", "DEBIT", "CREDIT"]
    ws.append(header)
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([8, 8, 40, 13, 13], "ABCDE"):
        ws.column_dimensions[col].width = w

    for l in data["lines"]:
        ws.append([l["code"], l["scode"], l["account"], float(l["debit"]) if l["debit"] else None,
                   float(l["credit"]) if l["credit"] else None])
        prefix = "TO " if l["dc"] == "D" else "BY "
        ws.append(["", "", prefix + l["narration1"], None, None])
        if l["narration2"]:
            ws.append(["", "", "   " + l["narration2"], None, None])

    last_row = ws.max_row
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=header_row, max_row=last_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border

    ws.append(["", "", "T O T A L", float(data["total_debit"]), float(data["total_credit"])])
    ws.append([])
    ws.append([])
    ws.append(["Signature", "", "", "", "Accountant"])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Voucher_{ttype}_{voucher_no}_{voucher_date:%d-%m-%Y}.xlsx"'
    return response
