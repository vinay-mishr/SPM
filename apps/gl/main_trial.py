"""Accounts > Trial Balance > Main Trial (VFP "MAIN TRIAL" report).

One row per ``gled_mst`` ledger as on a single RUN DATE: the same opening-
balance-style formula every other Ledger Printing report in this app
shares (``gyob`` + sum(Debit ``tamt``) - sum(Credit ``tamt``) for that
ledger's own ``tcode``), just dated ``<= RUN DATE`` instead of split into
"before start" / "in range" -- a trial balance is a single as-of snapshot,
not a period. A ledger with a net balance of exactly zero is dropped
entirely (the VFP source's own ``if mopbal>0 ... else if mopbal<0`` --
there's no branch for zero). A positive balance prints as DEBIT, negative
as CREDIT (``abs()``'d).

``gled_mst.subledger`` (a boolean flag marking an account that has its own
customer/vendor sub-ledger, printed as a leading ``*``) is used here exactly
as the pasted VFP source has it -- taken verbatim from this source like
apps.gl.main_ledger_group's own ``gl_type``/``group1`` were, not confirmed
against any other screen in this project (nothing else here has needed it).

Deliberately NOT reproduced: the VFP source's own side-effect writes back
``gytdd``/``gytdc`` onto each ``gled_mst`` row and accumulates ``amount``
onto the matching ``group1`` row (``gl_type``) as it scans -- those look
like they feed a separate "trial balance by group" screen that reads those
columns back later. This is a read-only report view; mutating master-table
columns as scratch state for every page load is not something a web report
should do, so those writes are skipped. Flagged, not guessed away quietly.
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows


def main_trial_rows(run_date):
    rows = []
    total_debit = total_credit = Decimal("0")
    for led in _rows("SELECT gcode, gdesc, gyob, subledger FROM gled_mst ORDER BY gdesc"):
        code = _clean(led["gcode"])
        balance = _d(led["gyob"])
        for r in _rows(
            "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tcode = %s AND tdate <= %s GROUP BY tdbcr",
            [code, run_date],
        ):
            if _clean(r["tdbcr"]) == "D":
                balance += _d(r["total"])
            elif _clean(r["tdbcr"]) == "C":
                balance -= _d(r["total"])

        if balance == 0:
            continue

        has_subledger = bool(led["subledger"])
        if balance > 0:
            total_debit += balance
            debit, credit = f"{balance:.2f}", ""
        else:
            total_credit += abs(balance)
            debit, credit = "", f"{abs(balance):.2f}"

        rows.append({
            "code": code, "desc": _clean(led["gdesc"]) or "", "has_subledger": has_subledger,
            "debit": debit, "credit": credit,
        })

    return rows, f"{total_debit:.2f}", f"{total_credit:.2f}", f"{abs(total_debit - total_credit):.2f}"


class MainTrialView(LoginRequiredMixin, TemplateView):
    template_name = "gl/main_trial.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


def _run_date(request):
    try:
        return date.fromisoformat(request.GET.get("run_date"))
    except (TypeError, ValueError):
        return date.today()


@login_required
@require_GET
def mt_data(request):
    run_date = _run_date(request)
    rows, total_debit, total_credit, difference = main_trial_rows(run_date)
    return JsonResponse({
        "run_date": run_date.isoformat(), "rows": rows,
        "total_debit": total_debit, "total_credit": total_credit, "difference": difference,
    })


@login_required
@require_GET
def mt_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    run_date = _run_date(request)
    rows, total_debit, total_credit, difference = main_trial_rows(run_date)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([])
    ws.append([f"Main Trial as on {run_date:%d/%m/%Y}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["Code", "Description", "Debit Amount", "Credit Amount"])
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 30, 15, 15], "ABCD"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([
            ("*" if r["has_subledger"] else " ") + r["code"], r["desc"],
            float(r["debit"]) if r["debit"] else None, float(r["credit"]) if r["credit"] else None,
        ])

    ws.append(["", "TOTAL AMOUNT", float(total_debit), float(total_credit)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append(["", "DIFFERENCE", float(difference)])
    ws.append(["", "*-ACCOUNT HAVING SUBLEDGER"])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Main_Trial_{run_date:%d-%m-%Y}.xlsx"'
    return response
