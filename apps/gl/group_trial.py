"""Accounts > Trial Balance > Group Trial (VFP "GROUPWISE MAIN TRIAL"
report).

Same per-ledger as-of-date balance formula as apps.gl.main_trial's own
"MAIN TRIAL" (``gyob`` + sum(Debit ``tamt``) - sum(Credit ``tamt``) for that
ledger's own ``tcode``, dated ``<= RUN DATE``, zero balances dropped,
positive=DEBIT/negative=CREDIT, ``gled_mst.subledger`` for the ``*``
prefix), just grouped under one ``<<group name>>`` heading per
``group1.grp_code`` (``gled_mst.gl_type``) -- the same group/gl_type
relationship already used by apps.gl.main_ledger_group. Groups with every
one of their ledgers at a zero balance are skipped entirely (no heading, no
subtotal -- the VFP source's own ``fl`` flag), same "nothing to show, print
nothing" rule as every other report in this app. Each group gets its own
DEBIT/CREDIT subtotal; the report ends with the grand total and the
difference.

``group1`` is iterated in its own natural/``grp_code`` order (the VFP
source's own ``use group1 in 0 order grp_code``, no explicit re-sort before
scanning); ledgers within a group by ``gcode`` (``set orde to gcode``).
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows


def group_trial_groups(run_date):
    groups = []
    grand_debit = grand_credit = Decimal("0")
    for grp in _rows("SELECT grp_code, grp_name FROM group1 ORDER BY grp_code"):
        grp_code = _clean(grp["grp_code"])
        ledgers = []
        total_debit = total_credit = Decimal("0")
        for led in _rows("SELECT gcode, gdesc, gyob, subledger FROM gled_mst WHERE gl_type = %s ORDER BY gcode", [grp_code]):
            code = _clean(led["gcode"])
            balance = _d(led["gyob"])
            for r in _rows(
                "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tcode = %s AND tdate <= %s GROUP BY tdbcr",
                [code, run_date],
            ):
                if _clean(r["tdbcr"]) == "D":
                    balance += _d(r["total"])
                elif _clean(r["tdbcr"]) == "C":
                    balance -= _d(r["total"])

            if balance == 0:
                continue

            has_subledger = bool(led["subledger"])
            if balance > 0:
                total_debit += balance
                debit, credit = f"{balance:.2f}", ""
            else:
                total_credit += abs(balance)
                debit, credit = "", f"{abs(balance):.2f}"

            ledgers.append({
                "code": code, "desc": _clean(led["gdesc"]) or "", "has_subledger": has_subledger,
                "debit": debit, "credit": credit,
            })

        if not ledgers:
            continue

        grand_debit += total_debit
        grand_credit += total_credit
        groups.append({
            "grp_code": grp_code, "grp_name": _clean(grp["grp_name"]) or "",
            "ledgers": ledgers, "total_debit": f"{total_debit:.2f}", "total_credit": f"{total_credit:.2f}",
        })

    return groups, f"{grand_debit:.2f}", f"{grand_credit:.2f}", f"{abs(grand_debit - grand_credit):.2f}"


class GroupTrialView(LoginRequiredMixin, TemplateView):
    template_name = "gl/group_trial.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


def _run_date(request):
    try:
        return date.fromisoformat(request.GET.get("run_date"))
    except (TypeError, ValueError):
        return date.today()


@login_required
@require_GET
def gt_data(request):
    run_date = _run_date(request)
    groups, grand_debit, grand_credit, difference = group_trial_groups(run_date)
    return JsonResponse({
        "run_date": run_date.isoformat(), "groups": groups,
        "grand_debit": grand_debit, "grand_credit": grand_credit, "difference": difference,
    })


@login_required
@require_GET
def gt_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    run_date = _run_date(request)
    groups, grand_debit, grand_credit, difference = group_trial_groups(run_date)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([])
    ws.append([f"Group Trial as on {run_date:%d/%m/%Y}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["Code", "Description", "Debit Amount", "Credit Amount"])
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 30, 15, 15], "ABCD"):
        ws.column_dimensions[col].width = w

    for g in groups:
        ws.append(["", f"<<{g['grp_name']}>>"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for l in g["ledgers"]:
            ws.append([
                ("*" if l["has_subledger"] else "'") + l["code"], l["desc"],
                float(l["debit"]) if l["debit"] else None, float(l["credit"]) if l["credit"] else None,
            ])
        ws.append(["", "T O T A L", float(g["total_debit"]), float(g["total_credit"])])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)

    ws.append([])
    ws.append(["", "T O T A L", float(grand_debit), float(grand_credit)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([f"Difference {difference}"])
    ws.append(["*-Account having Subledger"])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Group_Trial_{run_date:%d-%m-%Y}.xlsx"'
    return response
