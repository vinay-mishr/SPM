"""Accounts > Trial Balance > Month Trial (VFP "MONTH TRIAL" report).

Unlike apps.gl.main_trial's own "MAIN TRIAL" (a balance AS ON one date,
starting from ``gled_mst.gyob``), Month Trial is a pure NET MOVEMENT
between START DATE and END DATE -- the running total starts at 0 (no
opening balance folded in at all; the VFP source's own ``mopbal=0``), summed
from Debit/Credit ``tamt`` for that ledger's own ``tcode`` within the date
range. A ledger with a net movement of exactly zero is dropped, same "no
zero branch" rule as Main Trial; positive nets print as DEBIT, negative as
CREDIT. ``gled_mst.subledger`` (the ``*`` prefix) is the same confirmed-from-
this-source flag apps.gl.main_trial already uses.

Ledgers are iterated in ``gcode`` order here (the VFP source's own
``set orde to gcode``), unlike Main Trial's own ``gdesc`` order -- flagged
as a deliberate difference, not an inconsistency.

EXCEL / EXIT only -- the VFP form's own screenshot has no VIEW or PRINT.
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


def month_trial_rows(start, end):
    rows = []
    total_debit = total_credit = Decimal("0")
    for led in _rows("SELECT gcode, gdesc, subledger FROM gled_mst ORDER BY gcode"):
        code = _clean(led["gcode"])
        movement = Decimal("0")
        for r in _rows(
            "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tcode = %s AND tdate >= %s AND tdate <= %s GROUP BY tdbcr",
            [code, start, end],
        ):
            if _clean(r["tdbcr"]) == "D":
                movement += _d(r["total"])
            elif _clean(r["tdbcr"]) == "C":
                movement -= _d(r["total"])

        if movement == 0:
            continue

        has_subledger = bool(led["subledger"])
        if movement > 0:
            total_debit += movement
            debit, credit = f"{movement:.2f}", ""
        else:
            total_credit += abs(movement)
            debit, credit = "", f"{abs(movement):.2f}"

        rows.append({
            "code": code, "desc": _clean(led["gdesc"]) or "", "has_subledger": has_subledger,
            "debit": debit, "credit": credit,
        })

    return rows, f"{total_debit:.2f}", f"{total_credit:.2f}"


class MonthTrialView(LoginRequiredMixin, TemplateView):
    template_name = "gl/month_trial.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": _fy_start(today).isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), _fy_start(today))
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def mnt_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)

    rows, total_debit, total_credit = month_trial_rows(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([])
    ws.append([f" MONTHWISE TRIAL FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["Code", "Description", "Debit Amount", "Credit Amount"])
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 30, 15, 15], "ABCD"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([
            ("*" if r["has_subledger"] else "") + r["code"], r["desc"],
            float(r["debit"]) if r["debit"] else None, float(r["credit"]) if r["credit"] else None,
        ])

    ws.append(["", "T O T A L", float(total_debit), float(total_credit)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Month_Trial_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
