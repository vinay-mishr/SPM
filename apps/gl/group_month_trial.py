"""Accounts > Trial Balance > Groupwise Month Trial (VFP "GROUPWISE MAIN
TRIAL" report -- the date-range variant, ``grptrial.PRN``).

Same ``<<group name>>``-headed, per-group-subtotalled layout as
apps.gl.group_trial's own "GROUPWISE MAIN TRIAL" (as-of-one-date version),
but this is a pure NET MOVEMENT between START DATE and END DATE -- the
running total starts at 0, no ``gyob`` folded in at all (the VFP source's
own ``mopbal=0``), exactly the apps.gl.month_trial relationship to
apps.gl.main_trial. Zero-movement ledgers and groups-with-nothing-to-show
are both dropped, same rules as every other Trial Balance report.

``group1.grp_code`` is filtered to a fixed whitelist -- ``inli(grp_code,
'C01','C02','D02','D04','D03','D0102','D0101')`` -- taken verbatim from the
pasted VFP source, not guessed or generalized. If the list of groups this
report should cover ever changes, it has to change here too; it isn't
derived from anything else in the schema.

One more confirmed difference between this report's own PRN output and its
own Excel export (not normalized away): the printed/on-screen listing does
NOT show the ``gled_mst.subledger`` ``*`` marker at all (the VFP source's
own ``?gled_mst->gcode+...`` has no ``iif(subledger,...)``), while the
Excel export does. Both are reproduced exactly as given.
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows

GROUP_CODES = ["C01", "C02", "D02", "D04", "D03", "D0102", "D0101"]


def group_month_trial_groups(start, end):
    placeholders = ",".join(["%s"] * len(GROUP_CODES))
    groups = []
    grand_debit = grand_credit = Decimal("0")
    for grp in _rows(f"SELECT grp_code, grp_name FROM group1 WHERE grp_code IN ({placeholders}) ORDER BY grp_code", GROUP_CODES):
        grp_code = _clean(grp["grp_code"])
        ledgers = []
        total_debit = total_credit = Decimal("0")
        for led in _rows("SELECT gcode, gdesc, subledger FROM gled_mst WHERE gl_type = %s ORDER BY gcode", [grp_code]):
            code = _clean(led["gcode"])
            movement = Decimal("0")
            for r in _rows(
                "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tcode = %s AND tdate >= %s AND tdate <= %s GROUP BY tdbcr",
                [code, start, end],
            ):
                if _clean(r["tdbcr"]) == "D":
                    movement += _d(r["total"])
                elif _clean(r["tdbcr"]) == "C":
                    movement -= _d(r["total"])

            if movement == 0:
                continue

            has_subledger = bool(led["subledger"])
            if movement > 0:
                total_debit += movement
                debit, credit = f"{movement:.2f}", ""
            else:
                total_credit += abs(movement)
                debit, credit = "", f"{abs(movement):.2f}"

            ledgers.append({
                "code": code, "desc": _clean(led["gdesc"]) or "", "has_subledger": has_subledger,
                "debit": debit, "credit": credit,
            })

        if not ledgers:
            continue

        grand_debit += total_debit
        grand_credit += total_credit
        groups.append({
            "grp_code": grp_code, "grp_name": _clean(grp["grp_name"]) or "",
            "ledgers": ledgers, "total_debit": f"{total_debit:.2f}", "total_credit": f"{total_credit:.2f}",
        })

    return groups, f"{grand_debit:.2f}", f"{grand_credit:.2f}"


class GroupMonthTrialView(LoginRequiredMixin, TemplateView):
    template_name = "gl/group_month_trial.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def gmt_data(request):
    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    groups, grand_debit, grand_credit = group_month_trial_groups(start, end)
    return JsonResponse({
        "start_date": start.isoformat(), "end_date": end.isoformat(),
        "groups": groups, "grand_debit": grand_debit, "grand_credit": grand_credit,
    })


@login_required
@require_GET
def gmt_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)

    groups, grand_debit, grand_credit = group_month_trial_groups(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([])
    ws.append([f"Groupwise Trial from {start:%d/%m/%Y} to {end:%d/%m/%Y}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["Code", "Description", "Debit Amount", "Credit Amount"])
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 30, 15, 15], "ABCD"):
        ws.column_dimensions[col].width = w

    for g in groups:
        ws.append(["", f"<<{g['grp_name']}>>"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for l in g["ledgers"]:
            ws.append([
                ("*" if l["has_subledger"] else "'") + l["code"], l["desc"],
                float(l["debit"]) if l["debit"] else None, float(l["credit"]) if l["credit"] else None,
            ])
        ws.append(["", "T O T A L", float(g["total_debit"]), float(g["total_credit"])])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)

    ws.append([])
    ws.append(["", "T O T A L", float(grand_debit), float(grand_credit)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Groupwise_Month_Trial_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
