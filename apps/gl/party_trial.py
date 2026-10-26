"""Accounts > Trial Balance > Party Trial (VFP "PARTY TRIAL" report).

Six modes (radio buttons on the VFP form, ``thisform.tcom.value`` 1-6), all
sharing the exact same per-party balance formula apps.gl.party_ledger
already uses (``cyob`` + sum(Debit ``tamt``) - sum(Credit ``tamt``) for that
party's own ``tccode``, dated ``<= RUN DATE`` -- a trial-balance snapshot,
same "as of one date" shape as apps.gl.main_trial), just scoped to a
different subset of ``cust_mst`` and showing a different subset of columns:

  1. PARTY TRIAL BALANCE -- every party with a non-zero balance; both
     DEBIT and CREDIT columns.
  2. SUNDRY DR. STATEMENT -- parties in the Sundry Debtor GROUP
     (``cust_mst.code = PSDCODE``, confirmed as ``apps.parties.models.Party.
     group_code`` / ``settings.PARTY_DEBTOR_CODE``) with a debit balance;
     DEBIT column only.
  3. SUNDRY CR. STATEMENT -- Sundry Creditor GROUP
     (``settings.PARTY_CREDITOR_CODE``) with a credit balance; CREDIT only.
  4. ADVANCE FROM DR. -- Sundry Debtor GROUP parties with a CREDIT balance
     (a debtor who has paid in advance); CREDIT column only.
  5. ADVANCE TO SUPPLIER -- Sundry Creditor GROUP parties with a DEBIT
     balance (we've paid them in advance); DEBIT column only.
  6. SUB TRIAL LIST -- parties under one user-picked sub-ledger GROUP
     (any ``gled_mst`` row with ``subledger = True``, picked from a listbox
     the VFP form populates on selecting this mode); both DEBIT and CREDIT,
     same as mode 1.

Every mode orders parties by ``cname`` (the VFP source's own
``set orde to cname`` in every branch).
"""
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import _clean, _d, _rows

PARTY_DEBTOR_CODE = getattr(settings, "PARTY_DEBTOR_CODE", "S001")
PARTY_CREDITOR_CODE = getattr(settings, "PARTY_CREDITOR_CODE", "S002")

MODES = {
    1: {"title": "PARTY TRIAL AS ON", "columns": "both", "group_code": None, "include": "any"},
    2: {"title": "SUNDRY DEBTORS STATEMENT AS ON", "columns": "debit", "group_code": PARTY_DEBTOR_CODE, "include": "debit"},
    3: {"title": "SUNDRY CREDITORS STATEMENT AS ON", "columns": "credit", "group_code": PARTY_CREDITOR_CODE, "include": "credit"},
    4: {"title": "LIST OF ADVANCE FROM DEBTORS AS ON", "columns": "credit", "group_code": PARTY_DEBTOR_CODE, "include": "credit"},
    5: {"title": "LIST OF ADVANCE TO CREDITORS AS ON", "columns": "debit", "group_code": PARTY_CREDITOR_CODE, "include": "debit"},
    6: {"title": "LIST OF {name} AS ON", "columns": "both", "group_code": None, "include": "any"},
}


def search_subledger_groups():
    rows = _rows("SELECT gcode, gdesc FROM gled_mst WHERE subledger = 1 ORDER BY gcode")
    return [{"code": _clean(r["gcode"]), "name": _clean(r["gdesc"]) or ""} for r in rows]


def party_trial_rows(mode, run_date, group_code=None):
    spec = MODES[mode]
    filter_code = spec["group_code"] if mode != 6 else group_code
    if filter_code:
        parties = _rows("SELECT ccode, cname, cadd3, cyob FROM cust_mst WHERE code = %s ORDER BY cname", [filter_code])
    else:
        parties = _rows("SELECT ccode, cname, cadd3, cyob FROM cust_mst ORDER BY cname")

    rows = []
    total_debit = total_credit = Decimal("0")
    sno = 1
    for p in parties:
        ccode = _clean(p["ccode"])
        balance = _d(p["cyob"])
        for r in _rows(
            "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tccode = %s AND tdate <= %s GROUP BY tdbcr",
            [ccode, run_date],
        ):
            if _clean(r["tdbcr"]) == "D":
                balance += _d(r["total"])
            elif _clean(r["tdbcr"]) == "C":
                balance -= _d(r["total"])

        cytdd = balance if balance > 0 else Decimal("0")
        cytdc = abs(balance) if balance < 0 else Decimal("0")

        include = spec["include"]
        if include == "debit" and cytdd <= 0:
            continue
        if include == "credit" and cytdc <= 0:
            continue
        if include == "any" and cytdd <= 0 and cytdc <= 0:
            continue

        total_debit += cytdd
        total_credit += cytdc
        rows.append({
            "sno": sno, "code": ccode, "name": _clean(p["cname"]) or "", "city": _clean(p["cadd3"]) or "",
            "debit": f"{cytdd:.2f}" if cytdd > 0 else "", "credit": f"{cytdc:.2f}" if cytdc > 0 else "",
        })
        sno += 1

    return rows, f"{total_debit:.2f}", f"{total_credit:.2f}"


class PartyTrialView(LoginRequiredMixin, TemplateView):
    template_name = "gl/party_trial.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"mode": 1, "run_date": date.today().isoformat()}
        return context


def _params(request):
    try:
        mode = int(request.GET.get("mode") or 1)
    except ValueError:
        mode = 1
    if mode not in MODES:
        mode = 1
    try:
        run_date = date.fromisoformat(request.GET.get("run_date"))
    except (TypeError, ValueError):
        run_date = date.today()
    group_code = (request.GET.get("group_code") or "").strip()
    return mode, run_date, group_code


@login_required
@require_GET
def pt_groups(request):
    return JsonResponse({"groups": search_subledger_groups()})


@login_required
@require_GET
def pt_data(request):
    mode, run_date, group_code = _params(request)
    if mode == 6 and not group_code:
        return JsonResponse({"message": "Select a sub-ledger from the list."}, status=400)

    rows, total_debit, total_credit = party_trial_rows(mode, run_date, group_code)
    spec = MODES[mode]
    title = spec["title"]
    if mode == 6:
        group = _rows("SELECT gdesc FROM gled_mst WHERE gcode = %s", [group_code])
        group_name = _clean(group[0]["gdesc"]) if group else ""
        title = title.format(name=group_name)
    return JsonResponse({
        "run_date": run_date.isoformat(), "title": title, "columns": spec["columns"],
        "rows": rows, "total_debit": total_debit, "total_credit": total_credit,
    })


@login_required
@require_GET
def pt_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    mode, run_date, group_code = _params(request)
    if mode == 6 and not group_code:
        return JsonResponse({"message": "Select a sub-ledger from the list."}, status=400)

    rows, total_debit, total_credit = party_trial_rows(mode, run_date, group_code)
    spec = MODES[mode]
    title = spec["title"]
    if mode == 6:
        group = _rows("SELECT gdesc FROM gled_mst WHERE gcode = %s", [group_code])
        group_name = _clean(group[0]["gdesc"]) if group else ""
        title = title.format(name=group_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"{title} {run_date:%d/%m/%Y}"])
    ws.append([])
    header_row = ws.max_row + 1
    header = ["S.NO", "Ledger Description", "City"]
    if spec["columns"] in ("both", "debit"):
        header.append("Debit")
    if spec["columns"] in ("both", "credit"):
        header.append("Credit")
    ws.append(header)
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 35, 16, 14, 14], "ABCDE"):
        ws.column_dimensions[col].width = w

    for r in rows:
        row = [r["sno"], r["name"], r["city"]]
        if spec["columns"] in ("both", "debit"):
            row.append(float(r["debit"]) if r["debit"] else None)
        if spec["columns"] in ("both", "credit"):
            row.append(float(r["credit"]) if r["credit"] else None)
        ws.append(row)

    total_row = ["", "", "TOTAL"]
    if spec["columns"] in ("both", "debit"):
        total_row.append(float(total_debit))
    if spec["columns"] in ("both", "credit"):
        total_row.append(float(total_credit))
    ws.append(total_row)
    ws[ws.max_row][2].font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Party_Trial_{run_date:%d-%m-%Y}.xlsx"'
    return response
