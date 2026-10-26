"""Accounts > Data Entry > Cash In Hand (VFP "CASH IN HAND" report).

Daily running Cash ledger balance from a start date to an end date -- direct
port of the VFP PRN report. Two deliberate quirks carried over exactly as
the legacy source has them (not "fixed" here, flagged instead):

  * The OPENING balance (as on the day before Start Date) sums every
    ``trans`` row against the Cash ledger (``CASH_LEDGER_CODE``, apps.gl.
    cash_entry's own confirmed "C001"/"CASH IN HAND"), across every voucher
    type -- same formula as apps.gl.cash_entry.ledger_balance, just cut off
    before Start Date.
  * Each day's own MOVEMENT, however, only looks at ``ttype = 'CV'`` lines
    (apps.gl.cash_entry's own Cash Entry vouchers) whose own tcode ISN'T the
    Cash ledger itself -- i.e. the "other side" of every Cash Entry line,
    same convention apps.gl.cash_entry.list_entries already reads. A Bank
    Entry or Journal Entry line posted directly against the Cash ledger
    (legally possible -- nothing stops a user picking "CASH IN HAND" as a
    Bank Entry's own LEDGER/PARTY line, or a Journal Entry's own account)
    would NOT move this report's day-by-day balance, exactly matching the
    VFP source's own ``ttype='CV'`` filter -- not guessed, the pasted
    source really does scope only this one voucher type here, unlike the
    opening balance above it, which doesn't filter by ttype at all. This is
    a known inconsistency in the legacy report itself, carried over as-is.
"""
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cash_entry import CASH_LEDGER_CODE, _clean, _d, _rows


def cash_in_hand_rows(start, end):
    cash = _rows("SELECT gyob FROM gled_mst WHERE gcode = %s", [CASH_LEDGER_CODE])
    opening = _d(cash[0]["gyob"]) if cash else Decimal("0")

    dc = _rows(
        "SELECT tdbcr, SUM(tamt) AS total FROM trans WHERE tcode = %s AND tdate < %s GROUP BY tdbcr",
        [CASH_LEDGER_CODE, start],
    )
    for r in dc:
        if _clean(r["tdbcr"]) == "D":
            opening += _d(r["total"])
        elif _clean(r["tdbcr"]) == "C":
            opening -= _d(r["total"])

    daily = {}
    for r in _rows(
        "SELECT tdate, tdbcr, SUM(tamt) AS total FROM trans "
        "WHERE ttype = 'CV' AND tcode <> %s AND tdate >= %s AND tdate <= %s GROUP BY tdate, tdbcr",
        [CASH_LEDGER_CODE, start, end],
    ):
        day = r["tdate"]
        bucket = daily.setdefault(day, {"D": Decimal("0"), "C": Decimal("0")})
        bucket[_clean(r["tdbcr"])] = _d(r["total"])

    rows = []
    balance = opening
    sno = 1
    day = start
    while day <= end:
        bucket = daily.get(day, {"D": Decimal("0"), "C": Decimal("0")})
        balance = balance - bucket["D"] + bucket["C"]
        rows.append({
            "sno": sno, "date": day.isoformat(), "amount": f"{abs(balance):.2f}",
            "side": "Dr." if balance > 0 else "Cr." if balance < 0 else "",
        })
        sno += 1
        day += timedelta(days=1)
    return rows


class CashInHandView(LoginRequiredMixin, TemplateView):
    template_name = "gl/cash_in_hand.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        fy_start = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
        context["initial_data"] = {"start_date": fy_start.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    fy_start = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
    start = parse(request.GET.get("start_date"), fy_start)
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def cih_data(request):
    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse({"rows": cash_in_hand_rows(start, end)})


@login_required
@require_GET
def cih_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _params(request)
    rows = cash_in_hand_rows(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([settings.BRAND.get("address_line1", "")])
    ws.append([f"CASH IN HAND {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    header = ["S.NO", "DATE", "AMOUNT", ""]
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 14, 14, 6], "ABCD"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([r["sno"], r["date"], float(r["amount"]), r["side"]])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Cash_In_Hand_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
