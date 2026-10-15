"""Debit Note Statement: all debit notes dated in a range, with totals (VFP DNSTMNT)."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .models import DebitNote


def _parse(value, default):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return default


def statement_rows(start, end):
    rows, total = [], Decimal("0")
    for n in DebitNote.objects.filter(dn_date__gte=start, dn_date__lte=end).order_by("dn_date", "dn_no"):
        rows.append({"dn_no": n.dn_no, "dn_date": n.dn_date.isoformat(), "bill_no": n.invoice_no,
                     "bill_date": n.invoice_date.isoformat() if n.invoice_date else "",
                     "amount": f"{n.total_amount:.2f}", "party": n.party_name, "city": n.party_city})
        total += n.total_amount
    return rows, f"{total:.2f}"


class DebitNoteStatementView(LoginRequiredMixin, TemplateView):
    template_name = "debit_note/debit_note_statement.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def statement_data(request):
    today = date.today()
    start = _parse(request.GET.get("start_date"), today)
    end = _parse(request.GET.get("end_date"), today)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    rows, total = statement_rows(start, end)
    return JsonResponse({"rows": rows, "total": total})


@login_required
@require_GET
def statement_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    today = date.today()
    start = _parse(request.GET.get("start_date"), today)
    end = _parse(request.GET.get("end_date"), today)
    rows, total = statement_rows(start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"DEBIT NOTE STATEMENT FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["SNO", "D.N. NO.", "D.N. DATE", "BILL NO.", "BILL DATE", "PARTY NAME", "CITY", "AMOUNT"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 10, 12, 14, 12, 34, 18, 14], "ABCDEFGH"):
        ws.column_dimensions[col].width = w
    for i, r in enumerate(rows, start=1):
        ws.append([i, r["dn_no"], r["dn_date"], r["bill_no"], r["bill_date"], r["party"], r["city"], float(r["amount"])])
    ws.append([None, None, None, None, None, None, "TOTAL", float(total)])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Debit_Note_Statement.xlsx"'
    return response
