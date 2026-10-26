"""Money Receipt > Cheque Report (VFP "CHEQUE REPORT" form).

Every ``chq_file`` row (the same table apps.money_receipt.cheque_master
reads/writes) in a date range, split into SBI/HDFC columns by ``mop`` --
direct port of the VFP Excel macro. Unlike Cheque Master's own list, this
report's source has no ``type = 'S'`` filter of its own, so none is added
here either -- every cheque record in range is included, whichever screen
created it.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .cheque_master import BANKS
from .services import _clean, _d, _rows


def cheque_report_rows(start, end):
    rows = _rows(
        "SELECT f.chq_no, f.chq_dt, f.amount, f.mop, f.p_code, c.cname FROM chq_file f "
        "LEFT JOIN cust_mst c ON c.ccode = f.p_code "
        "WHERE f.chq_dt >= %s AND f.chq_dt <= %s ORDER BY f.chq_dt",
        [start, end],
    )
    out = []
    sbi_total = hdfc_total = Decimal("0")
    for r in rows:
        amount = _d(r["amount"])
        mop = int(_d(r["mop"])) if r["mop"] is not None else 0
        is_sbi = mop == 1
        if is_sbi:
            sbi_total += amount
        else:
            hdfc_total += amount
        out.append({
            "chq_no": _clean(r["chq_no"]) or "", "chq_dt": r["chq_dt"].isoformat() if r["chq_dt"] else "",
            "party_name": _clean(r["cname"]) or "" if _clean(r["p_code"]) else "",
            "bank": BANKS.get(mop, ""),
            "sbi_amount": f"{amount:.2f}" if is_sbi else "",
            "hdfc_amount": f"{amount:.2f}" if not is_sbi else "",
        })
    return out, {"sbi": f"{sbi_total:.2f}", "hdfc": f"{hdfc_total:.2f}"}


class ChequeReportView(LoginRequiredMixin, TemplateView):
    template_name = "money_receipt/cheque_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def cheque_report_data(request):
    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Starting Date cannot be later than Ending Date."}, status=400)
    rows, totals = cheque_report_rows(start, end)
    return JsonResponse({"rows": rows, "totals": totals})


@login_required
@require_GET
def cheque_report_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side

    start, end = _params(request)
    rows, totals = cheque_report_rows(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    header_row = ws.max_row + 1
    ws.append([f"CHEQUE REPORT FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    header_row += 1
    header = ["SNO", "CHEQUE NO", "CHEQUE DATE", "PARTY NAME", "BANK NAME", "SBI", "HDFC"]
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([5, 12, 14, 24, 14, 12, 12], "ABCDEFG"):
        ws.column_dimensions[col].width = w

    for i, r in enumerate(rows, start=1):
        ws.append([
            i, r["chq_no"], r["chq_dt"], r["party_name"], r["bank"],
            float(r["sbi_amount"]) if r["sbi_amount"] else None,
            float(r["hdfc_amount"]) if r["hdfc_amount"] else None,
        ])

    ws.append(["", "", "", "", "Total", float(totals["sbi"]), float(totals["hdfc"])])

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = border

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Cheque_Report_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
