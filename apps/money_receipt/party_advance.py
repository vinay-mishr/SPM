"""Money Receipt > Party Advance List (VFP "PARTY ADVANCE LIST" form).

Every party with an unadjusted advance (``adv_fil``, the same table
apps.money_receipt.services.previous_advance/apps.money_receipt.advance
already read/write), with its name/city joined from ``cust_mst`` -- a direct
port of the VFP PRN/Excel report. No filters in the source form itself
(screenshot: just VIEW/PRINT/EXCEL/EXIT) -- PRINT is dropped, same call
already made for every other report in this app (EXCEL is the modern
equivalent; nothing here can print to a physical printer).
"""
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import _clean, _d, _rows


def party_advance_rows():
    """adv_fil joined to cust_mst -- the VFP source skips any adv_fil row
    whose party isn't found in cust_mst (``if found()``), so this does too."""
    rows = _rows(
        "SELECT a.p_code, a.amount, c.cname, c.cadd3 FROM adv_fil a "
        "JOIN cust_mst c ON c.ccode = a.p_code ORDER BY c.cname",
    )
    out = []
    total = Decimal("0")
    for r in rows:
        amount = _d(r["amount"])
        total += amount
        out.append({
            "party_code": _clean(r["p_code"]), "party_name": _clean(r["cname"]) or "",
            "city": _clean(r["cadd3"]) or "", "amount": f"{amount:.2f}",
        })
    return out, f"{total:.2f}"


class PartyAdvanceListView(LoginRequiredMixin, TemplateView):
    template_name = "money_receipt/party_advance_list.html"


@login_required
@require_GET
def party_advance_data(request):
    rows, total = party_advance_rows()
    return JsonResponse({"rows": rows, "total": total})


@login_required
@require_GET
def party_advance_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side

    rows, total = party_advance_rows()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([settings.BRAND.get("address_line1", "")])
    ws.append(["PARTYWISE ADVANCE LIST"])
    ws.append([])
    header = ["PARTY NAME", "CITY", "AMOUNT"]
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([30, 20, 14], "ABC"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([r["party_name"], r["city"], float(r["amount"])])
    ws.append(["", "GRAND TOTAL", float(total)])

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.border = border

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Party_Advance_List.xlsx"'
    return response
