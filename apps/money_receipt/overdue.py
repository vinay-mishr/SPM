"""Money Receipt > Overdue Payment (VFP "OVERDUE PAYMENT" form).

Every pending bill (``bill``, same table/columns
apps.money_receipt.services.pending_bills already reads) that is past its
party's own credit period (``cust_mst.cr_days``, the same column
apps.sale_order.services already reads for this exact purpose), as on a
chosen date -- direct port of the VFP Excel macro. ALL/SINGLE party scope,
no date filter besides the as-on date, matching the screenshot.
"""
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import _clean, _d, _rows


def overdue_report(run_date, party_code=None):
    # Batched instead of one `bill` query per customer (the N+1 pattern that
    # made ALL take tens of seconds with hundreds of customers -- see the
    # identical fix in services.party_outstanding_report). One bill query,
    # grouped by party in Python, same formulas/results either way.
    party_sql = "SELECT ccode, cname, cadd3, cr_days FROM cust_mst"
    params = []
    if party_code:
        party_sql += " WHERE ccode = %s"
        params.append(party_code)
    party_sql += " ORDER BY cname"
    party_rows = _rows(party_sql, params)

    bills_by_party = {}
    bill_sql = "SELECT p_code, div, bill_no, b_date, n_amount, amt_rcd, c_note1, d_note1, int_amt, c_dis FROM bill WHERE pmt_rcd <> 'Y'"
    bill_params = []
    if party_code:
        bill_sql += " AND p_code = %s"
        bill_params.append(party_code)
    bill_sql += " ORDER BY p_code, b_date"
    for b in _rows(bill_sql, bill_params):
        bills_by_party.setdefault(_clean(b["p_code"]), []).append(b)

    groups = []
    grand_total = Decimal("0")
    for p in party_rows:
        ccode = _clean(p["ccode"])
        cr_days = int(_d(p["cr_days"]))
        bills = bills_by_party.get(ccode, [])
        rows_out = []
        party_total = Decimal("0")
        for b in bills:
            b_date = b["b_date"]
            if not b_date:
                continue
            due_date = date(b_date.year, b_date.month, b_date.day)
            delay_days = (run_date - (due_date + timedelta(days=cr_days))).days
            if delay_days <= 0:
                continue
            balance = (
                _d(b["n_amount"]) - _d(b["amt_rcd"]) - _d(b["c_note1"]) + _d(b["d_note1"])
                + _d(b["int_amt"]) - _d(b["c_dis"])
            )
            party_total += balance
            rows_out.append({
                "div": _clean(b["div"]) or "", "bill_no": _clean(b["bill_no"]),
                "bill_date": b_date.isoformat(), "bill_amount": f"{_d(b['n_amount']):.2f}",
                "amount_received": f"{_d(b['amt_rcd']):.2f}", "balance": f"{balance:.2f}",
                "delay_days": delay_days,
            })
        if not rows_out:
            continue
        groups.append({
            "party_code": ccode, "party_name": _clean(p["cname"]) or "", "city": _clean(p["cadd3"]) or "",
            "rows": rows_out, "total_balance": f"{party_total:.2f}",
        })
        grand_total += party_total

    return {"groups": groups, "grand_total": f"{grand_total:.2f}"}


class OverduePaymentView(LoginRequiredMixin, TemplateView):
    template_name = "money_receipt/overdue_payment.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


def _params(request):
    try:
        run_date = date.fromisoformat(request.GET.get("run_date") or "")
    except ValueError:
        run_date = date.today()
    party_code = (request.GET.get("party_code") or "").strip().upper() or None
    return run_date, party_code


@login_required
@require_GET
def overdue_data(request):
    run_date, party_code = _params(request)
    return JsonResponse(overdue_report(run_date, party_code))


@login_required
@require_GET
def overdue_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side

    run_date, party_code = _params(request)
    data = overdue_report(run_date, party_code)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([settings.BRAND.get("address_line2", "")])
    ws.append([f"OVERDUE PAYMENT AS ON {run_date:%d/%m/%Y}"])
    ws.append([])
    header = ["DIV", "BILL NO.", "BILL DATE", "BILL AMOUNT", "AMT.RECEIVED", "BAL.AMOUNT.", "DELAY DAYS."]
    header_row = ws.max_row + 1
    ws.append(header)
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([5, 10, 12, 13, 13, 12, 11], "ABCDEFG"):
        ws.column_dimensions[col].width = w

    for g in data["groups"]:
        ws.append([f"{g['party_name']},{g['city']},{g['party_code']}"])
        for r in g["rows"]:
            ws.append([r["div"], r["bill_no"], r["bill_date"], float(r["bill_amount"]),
                       float(r["amount_received"]), float(r["balance"]), r["delay_days"]])
        ws.append(["", "", "", "", "TOTAL", float(g["total_balance"]), ""])

    last_row = ws.max_row
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=header_row, max_row=last_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = border

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Overdue_Payment_{run_date:%d-%m-%Y}.xlsx"'
    return response
