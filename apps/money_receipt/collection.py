"""Money Receipt > Collection Report (VFP "COLLECTION REPORT" form).

One row per bill line actually collected (``col``, same table
apps.money_receipt.services already reads/writes -- see its own docstring
for the confirmed column list), for every receipt whose ``slip_dt`` falls in
the chosen range, optionally scoped to a single party -- a direct port of
the pasted VFP PRN report and its Excel-export twin (which agree with each
other once their aliasing is worked out: PRN's ``a->``/``b->``/``c->``/``d->``
are VFP's built-in work-area letters A-J, available for any open table
regardless of its own alias, for work areas 1-4 -- cust_mst/colh/col/bill in
that order -- while the Excel macro just spells the same areas out by table
name instead).

Each col row's own ``amount`` (the amount actually adjusted against that one
bill, not the receipt's total) is shown under CASH or CHEQUE depending on
the owning receipt's mode -- found via ``colh`` joined on
``(p_code, chq_no) = (col.p_code, col.nar)``, the same linking convention
apps.money_receipt.services.get_receipt/delete_receipt already use, and
``colh.slip_no`` is this app's own established place for the receipt mode
("CA"/"CH"/"RT" -- see services.save_receipt's own colh INSERT). A row's
``col.chk_dis`` says what kind of row it is: blank is a real Money Receipt
allocation (counted under CASH/CHEQUE), 'A' is an Advance Adjustment line
(apps.money_receipt.advance -- no actual money moved, so excluded from both
CASH/CHEQUE and the dishonour total, counted separately instead), and 'Y'
means apps.money_receipt.dishonour has since reversed it (moved to its own
"CHQ DISHONOUR AMT" total, marked with '#' on the row, same as the VFP
source).
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import _clean, _d, _rows


def collection_rows(start, end, party_code=None):
    sql = (
        "SELECT c.p_code, c.div, c.bill_no, c.b_date, c.b_amt, c.c_dis, c.int_amt, c.amount, c.chk_dis, "
        "c.slip_dt, p.cname, p.cadd3, h.slip_no, h.chq_no, h.chq_dt "
        "FROM col c LEFT JOIN cust_mst p ON p.ccode = c.p_code "
        "LEFT JOIN colh h ON h.p_code = c.p_code AND h.chq_no = c.nar "
        "WHERE c.slip_dt >= %s AND c.slip_dt <= %s"
    )
    params = [start, end]
    if party_code:
        sql += " AND c.p_code = %s"
        params.append(party_code)
    sql += " ORDER BY c.slip_dt, c.p_code"

    rows = []
    totals = {"cash": Decimal("0"), "cheque": Decimal("0"), "cash_discount": Decimal("0"),
              "interest": Decimal("0"), "dishonoured": Decimal("0"), "advance_adjusted": Decimal("0")}
    for r in _rows(sql, params):
        chk_dis = _clean(r["chk_dis"])
        dishonoured = chk_dis == "Y"
        is_advance = chk_dis == "A"
        amount = _d(r["amount"])
        is_cash = _clean(r["slip_no"]) == "CA"
        b_date = r["b_date"]
        slip_dt = r["slip_dt"]
        days = (slip_dt - b_date).days if slip_dt and b_date else None
        rows.append({
            "party_code": _clean(r["p_code"]), "party_name": _clean(r["cname"]) or "",
            "city": _clean(r["cadd3"]) or "",
            "cash_amount": f"{amount:.2f}" if is_cash and not is_advance else "",
            "cheque_amount": f"{amount:.2f}" if not is_cash and not is_advance else "",
            "chq_no": _clean(r["chq_no"]) or "", "chq_dt": r["chq_dt"].isoformat() if r["chq_dt"] else "",
            "div": _clean(r["div"]) or "", "bill_no": _clean(r["bill_no"]) or "",
            "bill_date": b_date.isoformat() if b_date else "",
            "bill_amount": f"{_d(r['b_amt']):.2f}", "cash_discount": f"{_d(r['c_dis']):.2f}",
            "interest_amt": f"{_d(r['int_amt']):.2f}", "days": days if days is not None else "",
            "dishonoured": dishonoured, "advance_adjusted": is_advance,
        })
        if dishonoured:
            totals["dishonoured"] += amount
        elif is_advance:
            totals["advance_adjusted"] += amount
        else:
            totals["cash" if is_cash else "cheque"] += amount
            totals["cash_discount"] += _d(r["c_dis"])
            totals["interest"] += _d(r["int_amt"])
    return rows, {k: f"{v:.2f}" for k, v in totals.items()}


class CollectionReportView(LoginRequiredMixin, TemplateView):
    template_name = "money_receipt/collection_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    party_code = (request.GET.get("party_code") or "").strip().upper() or None
    return start, end, party_code


@login_required
@require_GET
def collection_data(request):
    start, end, party_code = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    rows, totals = collection_rows(start, end, party_code)
    return JsonResponse({"rows": rows, "totals": totals})


@login_required
@require_GET
def collection_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end, party_code = _params(request)
    rows, totals = collection_rows(start, end, party_code)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"COLLECTION REPORT FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    header = ["Code", "Party Name", "City", "Cash", "Cheque", "Chq No/Receipt No", "Chq/Receipt Date",
               "Inv.No.", "Inv.Date", "Inv.Amt.", "C.Dis", "Days", "Chq Dis"]
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 28, 16, 10, 10, 12, 12, 10, 10, 10, 8, 6, 8], "ABCDEFGHIJKLM"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([
            r["party_code"], r["party_name"], r["city"],
            float(r["cash_amount"]) if r["cash_amount"] else None,
            float(r["cheque_amount"]) if r["cheque_amount"] else None,
            r["chq_no"], r["chq_dt"], r["bill_no"], r["bill_date"],
            float(r["bill_amount"]), float(r["cash_discount"]), r["days"],
            "#" if r["dishonoured"] else "",
        ])

    ws.append(["", "", "TOTAL", float(totals["cash"]), float(totals["cheque"]), "", "", "", "",
               "", "", "", ""])
    if float(totals["dishonoured"]) > 0:
        ws.append(["", "", "CHQ DISHONOUR AMT", "", float(totals["dishonoured"]),
                   "# - Dishonour Cheques", "", "", "", "", "", "", ""])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Collection_Report_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
