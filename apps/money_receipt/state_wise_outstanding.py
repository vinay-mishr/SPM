"""Reports > Sales Reports > State Wise Outstanding (VFP form titled
"STATEWISE REPWISE OUTSTANDING" on screen, but its own code and field
labels -- DIVISION / STATE / RUN DATE, ALL/SINGLE scoped to STATE, not a
rep -- are the "STATEWISE PARTY WISE OUTSTANDING" report; the window's own
title bar text doesn't match what the form actually does. Ported as what
the code does: state-wise, not rep-wise.

One row per outstanding ``bill`` line (``pmt_rcd <> 'Y'``, dated on or
before RUN DATE, scoped to one DIVISION or all of them), grouped STATE
(``cust_mst.cstate``, ALL or one via the STATE picker) then PARTY
(``cust_mst`` rows with that state, excluding ``active = 'N'``) --
preserved order and skip rules exactly: a bill with a zero balance is
skipped entirely (the newer, authoritative Excel block's own
``if mbalance <> 0``, unlike the older PRN print routine pasted alongside
it, which printed every row regardless); a party/state with no qualifying
bills prints no header or subtotal at all.

Aging is bucketed off ``bill.gr_dt`` (Goods Receipt Date), NOT the bill's
own date -- exactly what the VFP source reads (``daydiff = RUN DATE -
gr_dt``). A bill with no ``gr_dt`` has no well-defined day-difference in
the source either; this port buckets it as "Above 90 Days" rather than
crashing or silently dropping it, which the VFP source itself never had
to decide (date arithmetic against an empty date there raises its own
error).

DIVISION here is free text matched against ``bill.div`` (plain equality,
case preserved as typed) -- a VFP validation side-handler on the same
field hard-codes the choices 'SE'/'AP'/'VS'/'A', but no division code in
this project's real ``bill`` rows matches any of those (every row seen in
this live data carries 'SP' or 'OL' -- see apps.gst_invoice.services' own
``DIV = "SP"``), so that whitelist is not reproduced; "A" still means "all
divisions", matching the source's own special-case.

The VFP source also prints a per-division company name/address
(``&mcom``/``&mcity``, indirect macros reading globals like ``pcomSE``)
instead of the one fixed company header every other report in this
project uses (``settings.BRAND``) -- not reproduced; this report's Excel
header uses ``settings.BRAND`` like the rest, since there is nowhere in
this project's settings to keep a second company identity per division.
"""
from datetime import date

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import connections
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import _clean, _d, _rows

ALL_DIVISIONS = "A"
ALL_STATES = "A"


def _state_scope_rows(state_code):
    sql = "SELECT stcode, stname FROM state"
    params = []
    if state_code and state_code != ALL_STATES:
        sql += " WHERE stcode = %s"
        params.append(state_code)
    sql += " ORDER BY stname"
    return _rows(sql, params)


def state_wise_outstanding(run_date, division, state_code):
    states = _state_scope_rows(state_code)
    division = (division or ALL_DIVISIONS).strip()
    all_divisions = division.upper() == ALL_DIVISIONS

    groups = []
    grand_totals = {"below_45": 0, "between_45_60": 0, "between_60_90": 0, "above_90": 0, "balance": 0}

    for st in states:
        state_name = _clean(st["stname"]) or ""
        state_code_val = _clean(st["stcode"]) or ""
        customers = _rows(
            "SELECT ccode, cname, cadd3 FROM cust_mst WHERE cstate = %s AND COALESCE(active, '') <> 'N' "
            "ORDER BY cname",
            [state_code_val],
        )

        state_group = {"state_name": state_name, "state_code": state_code_val, "parties": [],
                        "totals": {"below_45": 0, "between_45_60": 0, "between_60_90": 0, "above_90": 0, "balance": 0}}

        for cust in customers:
            party_code = _clean(cust["ccode"]) or ""
            sql = (
                "SELECT bill_no, div, b_date, n_amount, amt_rcd, c_dis, int_amt, d_note1, c_note1, gr_dt "
                "FROM bill WHERE COALESCE(pmt_rcd, '') <> 'Y' AND p_code = %s AND b_date <= %s"
            )
            params = [party_code, run_date]
            if not all_divisions:
                sql += " AND div = %s"
                params.append(division)
            bills = _rows(sql, params)

            lines = []
            totals = {"below_45": 0, "between_45_60": 0, "between_60_90": 0, "above_90": 0, "balance": 0}
            for b in bills:
                balance = (
                    _d(b["n_amount"]) - _d(b["amt_rcd"]) - _d(b["c_note1"]) + _d(b["d_note1"])
                    + _d(b["int_amt"]) - _d(b["c_dis"])
                )
                if balance == 0:
                    continue
                day_diff = (run_date - b["gr_dt"]).days if b["gr_dt"] else 9999
                bucket = "below_45" if day_diff <= 45 else "between_45_60" if day_diff <= 60 \
                    else "between_60_90" if day_diff <= 90 else "above_90"
                lines.append({
                    "div": _clean(b["div"]) or "", "bill_no": _clean(b["bill_no"]) or "",
                    "bill_date": b["b_date"].isoformat() if b["b_date"] else "",
                    "bill_amount": f"{_d(b['n_amount']):.2f}", "amount_received": f"{_d(b['amt_rcd']):.2f}",
                    "cash_discount": f"{_d(b['c_dis']):.2f}", "interest": f"{_d(b['int_amt']):.2f}",
                    "dr_note": f"{_d(b['d_note1']):.2f}", "cr_note": f"{_d(b['c_note1']):.2f}",
                    "balance": f"{balance:.2f}", "bucket": bucket,
                })
                totals[bucket] += balance
                totals["balance"] += balance

            if not lines:
                continue
            state_group["parties"].append({
                "party_code": party_code, "party_name": _clean(cust["cname"]) or "",
                "city": _clean(cust["cadd3"]) or "", "lines": lines,
                "totals": {k: f"{v:.2f}" for k, v in totals.items()},
            })
            for k in totals:
                state_group["totals"][k] += totals[k]

        if not state_group["parties"]:
            continue
        state_group["totals"] = {k: f"{v:.2f}" for k, v in state_group["totals"].items()}
        groups.append(state_group)
        for k in grand_totals:
            grand_totals[k] += float(state_group["totals"][k])

    return groups, {k: f"{v:.2f}" for k, v in grand_totals.items()}


class StateWiseOutstandingView(LoginRequiredMixin, TemplateView):
    template_name = "money_receipt/state_wise_outstanding.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat(), "division": ALL_DIVISIONS}
        return context


@login_required
@require_GET
def swo_states(request):
    return JsonResponse({"states": [
        {"code": _clean(r["stcode"]), "name": _clean(r["stname"])} for r in _rows("SELECT stcode, stname FROM state ORDER BY stname")
    ]})


def _params(request):
    try:
        run_date = date.fromisoformat(request.GET.get("run_date"))
    except (TypeError, ValueError):
        run_date = date.today()
    division = (request.GET.get("division") or ALL_DIVISIONS).strip() or ALL_DIVISIONS
    scope = (request.GET.get("scope") or "all").strip().lower()
    state_code = (request.GET.get("state_code") or "").strip().upper() if scope == "single" else ALL_STATES
    return run_date, division, scope, state_code


@login_required
@require_GET
def swo_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    run_date, division, scope, state_code = _params(request)
    if scope == "single" and not state_code:
        return JsonResponse({"message": "Select a state first."}, status=400)

    groups, grand_totals = state_wise_outstanding(run_date, division, state_code)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"STATEWISE PARTY WISE OUTSTANDING AS ON {run_date:%d/%m/%Y}"])
    ws.append([])

    header = ["DIV", "BILL NO.", "BILL DATE", "BILL AMOUNT", "AMT.RECEIVED", "CASH DIS.", "INT.AMT.",
               "DR.NOTE AMT.", "CR.NOTE AMT.", "BALANCE AMT.", "BELOW 45 DAYS", "BETWEEN 45 TO 60",
               "BETWEEN 60 TO 90", "ABOVE 90"]
    widths = [20, 8, 10, 12, 12, 10, 10, 12, 12, 12, 13, 15, 15, 13]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip(widths, "ABCDEFGHIJKLMN"):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["C"].number_format = "dd/mm/yyyy"

    for g in groups:
        ws.append([f"{g['state_name']},{g['state_code']}"])
        for p in g["parties"]:
            ws.append([f"{p['party_name']},{p['city']},{p['party_code']}"])
            for l in p["lines"]:
                row = [l["div"], l["bill_no"], date.fromisoformat(l["bill_date"]) if l["bill_date"] else None,
                       float(l["bill_amount"]), float(l["amount_received"]) or None,
                       float(l["cash_discount"]) or None, float(l["interest"]) or None,
                       float(l["dr_note"]) or None, float(l["cr_note"]) or None, float(l["balance"])]
                slot = {"below_45": 10, "between_45_60": 11, "between_60_90": 12, "above_90": 13}[l["bucket"]]
                row += [None] * (slot - len(row))
                row.append(float(l["balance"]))
                ws.append(row)
            ws.append([
                "", "", "", "", "", "", "", "", "PARTY TOTAL", float(p["totals"]["balance"]),
                float(p["totals"]["below_45"]), float(p["totals"]["between_45_60"]),
                float(p["totals"]["between_60_90"]), float(p["totals"]["above_90"]),
            ])
            ws.cell(row=ws.max_row, column=9).font = Font(bold=True)
        ws.append([
            "", "", "", "", "", "", "", "", "STATE TOTAL", float(g["totals"]["balance"]),
            float(g["totals"]["below_45"]), float(g["totals"]["between_45_60"]),
            float(g["totals"]["between_60_90"]), float(g["totals"]["above_90"]),
        ])
        ws.cell(row=ws.max_row, column=9).font = Font(bold=True)

    ws.append([])
    ws.append(["", "", "", "", "S U M M A R Y"])
    ws.cell(row=ws.max_row, column=5).font = Font(bold=True)
    for label, key in [("BELOW 45 DAY", "below_45"), ("BETWEEN 45 TO 60", "between_45_60"),
                        ("BETWEEN 60 TO 90", "between_60_90"), ("ABOVE 90 DAYS", "above_90"),
                        ("TOTAL AMOUNT", "balance")]:
        ws.append(["", "", "", label, "", float(grand_totals[key])])
        ws.cell(row=ws.max_row, column=4).font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="State_Wise_Outstanding_{run_date:%d-%m-%Y}.xlsx"'
    return response
