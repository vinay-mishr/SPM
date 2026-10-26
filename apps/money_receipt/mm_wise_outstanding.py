"""Reports > Sales Reports > Mktg Manager Outstanding (VFP "MKTG MANAGER
WISE PARTY WISE OUTSTANDING" form, OLE-Excel export block).

Same shape as apps.money_receipt.state_wise_outstanding -- see that
module's own docstring for the shared aging/skip rules (``gr_dt``-based
45/60/90 buckets, a zero-balance bill skipped, an empty group prints
nothing) -- grouped Marketing Manager (``mm_mst``, ALL or one via the
``tpcode``/``mmhelp`` picker) then PARTY (every ``cust_mst`` row whose
``cmm`` names that manager, with no ``active`` exclusion here, unlike
State Wise Outstanding's own ``active <> 'N'``) then bill.

No DIVISION field or filter exists anywhere in this VFP source (unlike
State Wise Outstanding's own DIVISION box) -- every division's bills for
a matching party count here.
"""
from datetime import date

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import _clean, _d, _rows

ALL_MANAGERS = "A"


def _manager_scope_rows(manager_code):
    sql = "SELECT mcode, mname FROM mm_mst"
    params = []
    if manager_code and manager_code != ALL_MANAGERS:
        sql += " WHERE mcode = %s"
        params.append(manager_code)
    sql += " ORDER BY mname"
    return _rows(sql, params)


def mm_wise_outstanding(run_date, manager_code):
    managers = _manager_scope_rows(manager_code)

    groups = []
    grand_totals = {"below_45": 0, "between_45_60": 0, "between_60_90": 0, "above_90": 0, "balance": 0}

    for mm in managers:
        mm_name = _clean(mm["mname"]) or ""
        mm_code = _clean(mm["mcode"]) or ""
        customers = _rows(
            "SELECT ccode, cname, cadd3 FROM cust_mst WHERE cmm = %s ORDER BY cname", [mm_code],
        )

        mm_group = {"mm_name": mm_name, "mm_code": mm_code, "parties": [],
                    "totals": {"below_45": 0, "between_45_60": 0, "between_60_90": 0, "above_90": 0, "balance": 0}}

        for cust in customers:
            party_code = _clean(cust["ccode"]) or ""
            bills = _rows(
                "SELECT bill_no, b_date, n_amount, amt_rcd, c_dis, int_amt, d_note1, c_note1, gr_dt "
                "FROM bill WHERE COALESCE(pmt_rcd, '') <> 'Y' AND p_code = %s AND b_date <= %s",
                [party_code, run_date],
            )

            lines = []
            totals = {"below_45": 0, "between_45_60": 0, "between_60_90": 0, "above_90": 0, "balance": 0}
            for b in bills:
                balance = (
                    _d(b["n_amount"]) - _d(b["amt_rcd"]) - _d(b["c_note1"]) + _d(b["d_note1"])
                    + _d(b["int_amt"]) - _d(b["c_dis"])
                )
                if balance == 0:
                    continue
                day_diff = (run_date - b["gr_dt"]).days if b["gr_dt"] else 9999
                bucket = "below_45" if day_diff <= 45 else "between_45_60" if day_diff <= 60 \
                    else "between_60_90" if day_diff <= 90 else "above_90"
                lines.append({
                    "bill_no": _clean(b["bill_no"]) or "", "bill_date": b["b_date"].isoformat() if b["b_date"] else "",
                    "bill_amount": f"{_d(b['n_amount']):.2f}", "amount_received": f"{_d(b['amt_rcd']):.2f}",
                    "cash_discount": f"{_d(b['c_dis']):.2f}", "interest": f"{_d(b['int_amt']):.2f}",
                    "dr_note": f"{_d(b['d_note1']):.2f}", "cr_note": f"{_d(b['c_note1']):.2f}",
                    "balance": f"{balance:.2f}", "bucket": bucket,
                })
                totals[bucket] += balance
                totals["balance"] += balance

            if not lines:
                continue
            mm_group["parties"].append({
                "party_code": party_code, "party_name": _clean(cust["cname"]) or "",
                "city": _clean(cust["cadd3"]) or "", "lines": lines,
                "totals": {k: f"{v:.2f}" for k, v in totals.items()},
            })
            for k in totals:
                mm_group["totals"][k] += totals[k]

        if not mm_group["parties"]:
            continue
        mm_group["totals"] = {k: f"{v:.2f}" for k, v in mm_group["totals"].items()}
        groups.append(mm_group)
        for k in grand_totals:
            grand_totals[k] += float(mm_group["totals"][k])

    return groups, {k: f"{v:.2f}" for k, v in grand_totals.items()}


class MmWiseOutstandingView(LoginRequiredMixin, TemplateView):
    template_name = "money_receipt/mm_wise_outstanding.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


def _params(request):
    try:
        run_date = date.fromisoformat(request.GET.get("run_date"))
    except (TypeError, ValueError):
        run_date = date.today()
    scope = (request.GET.get("scope") or "all").strip().lower()
    manager_code = (request.GET.get("manager_code") or "").strip().upper() if scope == "single" else ALL_MANAGERS
    return run_date, scope, manager_code


@login_required
@require_GET
def mmo_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    run_date, scope, manager_code = _params(request)
    if scope == "single" and not manager_code:
        return JsonResponse({"message": "Select a marketing manager first."}, status=400)

    groups, grand_totals = mm_wise_outstanding(run_date, manager_code)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([settings.BRAND.get("address_line1", "")])
    ws.append([f"MKTG MANAGERWISE PARTY WISE OUTSTANDING AS ON {run_date:%d/%m/%Y}"])

    header = ["BILL NO.", "BILL DATE", "BILL AMOUNT", "AMT.RECEIVED", "CASH DIS.", "INT.AMT.", "DR.NOTE AMT.",
               "CR.NOTE AMT.", "BALANCE AMT.", "BELOW 45 DAYS", "BETWEEN 45 TO 60", "BETWEEN 60 TO 90", "ABOVE 90"]
    widths = [8, 10, 12, 12, 10, 10, 12, 12, 12, 13, 15, 15, 13]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip(widths, "ABCDEFGHIJKLM"):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["B"].number_format = "dd/mm/yyyy"

    for g in groups:
        ws.append([f"{g['mm_name']},{g['mm_code']}"])
        for p in g["parties"]:
            ws.append([f"{p['party_name']},{p['city']},{p['party_code']}"])
            for l in p["lines"]:
                row = [l["bill_no"], date.fromisoformat(l["bill_date"]) if l["bill_date"] else None,
                       float(l["bill_amount"]), float(l["amount_received"]) or None,
                       float(l["cash_discount"]) or None, float(l["interest"]) or None,
                       float(l["dr_note"]) or None, float(l["cr_note"]) or None, float(l["balance"])]
                slot = {"below_45": 9, "between_45_60": 10, "between_60_90": 11, "above_90": 12}[l["bucket"]]
                row += [None] * (slot - len(row))
                row.append(float(l["balance"]))
                ws.append(row)
            ws.append([
                "", "", "", "", "", "", "", "PARTY TOTAL", float(p["totals"]["balance"]),
                float(p["totals"]["below_45"]), float(p["totals"]["between_45_60"]),
                float(p["totals"]["between_60_90"]), float(p["totals"]["above_90"]),
            ])
            ws.cell(row=ws.max_row, column=8).font = Font(bold=True)
        ws.append([
            "", "", "", "", "", "", "", "M.M.  TOTAL", float(g["totals"]["balance"]),
            float(g["totals"]["below_45"]), float(g["totals"]["between_45_60"]),
            float(g["totals"]["between_60_90"]), float(g["totals"]["above_90"]),
        ])
        ws.cell(row=ws.max_row, column=8).font = Font(bold=True)

    ws.append(["", "", "", "S U M M A R Y"])
    ws.cell(row=ws.max_row, column=4).font = Font(bold=True)
    for label, key in [("BELOW 45 DAY", "below_45"), ("BETWEEN 45 TO 60", "between_45_60"),
                        ("BETWEEN 60 TO 90", "between_60_90"), ("ABOVE 90 DAYS", "above_90"),
                        ("TOTAL AMOUNT", "balance")]:
        ws.append(["", "", label, "", float(grand_totals[key])])
        ws.cell(row=ws.max_row, column=3).font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Mktg_Manager_Outstanding_{run_date:%d-%m-%Y}.xlsx"'
    return response
