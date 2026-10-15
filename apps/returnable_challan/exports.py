"""Party Wise Pending Challan -- Excel export (openpyxl)."""
from io import BytesIO


def build_pending_challan_workbook(data, company_name, start_date, end_date):
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side

    font = "Arial"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pending Challan"
    party = data["party"]
    ws.cell(row=1, column=1, value=company_name).font = Font(name=font, size=12, bold=True)
    ws.cell(row=2, column=1, value=f"PARTY WISE PENDING CHALLAN FROM {start_date.strftime('%d/%m/%Y')} TO {end_date.strftime('%d/%m/%Y')}").font = Font(name=font, size=10, bold=True)
    ws.cell(row=3, column=1, value=f"{party['name']}, {party['city']} ({party['code']})").font = Font(name=font, size=10, bold=True)
    headers = [("CH. NO", 10), ("CH. DATE", 12), ("ITEM CODE", 12), ("ITEM NAME", 36), ("QTY", 12)]
    for col, (h, w) in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col, value=h)
        cell.font = Font(name=font, size=9, bold=True)
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = w
    row = 6
    for r in data["rows"]:
        values = [r["chl_no"], r["date"], r["item_code"], r["item_name"], float(r["qty"])]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = Font(name=font, size=9)
            cell.border = border
            if col == 5:
                cell.number_format = "0.000"
        row += 1
    ws.freeze_panes = "A6"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
