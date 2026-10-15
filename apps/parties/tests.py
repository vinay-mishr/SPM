import json
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from . import exports, services
from .forms import PartyForm
from .models import Party

# A GSTIN with a correct check digit for state 03 (Punjab).
VALID_GSTIN = "03ABCDE1234F1Z" + services.gstin_check_digit("03ABCDE1234F1Z")


class CustMstTableMixin:
    """Party maps the old software's cust_mst (unmanaged), so tests create the table themselves."""

    @classmethod
    def setUpClass(cls):
        with connection.schema_editor() as editor:
            editor.create_model(Party)
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        with connection.schema_editor() as editor:
            editor.delete_model(Party)


class GstinTests(TestCase):
    def test_valid_gstin_passes(self):
        self.assertEqual(services.gstin_error(VALID_GSTIN), "")

    def test_published_example_check_digit(self):
        self.assertEqual(services.gstin_check_digit("27AAPFU0939F1Z"), "V")

    def test_bad_check_digit_and_format(self):
        wrong = VALID_GSTIN[:14] + ("A" if VALID_GSTIN[14] != "A" else "B")
        self.assertIn("check digit", services.gstin_error(wrong))
        self.assertIn("format", services.gstin_error("03ABCDE12341FZ5"))
        self.assertIn("15 characters", services.gstin_error("03ABC"))

    def test_mobile_normalisation(self):
        self.assertEqual(services.normalize_mobile("+91 98765-43210"), "9876543210")
        self.assertEqual(services.normalize_mobile("09876543210"), "9876543210")


@override_settings(COMPANY_STATE_CODE="03", PARTY_DEBTOR_CODE="S001", PARTY_CREDITOR_CODE="S002")
class PartyMasterTests(CustMstTableMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="operator", password="S3cure-pass!")

    def setUp(self):
        self.client.force_login(self.user)

    def post_party(self, url=None, **overrides):
        data = {
            "name": "  Sharma Packaging  ",
            "party_type": "SD",
            "address_1": "Plot 12, Phase 8",
            "city": "Mohali",
            "pin": "160055",
            "state_code": "03",
            "gst_no": VALID_GSTIN.lower(),
            "gst_category": "REGULAR",
            "mobile_1": "+91 98765 43210",
            "email": "Accounts@Sharma.in",
            "credit_days": "30",
            "is_active": "true",
            "supply_type": "",
            "permit": "false",
        }
        data.update(overrides)
        return self.client.post(url or reverse("parties:api_create"), data)

    def test_page_renders(self):
        response = self.client.get(reverse("parties:master"))
        self.assertContains(response, "Customer Name")
        self.assertContains(response, "Debtors</a>")

    def test_create_maps_to_cust_mst_columns(self):
        response = self.post_party()
        self.assertEqual(response.status_code, 200, response.content)
        party = Party.objects.get()
        self.assertEqual(party.party_code, "S001")  # first letter of the name + 001
        self.assertEqual(party.group_code, "S001")  # Sundry Debtor
        self.assertEqual(party.name, "SHARMA PACKAGING")
        self.assertEqual((party.city, party.pin), ("Mohali", "160055"))
        self.assertEqual((party.state_gst_code, party.state_name), ("03", "PUNJAB"))
        self.assertEqual((party.gst_no, party.gst_category, party.gst_category_code), (VALID_GSTIN, "REGULAR", "01"))
        self.assertEqual((party.mobile_1, party.email), ("9876543210", "accounts@sharma.in"))
        self.assertEqual((party.active, party.permit, party.supply_type), ("Y", "N", "W"))
        self.assertEqual(party.credit_days, Decimal("30"))

    def test_codes_follow_first_letter(self):
        self.post_party()
        self.post_party(name="Sunrise Soaps", gst_no="")
        self.post_party(name="Alpha Tools", gst_no="", party_type="SC")
        self.assertEqual(sorted(Party.objects.values_list("party_code", flat=True)), ["A001", "S001", "S002"])
        self.assertEqual(Party.objects.get(party_code="A001").group_code, "S002")

    def test_validation_messages(self):
        response = self.post_party(name=" ", pin="1234", mobile_1="12345", email="bad", credit_days="-2")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(response.json()["errors"]), {"name", "pin", "mobile_1", "email", "credit_days"})

    def test_gstin_must_match_state(self):
        response = self.post_party(state_code="06")
        self.assertEqual(response.status_code, 400)
        self.assertIn("selected state's code is 06", response.json()["errors"]["gst_no"][0]["message"])

    def test_state_filled_from_gstin(self):
        self.post_party(state_code="")
        self.assertEqual(Party.objects.get().state_gst_code, "03")

    def test_modify_same_row_and_old_values_tolerated(self):
        Party.objects.create(party_code="A034", group_code="S002", name="A&I ENGINEERS",
                             mobile_1="99880+2919", gst_no="03BAD", active="A", city="LUDHIANA")
        # Unchanged odd values (mobile, GST) must not block saving a change to the address.
        response = self.client.post(reverse("parties:api_update", args=["A034"]), {
            "name": "A&I ENGINEERS", "party_type": "SC", "address_1": "Focal Point", "city": "LUDHIANA",
            "mobile_1": "99880+2919", "gst_no": "03BAD", "is_active": "true", "permit": "false", "credit_days": "0",
        })
        self.assertEqual(response.status_code, 200, response.content)
        party = Party.objects.get(party_code="A034")
        self.assertEqual(
            (party.address_1, party.city, party.active, Party.objects.count()),
            ("Focal Point", "LUDHIANA", "A", 1),
        )

    def test_delete_blocked_by_transactions(self):
        self.post_party()
        with mock.patch.object(services, "has_transactions", return_value=True):
            response = self.client.post(reverse("parties:api_delete", args=["S001"]))
        self.assertEqual(response.status_code, 409)
        self.assertIn("Transactions Exist", response.json()["message"])
        self.assertTrue(Party.objects.exists())

    def test_delete_blocked_by_other_tables(self):
        self.post_party()
        with mock.patch.object(services, "tables_using", return_value=["order", "bill"]):
            response = self.client.post(reverse("parties:api_delete", args=["S001"]))
        self.assertEqual(response.status_code, 409)
        self.assertIn("order, bill", response.json()["message"])

    def test_delete(self):
        self.post_party()
        self.assertEqual(self.client.post(reverse("parties:api_delete", args=["S001"])).status_code, 200)
        self.assertFalse(Party.objects.exists())

    def test_next_code_preview(self):
        self.post_party()
        response = self.client.get(reverse("parties:api_next_code"), {"name": "Super Pack"})
        self.assertEqual(response.json()["code"], "S002")

    def test_writes_can_be_switched_off(self):
        with mock.patch.object(services, "writes_enabled", return_value=False):
            self.assertEqual(self.post_party().status_code, 403)
        self.assertFalse(Party.objects.exists())


class PartySerialiseTests(TestCase):
    def test_old_row_values_cleaned(self):
        party = Party(party_code="1001", group_code="S001", name="100 PERCENT", city="DELHI", phone_1="9953675804",
                      mobile_1=".", state_gst_code="07", state_name="DELHI", credit_days=Decimal("7"),
                      active="", supply_type="I", permit="N", email="x@y.in")
        data = services.serialize_party(party)
        self.assertEqual((data["id"], data["city"], data["phone_1"], data["mobile_1"]),
                         ("1001", "DELHI", "9953675804", ""))
        self.assertEqual((data["party_type"], data["is_active"], data["credit_days"]), ("SD", True, 7))

    def test_only_n_is_inactive(self):
        self.assertFalse(Party(active="N").is_active)
        self.assertTrue(Party(active="A").is_active)


class PartyExportTests(TestCase):
    def test_workbook_layout(self):
        rows = [{
            "code": "S001", "ccode": "A034", "cname": "A&I ENGINEERS", "cphno": "0172-2661100", "mobile": "9780496006",
            "cr_days": Decimal("30"), "pdc": "", "rpermit": "N", "gst_no": "03BSMPA1149R1ZK", "gsname": "03",
            "stname": "PUNJAB", "cadd1": "PLOT 12", "cadd2": "", "cadd3": "MOHALI", "cbadd3": "a@b.in",
        }]
        ws = load_workbook(BytesIO(exports.build_workbook("debtors", rows, "SP MACHINES"))).active
        self.assertEqual(ws["A1"].value, "SP MACHINES")
        self.assertEqual([c.value for c in ws[5]],
                         [1, "S001", "A034", "A&I ENGINEERS", "0172-2661100 9780496006", 30, None, "N",
                          "03BSMPA1149R1ZK", "03-PUNJAB"])
        self.assertEqual((ws["D6"].value, ws["E6"].value, ws["D7"].value), ("PLOT 12", "a@b.in", "MOHALI"))
        self.assertEqual((ws["I5"].number_format, ws["D5"].font.name, ws.column_dimensions["D"].width),
                         ("@", "Arial Narrow", 30))

    def test_export_needs_legacy_db(self):
        user = get_user_model().objects.create_user(username="exp", password="S3cure-pass!")
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse("parties:export", args=["debtors"])).status_code, 503)
        self.assertEqual(self.client.get(reverse("parties:export", args=["other"])).status_code, 404)


class PartyFormTests(TestCase):
    def test_unchanged_old_values_pass(self):
        original = {"gst_no": "03BAD", "mobile_1": ".", "pin": "12"}
        form = PartyForm({"name": "X", "party_type": "SD", "gst_no": "03BAD", "pin": "12",
                          "is_active": "true", "permit": "false"}, original=original)
        self.assertTrue(form.is_valid(), form.errors)


class OpeningBalanceTests(CustMstTableMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth.models import Permission

        User = get_user_model()
        cls.editor = User.objects.create_user(username="accounts", password="S3cure-pass!")
        cls.editor.user_permissions.add(Permission.objects.get(codename="change_opening_balance"))
        cls.viewer = User.objects.create_user(username="viewer", password="S3cure-pass!")

    def setUp(self):
        Party.objects.create(party_code="A002", group_code="S001", name="A.B.INDUSTRES", opening_balance=Decimal("79998"))
        Party.objects.create(party_code="A001", group_code="S002", name="A.S.ENGINEERS", opening_balance=Decimal("-180968"))
        Party.objects.create(party_code="1001", group_code="S001", name="100 PERCENT", opening_balance=None, gst_no=".")

    def save(self, changes, user=None, **extra):
        self.client.force_login(user or self.editor)
        return self.client.post(reverse("parties:opening_balance_save"), data=json.dumps({"changes": changes}),
                                content_type="application/json", **extra)

    def test_rows_split_debit_credit_sorted_by_name(self):
        from . import opening_balance as ob

        rows = ob.rows()
        self.assertEqual([r["code"] for r in rows], ["1001", "A002", "A001"])
        self.assertEqual((rows[1]["debit"], rows[1]["credit"]), ("79998.00", "0.00"))
        self.assertEqual((rows[2]["debit"], rows[2]["credit"]), ("0.00", "180968.00"))
        self.assertEqual((rows[0]["debit"], rows[0]["gst_no"]), ("0.00", ""))  # NULL -> 0, "." -> blank
        self.assertEqual(ob.totals(), {"debit": "79998.00", "credit": "180968.00"})

    def test_page_and_sidebar(self):
        self.client.force_login(self.viewer)
        response = self.client.get(reverse("parties:opening_balance"))
        self.assertContains(response, "Opening Balance Master")
        self.assertContains(response, "View only")
        self.client.force_login(self.editor)
        self.assertNotContains(self.client.get(reverse("parties:opening_balance")), "View only")

    def test_save_debit_and_credit(self):
        response = self.save([{"code": "A002", "debit": "0", "credit": "5000.50"},
                              {"code": "1001", "debit": "100", "credit": ""}])
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(Party.objects.get(pk="A002").opening_balance, Decimal("-5000.50"))
        self.assertEqual(Party.objects.get(pk="1001").opening_balance, Decimal("100.00"))
        self.assertEqual(response.json()["totals"], {"debit": "100.00", "credit": "185968.50"})

    def test_both_debit_and_credit_rejected_and_nothing_saved(self):
        response = self.save([{"code": "1001", "debit": "10", "credit": ""},
                              {"code": "A002", "debit": "10", "credit": "20"}])
        self.assertEqual(response.status_code, 400)
        self.assertIn("not both", response.json()["errors"]["A002"])
        self.assertIsNone(Party.objects.get(pk="1001").opening_balance)  # all-or-nothing

    def test_invalid_amounts_rejected(self):
        for bad in ("abc", "-5", "10.555", "1e3x", "123456789012"):
            response = self.save([{"code": "A002", "debit": bad, "credit": ""}])
            self.assertEqual(response.status_code, 400, bad)
        self.assertEqual(Party.objects.get(pk="A002").opening_balance, Decimal("79998"))

    def test_unknown_party_rejected(self):
        response = self.save([{"code": "ZZZZ", "debit": "1", "credit": ""}])
        self.assertEqual(response.status_code, 400)
        self.assertIn("ZZZZ", response.json()["errors"])

    def test_permission_required_to_save(self):
        response = self.save([{"code": "A002", "debit": "1", "credit": ""}], user=self.viewer)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(Party.objects.get(pk="A002").opening_balance, Decimal("79998"))

    def test_login_required(self):
        self.assertEqual(self.client.get(reverse("parties:opening_balance")).status_code, 302)
        self.assertEqual(self.client.get(reverse("parties:opening_balance_data")).status_code, 302)

    def test_csrf_enforced(self):
        from django.test import Client

        client = Client(enforce_csrf_checks=True)
        client.force_login(self.editor)
        response = client.post(reverse("parties:opening_balance_save"),
                               data=json.dumps({"changes": [{"code": "A002", "debit": "1"}]}),
                               content_type="application/json")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(Party.objects.get(pk="A002").opening_balance, Decimal("79998"))
