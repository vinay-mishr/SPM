"""Party Wise Pending Purchase Order — Excel export (openpyxl). Mirrors the old VFP
report's column layout (PARTY NAME, CITY, CODE, GRADE, ORD.NO, ORD.DATE, ITEM CODE,
ITEM NAME, PEND.QTY, RATE, VALUE, DES.DATE, DELAY DAYS)."""
from io import BytesIO

COLUMNS = [
    ("PARTY NAME", 30, "@"),
    ("CITY", 20, "@"),
    ("CODE", 8, "@"),
    ("GRADE", 8, "@"),
    ("ORD.NO", 10, "@"),
    ("ORD.DATE", 12, "dd/mm/yyyy"),
    ("ITEM CODE", 14, "@"),
    ("ITEM NAME", 30, "@"),
    ("PEND.QTY", 10, "0.00"),
    ("RATE", 10, "0.00"),
    ("VALUE", 12, "0.00"),
    ("DES.DATE", 12, "dd/mm/yyyy"),
    ("DELAY DAYS", 12, "0"),
]


def _date(iso):
    if not iso:
        return None
    from datetime import date
    y, m, d = iso.split("-")
    return date(int(y), int(m), int(d))


def build_workbook(groups, company_name, as_on_date):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pending PO"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=8, bold=True)
    ws.cell(row=2, column=1, value=f"PARTY WISE PENDING PURCHASE ORDER AS ON {as_on_date.strftime('%d/%m/%Y')}").font = Font(name=font_name, size=8, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=8, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for group in groups:
        for index, line in enumerate(group["lines"]):
            # Party Name/City/Code only on that party's first row -- not repeated
            # down every line of the same party.
            first = index == 0
            values = [
                group["party_name"] if first else "", group["city"] if first else "",
                group["party_code"] if first else "", group["grade"],
                line["ord_no"], _date(line["ord_date"]), line["item_code"], line["item_name"],
                float(line["pending_qty"]), float(line["rate"]), float(line["value"]),
                _date(line["despatch_date"]), line["delay_days"],
            ]
            for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, COLUMNS), start=1):
                cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
                cell.font = Font(name=font_name, size=8)
                cell.number_format = fmt
            row += 1

    for r in range(header_row, row):
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


DESPATCH_COLUMNS = [
    ("ORD.NO", 20, "@"),
    ("ORD.DATE", 12, "dd/mm/yyyy"),
    ("ITEM CODE", 14, "@"),
    ("ITEM NAME", 26, "@"),
    ("QTY", 10, "0.00"),
    ("DES.DATE", 12, "dd/mm/yyyy"),
    ("DELAY DAYS", 12, "0"),
]


def build_despatch_workbook(groups, company_name, run_date):
    """Party Wise Pending P.O. Despatch report. Matches the VFP layout: a party-heading
    row ("NAME,CITY,CODE" in column 1) followed by that party's item rows -- not one
    row per line with the party columns repeated."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pending Despatch"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=8, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"PARTY WISE PENDING PURCHASE ORDER DESPATCH REPORT AS ON {run_date.strftime('%d/%m/%Y')}",
    ).font = Font(name=font_name, size=8, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(DESPATCH_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=8, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for group in groups:
        heading_cell = ws.cell(
            row=row, column=1,
            value=f"{group['party_name']},{group['city']},{group['party_code']}",
        )
        heading_cell.font = Font(name=font_name, size=8, bold=True)
        row += 1
        for line in group["lines"]:
            values = [
                line["ord_no"], _date(line["ord_date"]), line["item_code"], line["item_name"],
                float(line["pending_qty"]), _date(line["despatch_date"]), line["delay_days"],
            ]
            for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, DESPATCH_COLUMNS), start=1):
                cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
                cell.font = Font(name=font_name, size=8)
                cell.number_format = fmt
            row += 1

    for r in range(header_row, row):
        for c in range(1, len(DESPATCH_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ADJUSTMENT_COLUMNS = [
    ("ORD.NO", 20, "@"),
    ("ORD.DATE", 12, "dd/mm/yyyy"),
    ("ITEM CODE", 14, "@"),
    ("ITEM NAME", 30, "@"),
    ("QTY", 10, "0.00"),
]


def build_po_adjustment_workbook(groups, company_name, run_date):
    """Party-Wise Purchase Order Adjustment report. Matches the VFP Excel macro's layout
    -- a party-heading row ("NAME,CITY,CODE" in column 1) followed by that party's
    already-adjusted lines, same convention as the despatch report above."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "PO Adjustment"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=8, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"PARTY WISE PURCHASE ORDER ADJUSTMENT AS ON {run_date.strftime('%d/%m/%Y')}",
    ).font = Font(name=font_name, size=8, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(ADJUSTMENT_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=8, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for group in groups:
        heading_cell = ws.cell(
            row=row, column=1,
            value=f"{group['party_name']},{group['city']},{group['party_code']}",
        )
        heading_cell.font = Font(name=font_name, size=8, bold=True)
        row += 1
        for line in group["lines"]:
            values = [
                line["ord_no"], _date(line["ord_date"]), line["item_code"], line["item_name"],
                float(line["pending_qty"]),
            ]
            for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, ADJUSTMENT_COLUMNS), start=1):
                cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
                cell.font = Font(name=font_name, size=8)
                cell.number_format = fmt
            row += 1

    for r in range(header_row, row):
        for c in range(1, len(ADJUSTMENT_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ITEM_WISE_COLUMNS = [
    ("ITEM NAME", 30, "@"),
    ("ORD.NO", 10, "@"),
    ("ORD.DATE", 12, "dd/mm/yyyy"),
    ("PARTY NAME", 26, "@"),
    ("CITY", 20, "@"),
    ("ORD.QTY", 10, "0.00"),
    ("REC.QTY", 10, "0.00"),
    ("BAL.QTY", 10, "0.00"),
    ("DES.DATE", 12, "dd/mm/yyyy"),
]


def build_item_wise_workbook(groups, company_name, run_date):
    """Item Wise Pending P.O. report. Matches the VFP layout: an item-heading row
    ("ITEM NAME,ITEM CODE" in column 1) followed by that item's PO lines."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial Narrow"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Item Wise Pending PO"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=10, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"ITEM WISE PENDING PURCHASE ORDER AS ON {run_date.strftime('%d/%m/%Y')}",
    ).font = Font(name=font_name, size=10, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(ITEM_WISE_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for group in groups:
        heading_cell = ws.cell(row=row, column=1, value=f"{group['part_name']},{group['part_code']}")
        heading_cell.font = Font(name=font_name, size=10, bold=True)
        row += 1
        for line in group["lines"]:
            values = [
                None,  # item name column stays blank on detail rows (the heading row carries it)
                line["ord_no"], _date(line["ord_date"]), line["party_name"], line["city"],
                float(line["ord_qty"]), float(line["rec_qty"]), float(line["bal_qty"]),
                _date(line["despatch_date"]),
            ]
            for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, ITEM_WISE_COLUMNS), start=1):
                cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
                cell.font = Font(name=font_name, size=10)
                cell.number_format = fmt
            row += 1

    for r in range(header_row, row):
        for c in range(1, len(ITEM_WISE_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


MONTHLY_ORDER_REPORT_COLUMNS = [
    ("ITEM NAME", 30, "@"),
    ("ORD.NO", 10, "@"),
    ("ORD.DATE", 12, "dd/mm/yyyy"),
    ("PARTY NAME", 26, "@"),
    ("CITY", 20, "@"),
    ("ORD.QTY", 10, "0.00"),
    ("REC.QTY", 10, "0.00"),
    ("BAL.QTY", 10, "0.00"),
    ("DES.DATE", 12, "dd/mm/yyyy"),
]


def build_monthly_order_report_workbook(rows, company_name, start_date, end_date):
    """Monthly Order Report. Flat rows (not grouped) -- every PO line placed in the date
    range, same column layout as Item Wise Pending P.O. but without the item-heading rows,
    since this report shows all lines (pending or not) rather than grouping by item."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Order Report"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=10, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"ORDER REPORT AS ON {start_date.strftime('%d-%m-%Y')} TO {end_date.strftime('%d-%m-%Y')}",
    ).font = Font(name=font_name, size=10, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(MONTHLY_ORDER_REPORT_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for line in rows:
        values = [
            line["item_name"], line["ord_no"], _date(line["ord_date"]), line["party_name"], line["city"],
            float(line["ord_qty"]), float(line["rec_qty"]), float(line["bal_qty"]), _date(line["despatch_date"]),
        ]
        for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, MONTHLY_ORDER_REPORT_COLUMNS), start=1):
            cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
            cell.font = Font(name=font_name, size=10)
            cell.number_format = fmt
        row += 1

    for r in range(header_row, row):
        for c in range(1, len(MONTHLY_ORDER_REPORT_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


MONTHLY_PURCHASE_COLUMNS = [
    ("CODE", 6, "@"),
    ("PARTY NAME", 25, "@"),
    ("BILL NO.", 8, "@"),
    ("BILL DATE", 10, "dd/mm/yyyy"),
    ("MRN NO.", 8, "@"),
    ("MRN DATE", 10, "dd/mm/yyyy"),
    ("GROSS AMT.", 12, "0.00"),
    ("NET AMT.", 12, "0.00"),
]


def build_monthly_purchase_workbook(rows, company_name, start_date, end_date):
    """Monthly Purchase Report. Matches the old VFP "PURCHASE REPORT" Excel layout."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial Narrow"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Report"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=10, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"PURCHASE REPORT FROM DATE {start_date.strftime('%d/%m/%Y')} TO {end_date.strftime('%d/%m/%Y')}",
    ).font = Font(name=font_name, size=10, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(MONTHLY_PURCHASE_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    total_gross = 0.0
    total_net = 0.0
    for r in rows:
        gross = float(r["gross_amount"])
        net = float(r["net_amount"])
        total_gross += gross
        total_net += net
        values = [
            r["code"], r["party_name"], r["bill_no"], _date(r["bill_date"]),
            r["mrn_no"], _date(r["mrn_date"]), gross, net,
        ]
        for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, MONTHLY_PURCHASE_COLUMNS), start=1):
            cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
            cell.font = Font(name=font_name, size=10)
            cell.number_format = fmt
        row += 1

    row += 1
    ws.cell(row=row, column=6, value="Total").font = Font(name=font_name, size=10, bold=True)
    gross_cell = ws.cell(row=row, column=7, value=total_gross)
    gross_cell.font = Font(name=font_name, size=10, bold=True)
    gross_cell.number_format = "0.00"
    net_cell = ws.cell(row=row, column=8, value=total_net)
    net_cell.font = Font(name=font_name, size=10, bold=True)
    net_cell.number_format = "0.00"

    for r in range(header_row, row - 1):
        for c in range(1, len(MONTHLY_PURCHASE_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ITEM_WISE_RATE_COLUMNS = [
    ("CODE", 20, "@"),
    ("PARTY NAME", 25, "@"),
    ("GRADE", 5, "@"),
    ("BILL NO", 6, "@"),
    ("DATE", 10, "dd/mm/yyyy"),
    ("QTY", 8, "0.00"),
    ("PO QTY", 8, "0.00"),
    ("RATE", 10, "0.00"),
]


def build_item_wise_rate_workbook(groups, company_name, start_date, end_date):
    """Item-Wise Rate List. Matches the VFP layout: an item-heading row
    ("ITEM NAME"+ITEM CODE in column 1) followed by that item's transaction rows."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial Narrow"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Item Wise Rate List"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=10, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"ITEM-WISE RATE LIST FROM DATE {start_date.strftime('%d/%m/%Y')} TO {end_date.strftime('%d/%m/%Y')}",
    ).font = Font(name=font_name, size=10, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(ITEM_WISE_RATE_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for group in groups:
        heading_cell = ws.cell(row=row, column=1, value=f"{group['item_name']}{group['item_code']}")
        heading_cell.font = Font(name=font_name, size=10, bold=True)
        row += 1
        for line in group["lines"]:
            values = [
                line["code"], line["party_name"], line["grade"], line["bill_no"], _date(line["bill_date"]),
                float(line["qty"]), float(line["po_qty"]), float(line["rate"]),
            ]
            for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, ITEM_WISE_RATE_COLUMNS), start=1):
                cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
                cell.font = Font(name=font_name, size=10)
                cell.number_format = fmt
            row += 1

    for r in range(header_row, row):
        for c in range(1, len(ITEM_WISE_RATE_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


PARTY_WISE_RATE_COLUMNS = [
    ("PARTY NAME", 30, "@"),
    ("GRADE", 5, "@"),
    ("CODE", 10, "@"),
    ("ITEM NAME", 20, "@"),
    ("BILL NO", 8, "@"),
    ("DATE", 10, "dd/mm/yyyy"),
    ("QTY", 8, "0.00"),
    ("PO QTY", 8, "0.00"),
    ("RATE", 10, "0.00"),
]


def build_party_wise_rate_workbook(groups, company_name, start_date, end_date):
    """Party-Wise Rate List. Matches the VFP layout: a party-heading row (name+city in
    column 1, grade in column 2) followed by that party's transaction rows."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Party Wise Rate List"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=10, bold=True)
    ws.cell(
        row=2, column=1,
        value=f"PARTY-WISE RATE LIST FROM DATE {start_date.strftime('%d/%m/%Y')} TO {end_date.strftime('%d/%m/%Y')}",
    ).font = Font(name=font_name, size=10, bold=True)

    header_row = 3
    for col, (heading, width, _fmt) in enumerate(PARTY_WISE_RATE_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for group in groups:
        # Old VFP: ALLTRIM(cname)+ALLTRIM(cadd3) -- name and city run together, no separator.
        ws.cell(row=row, column=1, value=f"{group['party_name']}{group['city']}").font = Font(name=font_name, size=10, bold=True)
        ws.cell(row=row, column=2, value=group["grade"]).font = Font(name=font_name, size=10, bold=True)
        row += 1
        for line in group["lines"]:
            values = [
                None, None,  # party name/grade columns stay blank on detail rows
                line["item_code"], line["item_name"], line["bill_no"], _date(line["bill_date"]),
                float(line["qty"]), float(line["po_qty"]), float(line["rate"]),
            ]
            for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, PARTY_WISE_RATE_COLUMNS), start=1):
                cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
                cell.font = Font(name=font_name, size=10)
                cell.number_format = fmt
            row += 1

    for r in range(header_row, row):
        for c in range(1, len(PARTY_WISE_RATE_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


PENDING_LETTER_COLUMNS = [
    ("ORD.NO.", 12, "@"),
    ("DATE", 12, "dd/mm/yyyy"),
    ("CODE", 14, "@"),
    ("ITEM NAME", 30, "@"),
    ("QTY", 8, "0.00"),
    ("DES.DT", 12, "dd/mm/yyyy"),
    ("D.DAYS", 8, "0"),
]


def build_pending_order_letter_workbook(party, lines, company_name, run_date):
    """Party Pending Order Letter. Matches the VFP "PENDING ORDER" business-letter layout,
    not a plain tabular report."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial Narrow"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pending Order Letter"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    def write(row, value, bold=False):
        cell = ws.cell(row=row, column=1, value=sanitize(value) if value not in ("", None) else None)
        cell.font = Font(name=font_name, size=10, bold=bold)
        return cell

    row = 1
    write(row, "PENDING ORDER", bold=True)
    row += 1
    write(row, f"{party['name']}                    DATE {run_date.strftime('%d/%m/%Y')}", bold=True)
    row += 1
    for addr in (party["address_1"], party["address_2"], party["address_3"]):
        if addr:
            write(row, addr)
            row += 1
    row += 1
    write(row, "Dear Sir")
    row += 1
    write(row, "Please refer to the following order")
    row += 2

    header_row = row
    for col, (heading, width, _fmt) in enumerate(PENDING_LETTER_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    row += 1

    for line in lines:
        values = [
            line["ord_no"], _date(line["ord_date"]), line["item_code"], line["item_name"],
            float(line["pending_qty"]), _date(line["despatch_date"]), line["delay_days"],
        ]
        for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, PENDING_LETTER_COLUMNS), start=1):
            cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
            cell.font = Font(name=font_name, size=10)
            cell.number_format = fmt
        row += 1

    for r in range(header_row, row):
        for c in range(1, len(PENDING_LETTER_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    row += 1
    write(row, "Kindly expedite the supplies which have already been delayed and intimate the despatch")
    row += 1
    write(row, "details immediately.")
    row += 2
    write(row, "Thanking you")
    row += 2
    write(row, f"For {company_name}")
    row += 2
    write(row, "Authorised Signatory")
    row += 1
    write(row, "Purchase Deptt.")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ITEMWISE_PARTY_COLUMNS = [
    ("CODE", 10, "@"),
    ("VENDOR NAME", 30, "@"),
    ("CITY", 18, "@"),
    ("GRADE", 10, "@"),
    ("RATE", 10, "0.00"),
]


def build_itemwise_party_list_workbook(part, rows, company_name):
    """Itemwise Party List. Matches the VFP layout: title + "<Item Name>, <Item Code>"
    heading, then the CODE/VENDOR NAME/CITY/RATE table."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial Narrow"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Itemwise Party List"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=10, bold=True)
    ws.cell(row=2, column=1, value="ITEMWISE PARTY LIST").font = Font(name=font_name, size=10, bold=True)
    ws.cell(row=3, column=1, value=f"{part['name']}, {part['code']}").font = Font(name=font_name, size=10, bold=True)

    header_row = 4
    for col, (heading, width, _fmt) in enumerate(ITEMWISE_PARTY_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for r in rows:
        values = [r["party_code"], r["party_name"], r["city"], r["grade"], float(r["rate"])]
        for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, ITEMWISE_PARTY_COLUMNS), start=1):
            cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
            cell.font = Font(name=font_name, size=10)
            cell.number_format = fmt
        row += 1

    for r in range(header_row, row):
        for c in range(1, len(ITEMWISE_PARTY_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


PARTY_WISE_ITEM_COLUMNS = [
    ("CODE", 12, "@"),
    ("ITEM NAME", 30, "@"),
    ("RATE", 10, "0.00"),
]


def build_party_wise_item_list_workbook(party, rows, company_name):
    """Party Wise Item List. Matches the VFP layout: company name, "PARTY WISE ITEM
    LIST" title, "<Party Name> GRADE <Grade>" heading, then CODE/ITEM NAME/RATE."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Border, Font, Side

    font_name = "Arial Narrow"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Party Wise Item List"

    def sanitize(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    ws.cell(row=1, column=1, value=company_name).font = Font(name=font_name, size=10, bold=True)
    ws.cell(row=2, column=1, value="PARTY WISE ITEM LIST").font = Font(name=font_name, size=10, bold=True)
    ws.cell(row=3, column=1, value=f"{party['name']} GRADE {party['grade']}").font = Font(name=font_name, size=10, bold=True)

    header_row = 5
    for col, (heading, width, _fmt) in enumerate(PARTY_WISE_ITEM_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=heading)
        cell.font = Font(name=font_name, size=10, bold=True)
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for r in rows:
        values = [r["code"], r["name"], float(r["rate"])]
        for col, (value, (_heading, _width, fmt)) in enumerate(zip(values, PARTY_WISE_ITEM_COLUMNS), start=1):
            cell = ws.cell(row=row, column=col, value=sanitize(value) if value not in ("", None) else None)
            cell.font = Font(name=font_name, size=10)
            cell.number_format = fmt
        row += 1

    for r in range(header_row, row):
        for c in range(1, len(PARTY_WISE_ITEM_COLUMNS) + 1):
            ws.cell(row=r, column=c).border = border

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
