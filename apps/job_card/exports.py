def build_job_card_listing_workbook(data, company_name, start_date, end_date):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([company_name])
    ws.append([f"JOB CARD LIST FROM {start_date:%d/%m/%Y} TO {end_date:%d/%m/%Y}"])
    ws.append([])
    headers = ["SNO", "JOB CARD", "JOB DATE", "CODE", "ITEM NAME", "Unit", "QTY", "PROD", "TIME REQD.",
               "TIME TAKEN", "CL", "CLOSING DATE", "SUPERVISOR NAME", "RATE", "AMOUNT"]
    ws.append(headers)
    header_row = ws.max_row
    widths = [5, 10, 12, 10, 28, 6, 10, 10, 10, 10, 5, 13, 22, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = w
    for c in ws[header_row]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for n, r in enumerate(data["rows"], start=1):
        ws.append([n, r["job_no"], r["date"], r["code"], r["item_name"], r["unit"],
                   float(r["qty"]), float(r["prod"]), float(r["time_reqd"]), float(r["time_taken"]),
                   r["closed"], r["close_date"], r["supervisor"], float(r["rate"]), float(r["amount"])])
    ws.append([None, None, None, None, None, "Total", None, None, None, None, None, None, None, None, float(data["total"])])
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        row[6].number_format = row[7].number_format = row[8].number_format = row[9].number_format = "0.00"
        row[13].number_format = row[14].number_format = "0.00"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_job_card_supervisor_workbook(data, company_name, start_date, end_date, supervisor_name):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([company_name])
    ws.append([f"JOB CARD LIST FROM {start_date:%d/%m/%Y} TO {end_date:%d/%m/%Y} ({supervisor_name})"])
    ws.append([])
    headers = ["SNO", "JOB CARD", "JOB DATE", "CODE", "ITEM NAME", "Unit", "QTY", "PROD", "WASTAGE",
               "CL", "CLOSING DATE", "SUPERVISOR NAME", "RATE", "AMOUNT"]
    ws.append(headers)
    header_row = ws.max_row
    widths = [5, 10, 12, 10, 28, 6, 10, 10, 10, 5, 13, 22, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = w
    for c in ws[header_row]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for n, r in enumerate(data["rows"], start=1):
        ws.append([n, r["job_no"], r["date"], r["code"], r["item_name"], r["unit"],
                   float(r["qty"]), float(r["prod"]), float(r["wastage"]), r["closed"], r["close_date"],
                   r["supervisor"], float(r["rate"]), float(r["amount"])])
    ws.append([None, None, None, None, None, "Total", None, None, None, None, None, None, None, float(data["total"])])
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for idx in (6, 7, 8, 12, 13):
            row[idx].number_format = "0.00"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_short_list_workbook(data, company_name):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([company_name])
    ws.append([f"SHORT-LIST FOR {data['item_name']} {data['item_code']}"])
    ws.append([f"QUANTITY {data['qty']}"])
    ws.append([])
    ws.append(["CODE", "ITEM NAME", "Unit", "INV.RATE", "STOCK", "REQD. QTY", "SHORTAGE"])
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([10, 30, 6, 10, 10, 12, 10], "ABCDEFG"):
        ws.column_dimensions[col].width = w
    for r in data["rows"]:
        ws.append([r["code"], r["name"], r["unit"], float(r["rate"]) if r["rate"] else None,
                   float(r["stock"]), float(r["reqd"]), float(r["shortage"]) if r["shortage"] else None])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
