"""Pending Job Card list -- legacy PENDING JOB CARD LIST: open, non-sub-assembly jobs in a date range."""
from datetime import date
from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from apps.items.models import Item
from apps.mrn.services import _clean

from .models import JobCard
from .views import _listing_range


def pending_jobs(start, end):
    items = {c: (n, u) for c, n, u in Item.objects.values_list("item_code", "item_name_1", "unit")}
    rows = []
    for j in JobCard.objects.filter(job_date__gte=start, job_date__lte=end).order_by("job_date", "job_no"):
        if _clean(j.sub_assembly) or _clean(j.close_flag) == "Y":
            continue
        code = _clean(j.item_code) or ""
        name, unit = items.get(code[:4], (None, None))
        rows.append({
            "div": _clean(j.div) or "", "job_no": _clean(j.job_no),
            "date": j.job_date.isoformat() if j.job_date else "",
            "item_code": code, "item_name": (_clean(name) or "")[:30], "unit": _clean(unit) or "",
            "qty": f"{(j.qty or 0):.2f}",
        })
    return rows


class PendingJobListView(LoginRequiredMixin, TemplateView):
    template_name = "job_card/pending_job_list.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.replace(day=1).isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def pending_job_data(request):
    start, end = _listing_range(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse({"rows": pending_jobs(start, end)})


@login_required
@require_GET
def pending_job_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _listing_range(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"PENDING JOB CARD LIST FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["DIV.", "JOB CARD", "JOB DATE", "CODE", "ITEM NAME", "UNIT", "QTY"])
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 10, 12, 10, 30, 6, 10], "ABCDEFG"):
        ws.column_dimensions[col].width = w
    for r in pending_jobs(start, end):
        ws.append([r["div"], r["job_no"], r["date"], r["item_code"], r["item_name"], r["unit"], float(r["qty"])])
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        row[6].number_format = "0.00"
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Pending_Job_List.xlsx"'
    return response
