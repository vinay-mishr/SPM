"""Reports > Sales Reports > Representative Wise Sale (VFP "REPRESENTATIVEWISE
SALE" form, OLE-Excel export block).

Direct port of the pasted VFP script's SECOND (OLE-Excel) block -- the
authoritative one, not the older PRN/``rwsale.PRN`` print routine above it
in the same paste (kept only for its filter logic, since the two agree on
everything but the rep code's source and the GST-only filter, both fixed
in the newer block). Differences the newer block introduces, preserved
here exactly:
  * The rep for a bill is ``bill.rep_cod`` (read straight off the bill
    itself), not ``cust_mst.crep`` (the party's own default rep) like the
    PRN version used.
  * Only GST bills count (``bill_type = 'G'``) -- the PRN version had no
    such filter.

Rows are grouped REP then STATE then by bill date (the VFP source's own
``rep+state+dtoc(b_date,1)`` index), each level printing a header/total
row: rep name, then within it each state name, then within THAT each
bill, then a state "TOTAL", rolling up into a rep "REP. TOTAL", rolling up
into one final "GRAND TOTAL" -- all preserved exactly, including that a
bill with no matching ``cust_mst`` row (so no state) still appears, just
grouped under a blank state.

ALL/SINGLE (``thisform.tsa.value``) scopes to one rep via
``thisform.tpcode.value`` -- ported as the ``rep_code`` parameter.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import connections
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

BILL_TYPE = "G"


def _rows(sql, params=None):
    with connections["legacy"].cursor() as cur:
        cur.execute(sql, params or [])
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]


def _clean(value):
    return (value or "").strip() if isinstance(value, str) or value is None else value


def _d(value):
    try:
        return Decimal(str(value if value is not None else 0))
    except Exception:
        return Decimal("0")


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


def rep_wise_sale_groups(start, end, rep_code=None):
    sql = (
        "SELECT b.div, b.bill_no, b.b_date, b.amount, b.rep_cod, c.cstate, c.cname, c.cadd3 "
        "FROM bill b LEFT JOIN cust_mst c ON c.ccode = b.p_code "
        "WHERE b.bill_type = %s AND b.b_date >= %s AND b.b_date <= %s"
    )
    params = [BILL_TYPE, start, end]
    if rep_code:
        sql += " AND b.rep_cod = %s"
        params.append(rep_code)
    sql += " ORDER BY b.rep_cod, c.cstate, b.b_date"
    bills = _rows(sql, params)

    rep_names = {_clean(r["rcode"]): _clean(r["rname"]) for r in _rows("SELECT rcode, rname FROM rep_mst")}
    state_names = {_clean(r["stcode"]): _clean(r["stname"]) for r in _rows("SELECT stcode, stname FROM state")}

    reps = []
    grand_total = Decimal("0")
    current_rep = None
    for b in bills:
        rep_code_val = _clean(b["rep_cod"]) or ""
        state_code_val = _clean(b["cstate"]) or ""
        if current_rep is None or current_rep["code"] != rep_code_val:
            current_rep = {
                "code": rep_code_val, "name": rep_names.get(rep_code_val, "") or rep_code_val,
                "states": [], "total": Decimal("0"),
            }
            reps.append(current_rep)
            current_state = None
        if current_state is None or current_state["code"] != state_code_val:
            current_state = {
                "code": state_code_val, "name": state_names.get(state_code_val, "") or state_code_val,
                "bills": [], "total": Decimal("0"),
            }
            current_rep["states"].append(current_state)

        amount = _d(b["amount"])
        current_state["bills"].append({
            "invoice_no": (_clean(b["div"]) or "") + (_clean(b["bill_no"]) or ""),
            "invoice_date": b["b_date"].isoformat() if b["b_date"] else "",
            "party_name": _clean(b["cname"]) or "", "party_city": _clean(b["cadd3"]) or "",
            "amount": f"{amount:.2f}",
        })
        current_state["total"] += amount
        current_rep["total"] += amount
        grand_total += amount

    for r in reps:
        r["total"] = f"{r['total']:.2f}"
        for s in r["states"]:
            s["total"] = f"{s['total']:.2f}"
    return reps, f"{grand_total:.2f}"


class RepWiseSaleView(LoginRequiredMixin, TemplateView):
    template_name = "reps/sale_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": _fy_start(today).isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), _fy_start(today))
    end = parse(request.GET.get("end_date"), today)
    scope = (request.GET.get("scope") or "all").strip().lower()
    rep_code = (request.GET.get("rep_code") or "").strip().upper() if scope == "single" else None
    return start, end, scope, rep_code


@login_required
@require_GET
def rws_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end, scope, rep_code = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    if scope == "single" and not rep_code:
        return JsonResponse({"message": "Select a representative first."}, status=400)

    reps, grand_total = rep_wise_sale_groups(start, end, rep_code)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"REPRESENTATIVE WISE SALE FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])

    header = ["INV.NO", "INV.DATE", "PARTY NAME", "PARTY CITY", "SALE AMT"]
    widths = [6, 10, 30, 20, 12]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip(widths, "ABCDE"):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["B"].number_format = "dd/mm/yy"
    ws.column_dimensions["E"].number_format = "0.00"

    for r in reps:
        ws.append([r["name"]])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for s in r["states"]:
            ws.append([s["name"]])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            for bi in s["bills"]:
                ws.append([
                    bi["invoice_no"], date.fromisoformat(bi["invoice_date"]) if bi["invoice_date"] else None,
                    bi["party_name"], bi["party_city"], float(bi["amount"]),
                ])
                ws.cell(row=ws.max_row, column=2).number_format = "dd/mm/yy"
                ws.cell(row=ws.max_row, column=5).number_format = "0.00"
            ws.append(["", "", "", "TOTAL", float(s["total"])])
            ws.cell(row=ws.max_row, column=4).font = Font(bold=True)
        ws.append(["", "", "", "REP. TOTAL", float(r["total"])])
        ws.cell(row=ws.max_row, column=4).font = Font(bold=True)

    ws.append(["", "", "", "GRAND TOTAL", float(grand_total)])
    ws.cell(row=ws.max_row, column=4).font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Representative_Wise_Sale_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
