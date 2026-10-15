import json
from datetime import date, datetime

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import JsonResponse
from django.views.decorators.http import require_GET, require_POST
from django.views.generic import TemplateView

from . import services


def _can_edit(user):
    return settings.LEGACY_WRITES_ENABLED


def _parse(value, default):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return default


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


class SaleOrderView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/sale_order.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["can_edit"] = _can_edit(self.request.user)
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def so_list(request):
    start = _parse(request.GET.get("start_date"), date.today())
    end = _parse(request.GET.get("end_date"), date.today())
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse({"rows": services.list_orders(start, end, request.GET.get("search", ""))})


@login_required
@require_GET
def so_next(request):
    return JsonResponse({"order_no": str(services.next_no())})


@login_required
@require_GET
def so_detail(request, order_no):
    data = services.load(order_no)
    if not data:
        return JsonResponse({"message": "Order not found."}, status=404)
    return JsonResponse(data)


@login_required
@require_POST
def so_save(request):
    if not _can_edit(request.user):
        return JsonResponse({"message": "You don't have permission to change sale orders."}, status=403)
    try:
        body = json.loads(request.body or b"{}")
    except ValueError:
        return JsonResponse({"message": "Invalid request."}, status=400)
    mode = body.get("mode")
    if mode not in ("add", "modify", "delete"):
        return JsonResponse({"message": "Unknown action."}, status=400)
    try:
        result = services.save(body, mode, username=request.user.get_username())
    except services.SaleOrderError as exc:
        return JsonResponse({"message": str(exc), "errors": exc.errors}, status=400)
    verb = {"add": "saved", "modify": "modified", "delete": "deleted"}[mode]
    return JsonResponse({"message": f"Sale order {result['order_no']} {verb}.", **result})


@login_required
@require_GET
def so_item(request):
    data = services.item_detail(request.GET.get("code", ""))
    if not data:
        return JsonResponse({"message": "Item not found."}, status=404)
    return JsonResponse(data)


@login_required
@require_GET
def so_party(request):
    data = services.party_detail(request.GET.get("code", ""))
    if not data:
        return JsonResponse({"message": "Party not found."}, status=404)
    return JsonResponse(data)


class SaleOrderPrintView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/print.html"

    def get_context_data(self, **kwargs):
        from apps.creditor_outstanding.remittance import amount_in_words
        from django.conf import settings
        from django.http import Http404

        context = super().get_context_data(**kwargs)
        data = services.load(kwargs["order_no"])
        if not data:
            raise Http404("Order not found.")
        context["order"] = data
        context["brand_name"] = settings.BRAND["name"].upper()
        context["in_words"] = amount_in_words(data["net_amount"])
        context["fill_rows"] = range(max(0, 20 - len(data["lines"])))
        context["terms"] = []
        return context


class SaleOrderGraphicalView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/graphical.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["today"] = date.today().strftime("%d/%m/%Y")
        return context


class SaleReportView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/sale_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def sale_report_data(request):
    start = _parse(request.GET.get("start_date"), date.today())
    end = _parse(request.GET.get("end_date"), date.today())
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse(services.sale_report(start, end))


@login_required
@require_GET
def sale_report_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start = _parse(request.GET.get("start_date"), date.today())
    end = _parse(request.GET.get("end_date"), date.today())
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    data = services.sale_report(start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"SALE ORDER REPORT FROM DATE {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["ORDER NO", "DATE", "PARTY NAME", "PARTY CITY", "ORDER VALUE"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([10, 12, 36, 20, 16], "ABCDE"):
        ws.column_dimensions[col].width = w
    for r in data["rows"]:
        ws.append([r["order_no"], r["date"], r["party_name"], r["city"], float(r["net_amount"])])
    ws.append(["", "", "", "Total", float(data["total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Sale_Order_Report_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response


class PartyPendingView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/party_pending.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"as_on": today.isoformat()}
        return context


def _pending_params(request):
    as_on = _parse(request.GET.get("as_on"), date.today())
    party = (request.GET.get("party_code") or "").strip().upper() or None
    return as_on, party


@login_required
@require_GET
def party_pending_data(request):
    as_on, party = _pending_params(request)
    return JsonResponse(services.party_pending(as_on, party))


@login_required
@require_GET
def party_pending_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    as_on, party = _pending_params(request)
    data = services.party_pending(as_on, party)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"PARTY WISE PENDING ORDER AS ON {as_on:%d/%m/%Y}"])
    ws.append([])
    ws.append(["ORD.NO", "ORD.DATE", "ITEM CODE", "ITEM NAME", "PENDING QTY", "CUR.STK", "VALUE"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([10, 12, 14, 36, 14, 12, 14], "ABCDEFG"):
        ws.column_dimensions[col].width = w
    for g in data["groups"]:
        ws.append([f"{g['party_name']}, {g['city']}, {g['party_code']}"])
        for l in g["lines"]:
            ws.append([l["order_no"], l["date"], l["item_code"], l["item_name"], float(l["pending"]),
                       float(l["current"]), float(l["value"])])
        ws.append(["", "", "", "TOTAL", None, None, float(g["total"])])
    ws.append(["", "", "", "GRAND TOTAL", None, None, float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Party_Wise_Pending_Order_{as_on:%d-%m-%Y}.xlsx"'
    return response


class ItemPendingView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/item_pending.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"as_on": date.today().isoformat()}
        return context


@login_required
@require_GET
def item_pending_data(request):
    as_on, _ = _pending_params(request)
    item = (request.GET.get("item_code") or "").strip().upper() or None
    return JsonResponse(services.item_pending(as_on, item))


@login_required
@require_GET
def item_pending_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    as_on, _ = _pending_params(request)
    item = (request.GET.get("item_code") or "").strip().upper() or None
    data = services.item_pending(as_on, item)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"ITEM WISE PENDING ORDER AS ON {as_on:%d/%m/%Y}"])
    ws.append([])
    ws.append(["ORD.NO", "ORD.DATE", "PARTY NAME", "CITY", "PENDING QTY", "CUR.STOCK"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([10, 12, 36, 20, 14, 12], "ABCDEF"):
        ws.column_dimensions[col].width = w
    for g in data["groups"]:
        ws.append([f"{g['item_name']}, {g['item_code']}"])
        for l in g["lines"]:
            ws.append([l["order_no"], l["date"], l["party_name"], l["city"], float(l["pending"])])
        ws.append(["", "", "", "TOTAL", float(g["total"]), float(g["current"])])
    ws.append(["", "", "", "GRAND TOTAL", float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Item_Wise_Pending_Order_{as_on:%d-%m-%Y}.xlsx"'
    return response


class AreaPendingView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/area_pending.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"as_on": date.today().isoformat()}
        return context


@login_required
@require_GET
def area_list(request):
    return JsonResponse({"areas": services.areas(request.GET.get("q", ""))})


@login_required
@require_GET
def area_pending_data(request):
    as_on, _ = _pending_params(request)
    area = (request.GET.get("area_code") or "").strip() or None
    return JsonResponse(services.area_pending(as_on, area))


@login_required
@require_GET
def area_pending_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    as_on, _ = _pending_params(request)
    area = (request.GET.get("area_code") or "").strip() or None
    data = services.area_pending(as_on, area)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"AREA WISE/PARTYWISE PENDING ORDER AS ON {as_on:%d/%m/%Y}"])
    ws.append([])
    ws.append(["ORD.NO", "ORD.DATE", "ITEM CODE", "ITEM NAME", "QTY"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([10, 12, 14, 36, 12], "ABCDE"):
        ws.column_dimensions[col].width = w
    for g in data["groups"]:
        ws.append([f"AREA {g['area_code']} {g['area_name']}"])
        for p in g["parties"]:
            ws.append([f"{p['party_name']}, {p['city']}, {p['party_code']}"])
            for l in p["lines"]:
                ws.append([l["order_no"], l["date"], l["item_code"], l["item_name"], float(l["pending"])])
        ws.append(["", "", "", "AREA TOTAL", float(g["total"])])
    ws.append(["", "", "", "GRAND TOTAL", float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Area_Wise_Pending_Order_{as_on:%d-%m-%Y}.xlsx"'
    return response


class RepOrdersView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/rep_orders.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def rep_list(request):
    return JsonResponse({"reps": services.reps(request.GET.get("q", ""))})


def _rep_params(request):
    start = _parse(request.GET.get("start_date"), date.today())
    end = _parse(request.GET.get("end_date"), date.today())
    rep = (request.GET.get("rep_code") or "").strip() or None
    return start, end, rep


@login_required
@require_GET
def rep_orders_data(request):
    start, end, rep = _rep_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse(services.rep_orders(start, end, rep))


@login_required
@require_GET
def rep_orders_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end, rep = _rep_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    data = services.rep_orders(start, end, rep)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"REP.WISE ORDER VALUE FROM DATE {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["PARTY NAME", "CITY", "ORD.NO", "ORD.DATE", "AMOUNT"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([36, 20, 12, 12, 14], "ABCDE"):
        ws.column_dimensions[col].width = w
    for g in data["groups"]:
        ws.append([g["rep_name"]])
        for o in g["orders"]:
            ws.append([o["party_name"], o["city"], o["div"] + o["order_no"], o["date"], float(o["amount"])])
        ws.append(["", "", "", "TOTAL", float(g["total"])])
    ws.append(["", "", "", "GRAND TOTAL", float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Rep_Wise_Order_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
class RepPendingView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/rep_pending.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def rep_list(request):
    return JsonResponse({"reps": services.reps(request.GET.get("q", ""))})


def _rep_params(request):
    start = _parse(request.GET.get("start_date"), date.today())
    end = _parse(request.GET.get("end_date"), date.today())
    rep = (request.GET.get("rep_code") or "").strip() or None
    return start, end, rep


@login_required
@require_GET
def rep_pending_data(request):
    start, end, rep = _rep_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse(services.rep_orders(start, end, rep, pending=True))


@login_required
@require_GET
def rep_pending_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end, rep = _rep_params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    data = services.rep_orders(start, end, rep, pending=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"REP.WISE PENDING ORDER VALUE FROM DATE {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["PARTY NAME", "CITY", "ORD.NO", "ORD.DATE", "AMOUNT"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([36, 20, 12, 12, 14], "ABCDE"):
        ws.column_dimensions[col].width = w
    for g in data["groups"]:
        ws.append([g["rep_name"]])
        for o in g["orders"]:
            ws.append([o["party_name"], o["city"], o["div"] + o["order_no"], o["date"], float(o["amount"])])
        ws.append(["", "", "", "TOTAL", float(g["total"])])
    ws.append(["", "", "", "GRAND TOTAL", float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Rep_Wise_Pending_Order_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response


class DespatchStatusView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/despatch_status.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"as_on": date.today().isoformat()}
        return context


@login_required
@require_GET
def despatch_status_data(request):
    as_on, _ = _pending_params(request)
    return JsonResponse(services.despatch_status(as_on))


@login_required
@require_GET
def despatch_status_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    as_on, _ = _pending_params(request)
    data = services.despatch_status(as_on)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"PENDING SALES ORDER STATUS REPORT REQUIRING DESPATCHES UP TO {as_on:%d/%m/%Y}"])
    ws.append([])
    ws.append(["SNO", "PARTY", "OA NO", "OA DATE", "DESPATCH DATE", "ORDER VALUE"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([6, 36, 12, 12, 14, 14], "ABCDEF"):
        ws.column_dimensions[col].width = w
    for n, o in enumerate(data["orders"], start=1):
        ws.append([n, o["party_name"], o["order_no"], o["order_date"], o["despatch_date"], float(o["amount"])])
    ws.append(["", "", "", "", "TOTAL", float(data["total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Pending_Despatch_Status_{as_on:%d-%m-%Y}.xlsx"'
    return response


class OrderAdjustmentView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/order_adjustment.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["can_edit"] = _can_edit(self.request.user)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


@login_required
@require_GET
def adjustment_data(request):
    party = (request.GET.get("party_code") or "").strip().upper()
    if not party:
        return JsonResponse({"message": "Select a party first."}, status=400)
    as_on, _ = _pending_params(request)
    return JsonResponse({"rows": services.adjustment_lines(party, as_on)})


@login_required
@require_POST
def adjustment_save(request):
    if not _can_edit(request.user):
        return JsonResponse({"message": "You don't have permission to adjust orders."}, status=403)
    try:
        body = json.loads(request.body or b"{}")
    except ValueError:
        return JsonResponse({"message": "Invalid request."}, status=400)
    party = (body.get("party_code") or "").strip().upper()
    lines = [l for l in body.get("lines") or [] if l.get("adjust")]
    if not party or not lines:
        return JsonResponse({"message": "Tick at least one order line to adjust."}, status=400)
    try:
        count = services.adjust_orders(party, lines)
    except services.SaleOrderError as exc:
        return JsonResponse({"message": str(exc)}, status=400)
    return JsonResponse({"message": f"{count} order line(s) adjusted."})


class AdjustmentReportView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/adjustment_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"as_on": date.today().isoformat()}
        return context


@login_required
@require_GET
def adjustment_report_data(request):
    as_on, party = _pending_params(request)
    return JsonResponse(services.adjustment_report(as_on, party))


@login_required
@require_GET
def adjustment_report_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    as_on, party = _pending_params(request)
    data = services.adjustment_report(as_on, party)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"ORDER ADJUSTMENT REPORT AS ON {as_on:%d/%m/%Y}"])
    ws.append([])
    ws.append(["ORD.NO", "ORD.DATE", "ITEM CODE", "ITEM NAME", "PENDING QTY", "VALUE"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([10, 12, 14, 36, 14, 14], "ABCDEF"):
        ws.column_dimensions[col].width = w
    for g in data["groups"]:
        ws.append([f"{g['party_name']}, {g['city']}, {g['party_code']}"])
        for l in g["lines"]:
            ws.append([l["order_no"], l["date"], l["item_code"], l["item_name"], float(l["pending"]), float(l["value"])])
        ws.append(["", "", "", "TOTAL", None, float(g["total"])])
    ws.append(["", "", "", "GRAND TOTAL", None, float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Order_Adjustment_Report_{as_on:%d-%m-%Y}.xlsx"'
    return response


class SpareOrderView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/spare_order.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["can_edit"] = _can_edit(self.request.user)
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def spare_list(request):
    start = _parse(request.GET.get("start_date"), date.today())
    end = _parse(request.GET.get("end_date"), date.today())
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    return JsonResponse({"rows": services.spare_list(start, end, request.GET.get("search", ""))})


@login_required
@require_GET
def spare_next(request):
    return JsonResponse({"order_no": str(services.spare_next_no())})


@login_required
@require_GET
def spare_detail(request, order_no):
    data = services.spare_load(order_no)
    if not data:
        return JsonResponse({"message": "Order not found."}, status=404)
    return JsonResponse(data)


@login_required
@require_POST
def spare_save_view(request):
    if not _can_edit(request.user):
        return JsonResponse({"message": "You don't have permission to change spare sale orders."}, status=403)
    try:
        body = json.loads(request.body or b"{}")
    except ValueError:
        return JsonResponse({"message": "Invalid request."}, status=400)
    mode = body.get("mode")
    if mode not in ("add", "modify", "delete"):
        return JsonResponse({"message": "Unknown action."}, status=400)
    try:
        result = services.spare_save(body, mode, username=request.user.get_username())
    except services.SaleOrderError as exc:
        return JsonResponse({"message": str(exc), "errors": exc.errors}, status=400)
    verb = {"add": "saved", "modify": "modified", "delete": "deleted"}[mode]
    return JsonResponse({"message": f"Spare sale order {result['order_no']} {verb}.", **result})


@login_required
@require_GET
def spare_item_view(request):
    data = services.spare_item(request.GET.get("code", ""))
    if not data:
        return JsonResponse({"message": "Item not found."}, status=404)
    return JsonResponse(data)


class SparePrintView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/spare_print.html"

    def get_context_data(self, **kwargs):
        from apps.creditor_outstanding.remittance import amount_in_words
        from django.conf import settings
        from django.http import Http404

        context = super().get_context_data(**kwargs)
        data = services.spare_load(kwargs["order_no"])
        if not data:
            raise Http404("Order not found.")
        context["order"] = data
        context["brand_name"] = settings.BRAND["name"].upper()
        context["in_words"] = amount_in_words(data["net_amount"])
        context["fill_rows"] = range(max(0, 20 - len(data["lines"])))
        context["terms"] = []
        return context


class SparePrintingView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/spare_printing.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["next_no"] = services.spare_next_no() - 1
        return context


class SparePartyPendingView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/spare_party_pending.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"as_on": date.today().isoformat()}
        return context


@login_required
@require_GET
def spare_party_pending_data(request):
    party = (request.GET.get("party_code") or "").strip().upper() or None
    return JsonResponse(services.spare_party_pending(party))


@login_required
@require_GET
def spare_party_pending_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    party = (request.GET.get("party_code") or "").strip().upper() or None
    data = services.spare_party_pending(party)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append(["PARTYWISE SPARE PART PENDING ORDER"])
    ws.append([])
    ws.append(["ORDER NO", "ORDER DATE", "PART NO", "ITEM DESCRIPTION", "ORDERED QTY", "ISSUE QTY", "PENDING ORDER"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([10, 12, 14, 36, 12, 12, 14], "ABCDEFG"):
        ws.column_dimensions[col].width = w
    for g in data["groups"]:
        ws.append([f"{g['party_name']}, {g['city']}, {g['party_code']}"])
        for l in g["lines"]:
            ws.append([l["order_no"], l["date"], l["item_code"], l["item_name"], float(l["ordered"]),
                       float(l["issued"]), float(l["pending"])])
        ws.append(["", "", "", "TOTAL", None, None, float(g["total"])])
    ws.append(["", "", "", "GRAND TOTAL", None, None, float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Partywise_Spare_Pending_Order.xlsx"'
    return response


class SpareItemPendingView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/spare_item_pending.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"as_on": date.today().isoformat()}
        return context


@login_required
@require_GET
def spare_item_pending_data(request):
    item = (request.GET.get("item_code") or "").strip().upper() or None
    return JsonResponse(services.spare_item_pending(item))


@login_required
@require_GET
def spare_item_pending_export(request):
    from io import BytesIO
    from django.conf import settings
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    item = (request.GET.get("item_code") or "").strip().upper() or None
    data = services.spare_item_pending(item)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append(["ITEMWISE SPARE PART PENDING ORDER"])
    ws.append([])
    ws.append(["PART NO", "ITEM DESCRIPTION", "PENDING ORDER"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([14, 40, 16], "ABC"):
        ws.column_dimensions[col].width = w
    for r in data["rows"]:
        ws.append([r["item_code"], r["item_name"], float(r["pending"])])
    ws.append(["", "TOTAL", float(data["grand_total"])])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Itemwise_Spare_Pending_Order.xlsx"'
    return response


class SpareAdjustmentView(LoginRequiredMixin, TemplateView):
    template_name = "sale_order/spare_adjustment.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["can_edit"] = _can_edit(self.request.user)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


@login_required
@require_GET
def spare_adjustment_data(request):
    party = (request.GET.get("party_code") or "").strip().upper()
    if not party:
        return JsonResponse({"message": "Select a party first."}, status=400)
    return JsonResponse({"rows": services.spare_adjustment_lines(party)})


@login_required
@require_POST
def spare_adjustment_save(request):
    if not _can_edit(request.user):
        return JsonResponse({"message": "You don't have permission to adjust orders."}, status=403)
    try:
        body = json.loads(request.body or b"{}")
    except ValueError:
        return JsonResponse({"message": "Invalid request."}, status=400)
    party = (body.get("party_code") or "").strip().upper()
    lines = [l for l in body.get("lines") or [] if l.get("adjust")]
    if not party or not lines:
        return JsonResponse({"message": "Tick at least one order line to adjust."}, status=400)
    try:
        count = services.spare_adjust(party, lines)
    except services.SaleOrderError as exc:
        return JsonResponse({"message": str(exc)}, status=400)
    return JsonResponse({"message": f"{count} order line(s) adjusted."})
