"""MRN > Short List New -- parts short of their reorder level (old VFP SHORT LIST).

Mode A lists every short part; mode P lists only short parts with nothing on order.
Pending sale order (VFP pord) is not in the legacy part table, so that column is blank.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from apps.parts.models import Part


def short_list(mode="A"):
    rows = []
    for p in Part.objects.order_by("part_number"):
        rop = p.reorder_qty or Decimal("0")
        stock = p.current_stock or Decimal("0")
        short = rop - stock
        if short <= 0:
            continue
        on_order = p.qty_on_order or Decimal("0")
        if mode == "P" and on_order > 0:
            continue
        rows.append({
            "code": (p.part_number or "").strip(),
            "name": (p.description_1 or "").strip(),
            "unit": (p.unit or "").strip(),
            "reorder": f"{rop:.2f}",
            "current": f"{stock:.0f}",
            "pending_sale": "",
            "short": f"{short:.2f}",
            "on_order": f"{on_order:.2f}" if on_order > 0 else "",
        })
    return rows


class ShortListView(LoginRequiredMixin, TemplateView):
    template_name = "mrn/short_list.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["run_date"] = date.today().strftime("%d-%m-%Y")
        context["vfp"] = _is_vfp(self.request)
        context["brand_address"] = settings.BRAND["address_line2"]
        return context


def _is_vfp(request):
    return request.GET.get("layout") == "vfp"


def _mode(request):
    if _is_vfp(request):
        return "A"
    return "P" if (request.GET.get("mode") or "A").upper() == "P" else "A"


@login_required
@require_GET
def short_list_data(request):
    return JsonResponse({"rows": short_list(_mode(request))})


@login_required
@require_GET
def short_list_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    vfp = _is_vfp(request)
    if vfp:
        ws.append([settings.BRAND["name"].upper()])
        ws.append([settings.BRAND["address_line2"]])
        ws.append(["SHORT LIST"])
        ws.append(["CODE", "ITEM NAME", "UNIT", "CUR.STK.", "REORDER LEVEL", "SHORT QTY", "QTY.ON ORDER"])
        widths = [12, 36, 7, 12, 14, 12, 14]
    else:
        ws.append([settings.BRAND["name"].upper()])
        ws.append([f"SHORT LIST AS ON {date.today():%d/%m/%Y}"])
        ws.append([])
        ws.append(["CODE", "PRODUCT NAME", "UNIT", "RE-ORD QTY", "CURRENT STOCK", "PENDING SALE ORD",
                   "SHORT QTY", "QTY ON ORDERED"])
        widths = [12, 36, 7, 12, 14, 16, 12, 14]
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip(widths, "ABCDEFGH"):
        ws.column_dimensions[col].width = w
    for r in short_list(_mode(request)):
        if vfp:
            ws.append([r["code"], r["name"], r["unit"], float(r["current"]), float(r["reorder"]),
                       float(r["short"]), float(r["on_order"]) if r["on_order"] else None])
        else:
            ws.append([r["code"], r["name"], r["unit"], float(r["reorder"]), float(r["current"]),
                       float(r["pending_sale"]) if r["pending_sale"] else None, float(r["short"]),
                       float(r["on_order"]) if r["on_order"] else None])
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(size=12, bold=c.font.bold)
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Short_List_New.xlsx"'
    return response
