"""Reports > Sales Reports > Stock Statement (VFP "STOCK STATEMENT" form,
OLE-Excel export block).

Driven by ``item`` directly, one row per item code, NOT by the VFP
source's own ``partmst``-driven outer loop. That loop opens with
``sele partmst`` then immediately reads ``mopbal=op_bal`` -- but the real
``part-mst`` table (apps.parts.models.Part's own table; note the real
table's hyphen) HAS NO ``op_bal`` column (confirmed against
INFORMATION_SCHEMA; it has ``op_stk``/``gopstk`` instead), while ``item``
has exactly that column plus every other field this report reads
(``rate``, ``pre_bal``, ``order``, ``dummy_stk``) under the same names the
VFP source uses unqualified. The outer loop also ends each pass with a
bare ``sele item`` + ``skip`` -- advancing ITEM in lockstep with
PARTMST's own row position, assuming the two tables have identical row
order/count, with no ``seek`` ever lining item up to the current part
code. Both of these only make sense if the intended driving table was
``item`` all along; porting the partmst/item parallel-position construct
literally would silently misattribute rate/description/pending-order
data between unrelated item rows whenever the two tables' row counts or
order ever diverge. Flagged here rather than guessed past silently.

``order1``/``mrn_dat``/``cn_dat`` are opened by the VFP source but never
read by its own Excel block (``order1``: Pending Sale Order Qty actually
comes straight from ``item.order``, a precomputed stored field, not a
live ``order1`` sum; ``cn_dat`` rows ARE appended into the working
cursor, but the block's own ``do case`` only ever branches on
``stts`` IN ('P','A','S','I','M') -- ``cn_dat``'s own stamped stts='R' never
matches any branch, so those rows are mathematically inert) -- all three
left unread here too, not guessed into use.

The working dataset ("bb" in the VFP source) is ``i_master`` rows (the
same finished-goods GST table apps.gst_invoice.hsn_report reads) PLUS
``chl_dat`` (Returnable Challan despatch) rows re-stamped ``stts='S'`` --
counted as sales by quantity, but at a BLANK rate (the VFP source's own
``appe blan`` clears every field before copying over only
p_code/i_code/date/bill_no/qty/stts; ``chl_dat.rate`` is never copied), so
challan-derived "sales" contribute quantity but zero amount everywhere
rate is read off the row itself. Preserved exactly, not "fixed" to use
``chl_dat.rate``.

Column labels and what they actually compute diverge in several places in
the VFP source itself -- preserved verbatim (not relabeled) since this is
a direct port, each one confirmed against the source rather than assumed:
  * "Sale Amt @LP" (column G) is actually the booked sale amount at each
    line's own rate (``msamt`` -- zero for challan-derived lines, per
    above); "Sale Amt" (column H) is actually sale+issue+sample quantity
    priced at the ITEM's current list rate (``msamt1``) -- the two
    labels are swapped from what they compute.
  * "Last.Yr 30 days Avg. Sale" (column N) is simply ``item.pre_bal``,
    not any computed average.
  * "Shortage" (column U) is always 0 -- the VFP source's own live
    if/else only ever sets ``mexcess``; ``mshort`` is initialised to 0
    and never reassigned outside a commented-out dead branch.
  * The TOTAL row's own "Receipt Amt @LP" cell is actually the summed
    RECEIPT QUANTITY (``tramt`` accumulates ``mrqty``, not ``mramt``).
A row is skipped entirely when opening balance, receipt, sale, issue AND
sample quantities are all exactly zero -- same "nothing happened" rule
the VFP source applies before printing anything for that item.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import connections
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

RECEIPT_CODES = {"P", "A"}


def _rows(sql, params=None):
    with connections["legacy"].cursor() as cur:
        cur.execute(sql, params or [])
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]


def _clean(value):
    return (value or "").strip() if isinstance(value, str) or value is None else value


def _d(value):
    try:
        return Decimal(str(value if value is not None else 0))
    except Exception:
        return Decimal("0")


def _fy_start(today):
    return date(today.year if today.month >= 4 else today.year - 1, 4, 1)


def _round0(value):
    return value.quantize(Decimal("1"))


def stock_statement_rows(start, end):
    mavg = Decimal((end - start).days) or Decimal("1")

    i_master_lines = _rows("SELECT i_code, date, qty, rate, stts FROM i_master WHERE date <= %s", [end])
    chl_lines = _rows("SELECT i_code, date, qty FROM chl_dat WHERE date <= %s", [end])

    by_item = {}
    for l in i_master_lines:
        code = _clean(l["i_code"]) or ""
        by_item.setdefault(code, []).append({
            "date": l["date"], "qty": _d(l["qty"]), "rate": _d(l["rate"]), "stts": _clean(l["stts"]) or "",
        })
    for l in chl_lines:
        code = _clean(l["i_code"]) or ""
        by_item.setdefault(code, []).append({
            "date": l["date"], "qty": _d(l["qty"]), "rate": Decimal("0"), "stts": "S",
        })

    items = _rows(
        "SELECT i_code, i_name, op_bal, rate, pre_bal, [order], dummy_stk FROM item ORDER BY i_code"
    )
    gip = {}
    for r in _rows("SELECT i_code, SUM(qty) AS qty FROM new_job WHERE [close] <> 'Y' GROUP BY i_code"):
        gip[_clean(r["i_code"]) or ""] = _d(r["qty"])

    rows = []
    total_receipt_qty = total_sale_amt = total_sale_amt_lp = total_issue_amt = total_sample_amt = Decimal("0")

    for item in items:
        code = _clean(item["i_code"]) or ""
        item_rate = _d(item["rate"])
        pre_bal = _d(item["pre_bal"])
        pending_order = _d(item["order"])
        lines = by_item.get(code, [])

        mopbal = _d(item["op_bal"])
        for l in lines:
            if l["date"] and l["date"] < start:
                if l["stts"] in RECEIPT_CODES:
                    mopbal += l["qty"]
                elif l["stts"] in ("S", "I", "M"):
                    mopbal -= l["qty"]

        receipt_qty = receipt_amt = sale_qty = sale_amt = issue_qty = issue_amt = Decimal("0")
        sample_qty = sample_amt = Decimal("0")
        for l in lines:
            if not l["date"] or l["date"] < start or l["date"] > end:
                continue
            if l["stts"] in RECEIPT_CODES:
                receipt_qty += l["qty"]
                receipt_amt += _round0(l["qty"] * item_rate)
            elif l["stts"] == "S":
                sale_qty += l["qty"]
                sale_amt += _round0(l["qty"] * l["rate"])
            elif l["stts"] == "I":
                issue_qty += l["qty"]
                issue_amt += _round0(l["qty"] * l["rate"])
            elif l["stts"] == "M":
                sample_qty += l["qty"]
                sample_amt += l["qty"] * l["rate"]

        if mopbal == 0 and receipt_qty == 0 and sale_qty == 0 and issue_qty == 0 and sample_qty == 0:
            continue

        balance_qty = mopbal + receipt_qty - sale_qty - issue_qty - sample_qty
        sale_avg = _round0(((sale_qty + issue_qty + sample_qty) / mavg) * 30)
        sale_amt_lp = (sale_qty + issue_qty + sample_qty) * item_rate
        goods_in_process = gip.get(code, Decimal("0"))

        if sale_avg > pre_bal:
            net = sale_avg + pending_order - balance_qty - goods_in_process
        else:
            net = pre_bal + pending_order - balance_qty - goods_in_process
        if net > 0:
            reqd_qty, excess = net, Decimal("0")
        else:
            reqd_qty, excess = Decimal("0"), abs(net)

        rows.append({
            "item_code": code, "item_name": _clean(item["i_name"]) or "",
            "opening": f"{mopbal:.2f}", "receipt_qty": f"{receipt_qty:.2f}", "receipt_amt_lp": f"{receipt_amt:.2f}",
            "sale_qty": f"{sale_qty:.2f}", "sale_amt_lp": f"{sale_amt:.2f}", "sale_amt": f"{sale_amt_lp:.2f}",
            "issued_qty": f"{issue_qty:.2f}", "issued_amt": f"{issue_amt:.2f}",
            "sample_qty": f"{sample_qty:.2f}", "sample_amt": f"{sample_amt:.2f}",
            "cur_yr_avg_sale": f"{sale_avg:.2f}", "last_yr_avg_sale": f"{pre_bal:.2f}",
            "goods_in_process": f"{goods_in_process:.2f}", "closing_bal": f"{balance_qty:.2f}",
            "pending_order_qty": f"{pending_order:.2f}", "reqd_qty": f"{reqd_qty:.2f}",
            "excess": f"{excess:.2f}", "shortage": "0.00",
        })
        total_receipt_qty += receipt_qty
        total_sale_amt_lp += sale_amt
        total_sale_amt += sale_amt_lp
        total_issue_amt += issue_amt
        total_sample_amt += sample_amt

    totals = {
        "receipt_qty": f"{total_receipt_qty:.2f}", "sale_amt_lp": f"{total_sale_amt_lp:.2f}",
        "sale_amt": f"{total_sale_amt:.2f}", "issue_amt": f"{total_issue_amt:.2f}",
        "sample_amt": f"{total_sample_amt:.2f}",
    }
    return rows, totals


class StockStatementView(LoginRequiredMixin, TemplateView):
    template_name = "items/stock_statement.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": _fy_start(today).isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), _fy_start(today))
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def ss_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _params(request)
    if start >= end:
        return JsonResponse({"message": "Start Date must be before End Date."}, status=400)

    rows, totals = stock_statement_rows(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append([f"STOCK STATEMENT FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])

    header = ["Code", "Item Description", "Opening", "Receipt Qty", "Receipt Amt @LP", "Sale Qty",
               "Sale Amt @LP", "Sale Amt", "Issued Qty", "Issued Amt", "Sample Qty", "Sample Amt",
               "Cur.Yr 30 days Avg. Sale", "Last.Yr 30 days Avg. Sale", "Goods in Process", "Closing Bal",
               "Pending Sale Order Qty", "Reqd Qty", "Planned Qty", "Excess", "Shortage"]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
        c.alignment = c.alignment.copy(wrap_text=True)
    for col in "CDEFGHIJKLMNOPQRSTU":
        ws.column_dimensions[col].width = 10
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["N"].width = 15

    for r in rows:
        ws.append([
            r["item_code"], r["item_name"], float(r["opening"]), float(r["receipt_qty"]),
            float(r["receipt_amt_lp"]), float(r["sale_qty"]), float(r["sale_amt_lp"]), float(r["sale_amt"]),
            float(r["issued_qty"]), float(r["issued_amt"]), float(r["sample_qty"]), float(r["sample_amt"]),
            float(r["cur_yr_avg_sale"]), float(r["last_yr_avg_sale"]), float(r["goods_in_process"]),
            float(r["closing_bal"]), float(r["pending_order_qty"]), float(r["reqd_qty"]), None,
            float(r["excess"]), float(r["shortage"]),
        ])

    ws.append(["", "TOTAL", "", "", float(totals["receipt_qty"]), "", float(totals["sale_amt_lp"]),
               float(totals["sale_amt"]), "", float(totals["issue_amt"]), "", float(totals["sample_amt"])])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Stock_Statement_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
