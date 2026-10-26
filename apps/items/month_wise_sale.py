"""Reports > Sales Reports > Month Wise Sale (VFP "PRODUCTWISE MONTHLY SALE" /
MWSALE.PRN form, OLE-Excel export block).

One row per item sold (``i_master.stts = 'S'``, the "sold" line-status flag)
within Start Date/End Date, read off ``i_master.date`` (the line's own
date, not any bill date -- same convention as apps.gst_invoice.hsn_report).

Quantities are bucketed by CALENDAR month number (``MONTH(date)``) into a
FISCAL-year column order: Apr..Dec (months 4-12) then Jan..Mar (months
1-3) -- exactly the VFP source's own ``for i=4 to 12`` then ``for i=1 to
3`` loops. This only reads sensibly when Start Date/End Date span one
fiscal year; like the VFP source itself, nothing here enforces that.

Average divides the item's total quantity by a month-count the VFP source
reads from ``thisform.tmonth.value`` -- a plain number box with no visible
default anywhere in the pasted form init, so this port exposes it as its
own "No. of Months" field, defaulting to the whole months between Start
Date and End Date (at least 1, so Average never divides by zero).

Pending Order sums ``order1.stock - order1.issue`` for that item across
every row whose ``adj <> 'Y'``, same as the VFP source. Current Stock
reads ``stock.stock`` for that item -- the VFP source's own ``scan`` loop
overwrites the cell on every matching row rather than summing, so (same as
it) only the LAST matching ``stock`` row for that item wins; real data
carries at most one ``stock`` row per item, so this is never ambiguous in
practice.

Previous Balance / Target (``item.pre_bal``/``item.target``) are read by
the PRN report but commented out of the Excel block itself -- left out
here too, not guessed back in.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import connections
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

FISCAL_MONTHS = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
MONTH_LABELS = ["Apr", "May", "Jun", "July", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]


def _rows(sql, params=None):
    with connections["legacy"].cursor() as cur:
        cur.execute(sql, params or [])
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]


def _clean(value):
    return (value or "").strip() if isinstance(value, str) or value is None else value


def _d(value):
    try:
        return Decimal(str(value if value is not None else 0))
    except Exception:
        return Decimal("0")


def _month_count(start, end):
    months = (end.year - start.year) * 12 + (end.month - start.month) + 1
    return max(months, 1)


def month_wise_sale_rows(start, end, months):
    lines = _rows(
        "SELECT i_code, date, qty FROM i_master WHERE date >= %s AND date <= %s AND stts = 'S' ORDER BY i_code",
        [start, end],
    )
    by_item = {}
    item_codes = []
    for l in lines:
        code = _clean(l["i_code"]) or ""
        if code not in by_item:
            by_item[code] = [Decimal("0")] * 12
            item_codes.append(code)
        by_item[code][l["date"].month - 1] += _d(l["qty"])

    items = {_clean(r["i_code"]): r for r in _rows(
        "SELECT i_code, i_name FROM item WHERE i_code IN (%s)" % ",".join(["%s"] * len(item_codes)),
        item_codes,
    )} if item_codes else {}

    pending = {}
    for r in _rows(
        "SELECT i_code, SUM(stock - issue) AS pending FROM order1 WHERE adj <> 'Y' AND i_code IN (%s) "
        "GROUP BY i_code" % ",".join(["%s"] * len(item_codes)),
        item_codes,
    ) if item_codes else []:
        pending[_clean(r["i_code"])] = _d(r["pending"])

    current_stock = {}
    for r in _rows(
        "SELECT i_code, stock FROM stock WHERE i_code IN (%s)" % ",".join(["%s"] * len(item_codes)),
        item_codes,
    ) if item_codes else []:
        current_stock[_clean(r["i_code"])] = _d(r["stock"])  # last row wins, same as VFP source

    rows = []
    for code in item_codes:
        months_qty = by_item[code]
        fiscal_qty = [months_qty[m - 1] for m in FISCAL_MONTHS]
        total_qty = sum(fiscal_qty)
        item = items.get(code)
        rows.append({
            "item_code": code, "item_name": _clean(item["i_name"]) if item else "",
            "months": [f"{q:.2f}" for q in fiscal_qty], "total": f"{total_qty:.2f}",
            "average": f"{(total_qty / months):.2f}",
            "pending_order": f"{pending.get(code, Decimal('0')):.2f}",
            "current_stock": f"{current_stock.get(code, Decimal('0')):.2f}",
        })
    return rows


class MonthWiseSaleView(LoginRequiredMixin, TemplateView):
    template_name = "items/month_wise_sale.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        start = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
        context["initial_data"] = {
            "start_date": start.isoformat(), "end_date": today.isoformat(),
            "months": _month_count(start, today),
        }
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    fy_start = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
    start = parse(request.GET.get("start_date"), fy_start)
    end = parse(request.GET.get("end_date"), today)
    try:
        months = max(int(request.GET.get("months") or _month_count(start, end)), 1)
    except ValueError:
        months = _month_count(start, end)
    return start, end, months


@login_required
@require_GET
def mws_export(request):
    from io import BytesIO

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end, months = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)

    rows = month_wise_sale_rows(start, end, months)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"PRODUCTWISE MONTHLY SALE FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])

    header = ["Item Description"] + MONTH_LABELS + ["Total", "Average", "Pend Ord", "Cur Stk"]
    widths = [25] + [8] * 12 + [8, 8, 8, 8]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True)
    for w, col in zip(widths, "ABCDEFGHIJKLMNOPQ"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append(
            [r["item_name"]] + [float(q) for q in r["months"]] +
            [float(r["total"]), float(r["average"]), float(r["pending_order"]), float(r["current_stock"])]
        )

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Month_Wise_Sale_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
