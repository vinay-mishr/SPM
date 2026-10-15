"""Creditor Outstanding > Payment Report: payments made between two dates, optionally for one vendor."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from apps.parties.models import Party

from .models import PaymentDetail


def _parse(value, default):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return default


def payment_rows(start, end, party_code=None):
    qs = PaymentDetail.objects.filter(date__gte=start, date__lte=end)
    if party_code:
        qs = qs.filter(p_code=party_code)
    groups = {}
    for p in qs.values("p_code", "nar", "date", "amount", "c_dis", "int_amt"):
        key = (p["p_code"] or "", p["nar"] or "")
        g = groups.setdefault(key, {"date": p["date"], "amount": Decimal("0"), "cdis": Decimal("0"),
                                    "int": Decimal("0"), "bills": 0})
        g["amount"] += Decimal(str(p["amount"] or 0))
        g["cdis"] += Decimal(str(p["c_dis"] or 0))
        g["int"] += Decimal(str(p["int_amt"] or 0))
        g["bills"] += 1
        if p["date"] and (g["date"] is None or p["date"] > g["date"]):
            g["date"] = p["date"]
    names = dict(Party.objects.values_list("party_code", "name"))
    out = []
    for (code, chq), g in groups.items():
        out.append({"date": g["date"].isoformat() if g["date"] else "", "party_code": code,
                    "party_name": (names.get(code) or "").strip(), "chq_no": chq, "bills": g["bills"],
                    "c_dis": f"{g['cdis']:.2f}", "interest": f"{g['int']:.2f}", "amount": f"{g['amount']:.2f}"})
    out.sort(key=lambda r: (r["date"], r["party_name"]), reverse=True)
    total = sum((Decimal(r["amount"]) for r in out), Decimal("0"))
    return out, f"{total:.2f}"


class PaymentReportView(LoginRequiredMixin, TemplateView):
    template_name = "creditor_outstanding/payment_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.replace(day=1).isoformat(), "end_date": today.isoformat()}
        return context


@login_required
@require_GET
def pr_data(request):
    start = _parse(request.GET.get("start_date"), date.today().replace(day=1))
    end = _parse(request.GET.get("end_date"), date.today())
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)
    rows, total = payment_rows(start, end, (request.GET.get("party_code") or "").strip().upper() or None)
    return JsonResponse({"rows": rows, "total": total})


@login_required
@require_GET
def pr_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    start = _parse(request.GET.get("start_date"), date.today().replace(day=1))
    end = _parse(request.GET.get("end_date"), date.today())
    rows, total = payment_rows(start, end, (request.GET.get("party_code") or "").strip().upper() or None)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"PAYMENT REPORT FROM {start:%d/%m/%Y} TO {end:%d/%m/%Y}"])
    ws.append([])
    ws.append(["DATE", "VENDOR", "PARTY NAME", "CHQ/NEFT/RCPT NO", "BILLS", "C.DIS", "INTEREST", "AMOUNT"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([12, 10, 34, 18, 8, 12, 12, 14], "ABCDEFGH"):
        ws.column_dimensions[col].width = w
    for r in rows:
        ws.append([r["date"], r["party_code"], r["party_name"], r["chq_no"], r["bills"],
                   float(r["c_dis"]), float(r["interest"]), float(r["amount"])])
    ws.append([None, None, None, None, None, None, "TOTAL", float(total)])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Payment_Report.xlsx"'
    return response
