"""Creditor Outstanding > Outstanding Statement (VFP CROUTSTAND).

Bills in cr_out that are not fully paid (pmt_mad not 'Y') and dated on or before the run date,
grouped by vendor. Excel adds aging buckets by days between bill date and run date.
"""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from apps.parties.models import Party

from .models import CreditorBill


def _d(v):
    return Decimal(str(v or 0))


def statement(run_date, party_code=None):
    qs = CreditorBill.objects.exclude(pmt_mad="Y").filter(inv_dt__lte=run_date).order_by("p_code", "inv_dt")
    if party_code:
        qs = qs.filter(p_code=party_code)
    bills = list(qs.values("p_code", "inv_no", "inv_dt", "div", "mrn_no", "cl_date", "amount", "tds_amt",
                           "pmt_rcd", "dr_amt", "c_dis"))
    parties = {c: (n, city) for c, n, city in Party.objects.values_list("party_code", "name", "city")}
    groups = {}
    for b in bills:
        bal = _d(b["amount"]) - _d(b["pmt_rcd"]) - _d(b["dr_amt"]) - _d(b["c_dis"]) - _d(b["tds_amt"])
        if bal <= 0:
            continue
        groups.setdefault(b["p_code"], []).append((b, bal))
    out = []
    for code, rows in groups.items():
        name, city = parties.get(code, ("", ""))
        lines, total = [], Decimal("0")
        buckets = [Decimal("0")] * 4
        for b, bal in rows:
            days = (run_date - b["inv_dt"]).days if b["inv_dt"] else 0
            idx = 0 if days <= 45 else 1 if days <= 60 else 2 if days <= 90 else 3
            buckets[idx] += bal
            total += bal
            lines.append({
                "mrn_no": (b["mrn_no"] or "").strip(),
                "cl_date": b["cl_date"].isoformat() if b["cl_date"] else "",
                "bucket": idx,
                "bill_no": f"{(b['div'] or '').strip()}{(b['inv_no'] or '').strip()}",
                "bill_date": b["inv_dt"].isoformat() if b["inv_dt"] else "",
                "bill_amount": f"{_d(b['amount']):.2f}", "tds": f"{_d(b['tds_amt']):.2f}",
                "paid": f"{_d(b['pmt_rcd']):.2f}", "dr_note": f"{_d(b['dr_amt']):.2f}",
                "c_dis": f"{_d(b['c_dis']):.2f}", "balance": f"{bal:.2f}", "days": days,
            })
        out.append({"party_code": code, "party_name": (name or "").strip(), "city": (city or "").strip(),
                    "lines": lines, "total": f"{total:.2f}", "buckets": [f"{x:.2f}" for x in buckets]})
    out.sort(key=lambda x: x["party_name"])
    return out


def _parse(value, default):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return default


class OutstandingStatementView(LoginRequiredMixin, TemplateView):
    template_name = "creditor_outstanding/outstanding_statement.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


@login_required
@require_GET
def statement_data(request):
    run = _parse(request.GET.get("run_date"), date.today())
    party = (request.GET.get("party_code") or "").strip().upper() or None
    return JsonResponse({"groups": statement(run, party)})


@login_required
@require_GET
def statement_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    run = _parse(request.GET.get("run_date"), date.today())
    party = (request.GET.get("party_code") or "").strip().upper() or None
    groups = statement(run, party)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"CREDITOR OUTSTANDING STATEMENT AS ON {run:%d/%m/%Y}"])
    ws.append([])
    headers = ["PARTY NAME", "CITY", "MRN NO.", "CL. DATE", "BILL NO.", "BILL DATE", "BILL AMT.", "TDS AMT.",
               "AMT.PAID", "DR.NOTE AMT.", "C.DIS", "BAL.AMT", "BELOW 45 DAYS", "BELOW 45 TO 60",
               "BELOW 60 TO 90", "ABOVE 90"]
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([30, 16, 10, 12, 14, 12, 12, 10, 12, 12, 10, 12, 13, 13, 13, 12], "ABCDEFGHIJKLMNOP"):
        ws.column_dimensions[col].width = w
    grand = [Decimal("0")] * 5
    for g in groups:
        for ln in g["lines"]:
            buckets = [float(ln["balance"]) if ln["bucket"] == i else None for i in range(4)]
            ws.append([g["party_name"], g["city"], ln["mrn_no"], ln["cl_date"], ln["bill_no"], ln["bill_date"],
                       float(ln["bill_amount"]), float(ln["tds"]), float(ln["paid"]), float(ln["dr_note"]),
                       float(ln["c_dis"]), float(ln["balance"])] + buckets)
        bk = g["buckets"]
        ws.append([g["party_name"], "", None, None, "TOTAL", None, None, None, None, None, None,
                   float(g["total"])] + [float(x) for x in bk])
        grand[0] += Decimal(g["total"])
        for i in range(4):
            grand[i + 1] += Decimal(bk[i])
    ws.append(["GRAND TOTAL"] + [None] * 10 + [float(grand[0])] + [float(x) for x in grand[1:]])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Creditor_Outstanding_Statement.xlsx"'
    return response


def _credit_days():
    from django.db import connections

    cur = connections["legacy"].cursor()
    cur.execute("SELECT code, cr_days FROM cust_mst")
    return {(c or "").strip(): int(d or 0) for c, d in cur.fetchall()}


def report(run_date, party_code=None):
    days_map = _credit_days()
    groups = statement(run_date, party_code)
    for g in groups:
        credit = days_map.get(g["party_code"], 0)
        for ln in g["lines"]:
            due = ln["days"] - credit
            ln["due_days"] = due
    return groups


class OutstandingReportView(LoginRequiredMixin, TemplateView):
    template_name = "creditor_outstanding/outstanding_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["initial_data"] = {"run_date": date.today().isoformat()}
        return context


@login_required
@require_GET
def report_data(request):
    run = _parse(request.GET.get("run_date"), date.today())
    party = (request.GET.get("party_code") or "").strip().upper() or None
    return JsonResponse({"groups": report(run, party)})


@login_required
@require_GET
def report_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    run = _parse(request.GET.get("run_date"), date.today())
    party = (request.GET.get("party_code") or "").strip().upper() or None
    groups = report(run, party)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([settings.BRAND["name"].upper()])
    ws.append([f"CREDITOR OUTSTANDING REPORT AS ON {run:%d/%m/%Y}"])
    ws.append([])
    headers = ["PARTY NAME", "CITY", "MRN NO.", "BILL NO.", "BILL DATE", "BILL AMT.", "TDS AMT.", "AMT.PAID",
               "DR.NOTE AMT.", "C.DIS", "BAL.AMT", "DEL.DATE"]
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for w, col in zip([30, 16, 10, 14, 12, 12, 10, 12, 12, 10, 12, 10], "ABCDEFGHIJKL"):
        ws.column_dimensions[col].width = w
    grand = Decimal("0")
    for g in groups:
        for ln in g["lines"]:
            ws.append([g["party_name"], g["city"], ln["mrn_no"], ln["bill_no"], ln["bill_date"],
                       float(ln["bill_amount"]), float(ln["tds"]), float(ln["paid"]), float(ln["dr_note"]),
                       float(ln["c_dis"]), float(ln["balance"]), ln["due_days"]])
        ws.append([g["party_name"], "TOTAL", None, None, None, None, None, None, None, None, float(g["total"])])
        grand += Decimal(g["total"])
    ws.append(["GRAND TOTAL", None, None, None, None, None, None, None, None, None, float(grand)])
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Creditor_Outstanding_Report.xlsx"'
    return response
