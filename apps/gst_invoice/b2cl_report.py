"""GST Reports > Business to Consumer Large (VFP "B2CL" form).

Direct port of the pasted VFP B2CL(5) Excel-export block. Same ``bill``/
``i_master`` tables and ``bill_type = 'G'`` convention as
apps.gst_invoice.b2b_report (see its own docstring for the shared pieces).

A bill counts as B2CL when it's a GST bill, interstate (``cstat = 'I'``,
the same flag apps.gst_invoice carries through to ``i_master.cstat``),
REGULAR category (``cat_cd = '01'``, see apps.gst_invoice.b2b_report's own
note on this code) and its net amount is >= Rs. 2,50,000 -- preserved
exactly from the VFP source's own filter, including that it does NOT check
``gst_no`` here (unlike B2B's own filter).

Each matching bill contributes one Excel row PER TAX RATE, grouped the same
"do while gst==mgst" way as B2B -- but unlike B2B, Taxable Value here sums
plain ``amount`` (no ``dis_amt`` subtracted), matching the VFP source's own
``mamt=mamt+amount`` line exactly.

The VFP source computes a "d" (distinct-GSTIN) count but never writes it to
the sheet (its own ``OleApp.Cells(3,1).Value=d`` line is commented out,
overwritten by ``=c`` right after) -- so it's not reproduced here either.
Cess has no column read anywhere in the VFP source, so it's always 0.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import BILL_TYPE, _clean, _d, _rows

CATEGORY_CODE = "01"  # REGULAR -- see apps.gst_invoice.b2b_report
MIN_INVOICE_VALUE = Decimal("250000.00")


def b2cl_rows(start, end):
    bills = _rows(
        "SELECT bill_no, b_date, n_amount, amt_tf, stname FROM bill "
        "WHERE bill_type = %s AND cstat = 'I' AND cat_cd = %s AND n_amount >= %s "
        "AND b_date >= %s AND b_date <= %s ORDER BY b_date, bill_no",
        [BILL_TYPE, CATEGORY_CODE, MIN_INVOICE_VALUE, start, end],
    )

    rows = []
    invoice_total = taxable_total = Decimal("0")

    for b in bills:
        bill_no = _clean(b["bill_no"])
        invoice_value = _d(b["n_amount"]) - _d(b["amt_tf"])
        invoice_total += invoice_value

        lines = _rows(
            "SELECT gst, amount FROM i_master WHERE bill_no = %s AND bill_type = %s AND gst > 0 "
            "ORDER BY gst",
            [bill_no, BILL_TYPE],
        )
        by_rate = {}
        for l in lines:
            rate = _d(l["gst"])
            by_rate[rate] = by_rate.get(rate, Decimal("0")) + _d(l["amount"])

        for rate in sorted(by_rate):
            taxable_value = by_rate[rate]
            taxable_total += taxable_value
            rows.append({
                "bill_no": bill_no, "bill_date": b["b_date"].isoformat() if b["b_date"] else "",
                "invoice_value": f"{invoice_value:.2f}", "place_of_supply": _clean(b["stname"]) or "",
                "rate": f"{rate:.2f}", "taxable_value": f"{taxable_value:.2f}",
            })

    totals = {
        "invoices": len(bills), "invoice_value": f"{invoice_total:.2f}",
        "taxable_value": f"{taxable_total:.2f}",
    }
    return rows, totals


class B2clReportView(LoginRequiredMixin, TemplateView):
    template_name = "gst_invoice/b2cl_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def b2cl_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)

    rows, totals = b2cl_rows(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(["Summary For B2CL(5)"])
    ws["A1"].font = Font(bold=True, size=11)
    ws.append(["No. of Invoices", "", "Total Invoice Value", "", "", "Total Taxable Value", "Total Cess"])
    for c in ws[2]:
        c.font = Font(bold=True, size=11)
    ws.append([totals["invoices"], "", float(totals["invoice_value"]), "", "", float(totals["taxable_value"]), 0])
    ws.append([])

    header = ["Invoice Number", "Invoice Date", "Invoice Value", "Place of Supply", "RATE", "Taxable Value",
               "Cess Amount", "E-COMMERCE GSTIN"]
    widths = [18, 14, 21, 30, 6, 21, 13, 25]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.alignment = c.alignment.copy(horizontal="center", vertical="center", wrap_text=True)
        c.font = Font(bold=True, size=11)
    for w, col in zip(widths, "ABCDEFGH"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([
            r["bill_no"], date.fromisoformat(r["bill_date"]) if r["bill_date"] else None,
            float(r["invoice_value"]), r["place_of_supply"], float(r["rate"]), float(r["taxable_value"]),
            None, "",
        ])
        ws[f"B{ws.max_row}"].number_format = "dd-mmm-yyyy"
        for col in ("C", "E", "F"):
            ws[f"{col}{ws.max_row}"].number_format = "0.00"

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="B2CL_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
