"""GST Reports > HSN Wise Report (VFP "HSN" form).

Direct port of the pasted VFP HSN(12) Excel-export block. Reads
``i_master`` only (no ``bill`` join), filtered by the line's OWN ``date``
column (not ``bill.b_date`` -- unlike apps.gst_invoice.b2b_report/
b2cl_report/b2cs_report, which all filter on the bill's date) and
``bill_type = 'G'``, with no GST-rate filter (unlike those three reports,
this one includes GST = 0 lines too -- the VFP source's own ``SET FILTER``
line has no ``gst > 0`` clause).

One Excel row per (HS code, unit) pair -- the VFP source's own nested
"do while hs_code==micode" / "do while hs_code==micode AND unit==munit"
loops, walking the ``HSUNIT`` index (hs_code then unit). "No. of HSN"
(summary ``C``) counts distinct HS codes only, incremented once per outer
loop pass, not once per (HS code, unit) row.

Description (column B) is never assigned a value anywhere in the VFP
source despite the header -- preserved as blank here too, not guessed from
apps.items.Item.

Total Value sums ``amount + cgst_amt + sgst_amt + igst_amt``; Taxable Value
sums plain ``amount`` -- both exactly as the VFP source computes them, not
"corrected". Cess has no column read anywhere in the VFP source, so it's
always 0.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import BILL_TYPE, _clean, _d, _rows


def hsn_rows(start, end):
    lines = _rows(
        "SELECT hs_code, unit, qty, amount, cgst_amt, sgst_amt, igst_amt FROM i_master "
        "WHERE bill_type = %s AND date >= %s AND date <= %s ORDER BY hs_code, unit",
        [BILL_TYPE, start, end],
    )

    groups = {}
    hs_codes = []
    for l in lines:
        hs_code = _clean(l["hs_code"]) or ""
        unit = _clean(l["unit"]) or ""
        if hs_code not in groups:
            groups[hs_code] = {}
            hs_codes.append(hs_code)
        bucket = groups[hs_code].setdefault(unit, {
            "qty": Decimal("0"), "amount": Decimal("0"),
            "cgst": Decimal("0"), "sgst": Decimal("0"), "igst": Decimal("0"),
        })
        bucket["qty"] += _d(l["qty"])
        bucket["amount"] += _d(l["amount"])
        bucket["cgst"] += _d(l["cgst_amt"])
        bucket["sgst"] += _d(l["sgst_amt"])
        bucket["igst"] += _d(l["igst_amt"])

    rows = []
    total_value = total_taxable = total_igst = total_cgst = total_sgst = Decimal("0")
    for hs_code in hs_codes:
        for unit in groups[hs_code]:
            b = groups[hs_code][unit]
            value = b["amount"] + b["cgst"] + b["sgst"] + b["igst"]
            total_value += value
            total_taxable += b["amount"]
            total_igst += b["igst"]
            total_cgst += b["cgst"]
            total_sgst += b["sgst"]
            rows.append({
                "hs_code": hs_code, "unit": unit, "qty": f"{b['qty']:.2f}",
                "value": f"{value:.2f}", "taxable_value": f"{b['amount']:.2f}",
                "igst": f"{b['igst']:.2f}", "cgst": f"{b['cgst']:.2f}", "sgst": f"{b['sgst']:.2f}",
            })

    totals = {
        "hsn_count": len(hs_codes), "value": f"{total_value:.2f}", "taxable_value": f"{total_taxable:.2f}",
        "igst": f"{total_igst:.2f}", "cgst": f"{total_cgst:.2f}", "sgst": f"{total_sgst:.2f}",
    }
    return rows, totals


class HsnReportView(LoginRequiredMixin, TemplateView):
    template_name = "gst_invoice/hsn_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def hsn_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)

    rows, totals = hsn_rows(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(["Summary For HSN(12)"])
    ws["A1"].font = Font(bold=True, size=11)
    ws.append(["No. of HSN", "", "", "", "Total Value", "Total Taxable Value", "Total Integrated Tax",
               "Total Central Tax", "Total State/UT Tax", "Total Cess"])
    for c in ws[2]:
        c.font = Font(bold=True, size=11)
    ws.append([totals["hsn_count"], "", "", "", float(totals["value"]), float(totals["taxable_value"]),
               float(totals["igst"]), float(totals["cgst"]), float(totals["sgst"]), 0])
    ws.append([])

    header = ["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Taxable Value",
               "Integrated Tax Amount", "Central Tax Amount", "State/UT Tax Amount", "Cess Amount"]
    widths = [20, 20, 20, 16, 12, 21, 25, 25, 25, 13]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.alignment = c.alignment.copy(horizontal="center", vertical="center", wrap_text=True)
        c.font = Font(bold=True, size=11)
    for w, col in zip(widths, "ABCDEFGHIJ"):
        ws.column_dimensions[col].width = w

    for r in rows:
        ws.append([
            r["hs_code"], "", r["unit"], float(r["qty"]), float(r["value"]), float(r["taxable_value"]),
            float(r["igst"]), float(r["cgst"]), float(r["sgst"]), None,
        ])
        for col in ("D", "E", "F", "G", "H", "I"):
            ws[f"{col}{ws.max_row}"].number_format = "0.00"

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="HSN_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
