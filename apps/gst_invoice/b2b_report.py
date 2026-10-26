"""GST Reports > Business to Business (VFP "B2B" form).

Direct port of the pasted VFP B2B(4) Excel-export block. Reuses the same
physical ``bill``/``i_master`` tables and ``bill_type = 'G'`` convention as
apps.gst_invoice.services (see its own module docstring).

A bill counts as B2B when it's a GST bill (``bill_type = 'G'``) for a
registered recipient (``gst_no`` not blank) in the "REGULAR" customer
category (``cat_cd = '01'`` -- the same code apps.parties.services.
GST_CATEGORY_CODES["REGULAR"] copies onto cust_mst, and onto the bill itself
at invoice time), dated within Start Date/End Date.

Each such bill contributes one Excel row PER TAX RATE it carries, not one
row per bill: its ``i_master`` lines (``bill_type = 'G'``, ``gst > 0``) are
grouped by rate (``gst``), and each group's own (amount - dis_amt) is summed
into that row's Taxable Value -- exactly the VFP source's own "do while
gst==mgst" grouping (its table is ordered by ``gst``, so equal rates for the
same bill are contiguous).

Place of Supply is read straight off ``bill.stname`` (not derived from a
state code) and Reverse Charge straight off ``bill.y_no`` -- both preserved
exactly as the VFP source reads them, not reinterpreted.

Summary (B2B(4)) counts, also preserved from the VFP source: No. of
Recipients is the count of DISTINCT gst_no among matching bills (not
invoices); No. of Invoices is the count of matching bills; Total Invoice
Value sums (n_amount - amt_tf) per bill; Total Taxable Value sums every
rate-group's own Taxable Value. Cess has no column read anywhere in the VFP
source, so it's always 0 here too.
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_GET
from django.views.generic import TemplateView

from .services import BILL_TYPE, _clean, _d, _rows

CATEGORY_CODE = "01"  # REGULAR -- apps.parties.services.GST_CATEGORY_CODES["REGULAR"]


def b2b_rows(start, end):
    bills = _rows(
        "SELECT bill_no, b_date, n_amount, amt_tf, stname, y_no, cat_na, gst_no FROM bill "
        "WHERE bill_type = %s AND cat_cd = %s AND gst_no IS NOT NULL AND LTRIM(RTRIM(gst_no)) <> '' "
        "AND b_date >= %s AND b_date <= %s ORDER BY b_date, bill_no",
        [BILL_TYPE, CATEGORY_CODE, start, end],
    )

    rows = []
    recipients = set()
    invoice_total = taxable_total = Decimal("0")

    for b in bills:
        bill_no = _clean(b["bill_no"])
        recipients.add(_clean(b["gst_no"]))
        invoice_value = _d(b["n_amount"]) - _d(b["amt_tf"])
        invoice_total += invoice_value

        lines = _rows(
            "SELECT gst, amount, dis_amt FROM i_master WHERE bill_no = %s AND bill_type = %s AND gst > 0 "
            "ORDER BY gst",
            [bill_no, BILL_TYPE],
        )
        by_rate = {}
        for l in lines:
            rate = _d(l["gst"])
            by_rate[rate] = by_rate.get(rate, Decimal("0")) + _d(l["amount"]) - _d(l["dis_amt"])

        for rate in sorted(by_rate):
            taxable_value = by_rate[rate]
            taxable_total += taxable_value
            rows.append({
                "gst_no": _clean(b["gst_no"]) or "", "bill_no": bill_no,
                "bill_date": b["b_date"].isoformat() if b["b_date"] else "",
                "invoice_value": f"{invoice_value:.2f}", "place_of_supply": _clean(b["stname"]) or "",
                "reverse_charge": _clean(b["y_no"]) or "", "invoice_type": _clean(b["cat_na"]) or "",
                "rate": f"{rate:.2f}", "taxable_value": f"{taxable_value:.2f}",
            })

    totals = {
        "recipients": len(recipients), "invoices": len(bills),
        "invoice_value": f"{invoice_total:.2f}", "taxable_value": f"{taxable_total:.2f}",
    }
    return rows, totals


class B2bReportView(LoginRequiredMixin, TemplateView):
    template_name = "gst_invoice/b2b_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        context["initial_data"] = {"start_date": today.isoformat(), "end_date": today.isoformat()}
        return context


def _params(request):
    def parse(value, default):
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            return default

    today = date.today()
    start = parse(request.GET.get("start_date"), today)
    end = parse(request.GET.get("end_date"), today)
    return start, end


@login_required
@require_GET
def b2b_export(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    start, end = _params(request)
    if start > end:
        return JsonResponse({"message": "Start Date cannot be later than End Date."}, status=400)

    rows, totals = b2b_rows(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.column_dimensions["A"].width = 25

    ws.append(["Summary For B2B(4)"])
    ws["A1"].font = Font(bold=True, size=11)
    ws.append(["No. of Recipients", "No. of Invoices", "", "Total Invoice Value", "", "", "", "", "",
               "Total Taxable Value", "Total Cess"])
    for c in ws[2]:
        c.font = Font(bold=True, size=11)
    ws.append([totals["recipients"], totals["invoices"], "", float(totals["invoice_value"]), "", "", "", "", "",
               float(totals["taxable_value"]), 0])
    ws.append([])

    header = ["GSTIN/UIN of Recipient", "Invoice Number", "Invoice Date", "Invoice Value", "Place of Supply",
               "Reverse Charge", "Invoice Type", "E-COMMERCE GSTIN", "RATE", "Taxable Value", "Cess Amount"]
    widths = [25, 18, 14, 21, 30, 18, 13, 25, 6, 21, 13]
    ws.append(header)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.alignment = c.alignment.copy(horizontal="center", vertical="center", wrap_text=True)
        c.font = Font(bold=True, size=11)
    for w, col in zip(widths, "ABCDEFGHIJK"):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["C"].number_format = "dd-mmm-yyyy"

    for r in rows:
        ws.append([
            r["gst_no"], r["bill_no"], date.fromisoformat(r["bill_date"]) if r["bill_date"] else None,
            float(r["invoice_value"]), r["place_of_supply"], r["reverse_charge"], r["invoice_type"], "",
            float(r["rate"]), float(r["taxable_value"]), None,
        ])
        for col in ("C", "D", "I", "J"):
            ws[f"{col}{ws.max_row}"].number_format = "0.00" if col != "C" else "dd-mmm-yyyy"

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="B2B_{start:%d-%m-%Y}_to_{end:%d-%m-%Y}.xlsx"'
    return response
